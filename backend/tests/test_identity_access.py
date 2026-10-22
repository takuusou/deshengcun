from __future__ import annotations

from contextlib import ExitStack
from datetime import datetime, timedelta, timezone
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook

from app import admin_auth, admin_routes, admin_services, db, main, schemas, services


class IdentityAccessIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.environment = patch.dict(
            os.environ,
            {"SENSITIVE_LOOKUP_HMAC_KEY_HEX": "84" * 32},
            clear=False,
        )
        cls.environment.start()
        cls.temp_dir = TemporaryDirectory()
        cls.original_db_path = db.DB_PATH
        cls.original_data_dir = db.DATA_DIR
        cls.original_db_backend = db.DB_BACKEND
        db.DB_PATH = Path(cls.temp_dir.name) / "identity-access.db"
        db.DATA_DIR = db.DB_PATH.parent
        db.DB_BACKEND = "sqlite"
        db.init_db()
        cls._seed_data()

    @classmethod
    def tearDownClass(cls) -> None:
        db.DB_PATH = cls.original_db_path
        db.DATA_DIR = cls.original_data_dir
        db.DB_BACKEND = cls.original_db_backend
        cls.temp_dir.cleanup()
        cls.environment.stop()

    @classmethod
    def _seed_data(cls) -> None:
        now = datetime.now(timezone.utc)
        now_text = now.isoformat()
        expires_at = (now + timedelta(hours=2)).isoformat()
        conn = db.get_connection()
        try:
            users = [
                (
                    "openid-a",
                    "王五",
                    "男",
                    "100****8888",
                    "phone-a",
                    "000000************",
                    services._id_hash("00000019900101123X"),
                    "id-a",
                    "第一村民小组",
                ),
                (
                    "openid-b",
                    "李四",
                    "女",
                    "101****9999",
                    "phone-b",
                    "000004*********",
                    services._id_hash("000004670401001"),
                    "id-b",
                    "第二村民小组",
                ),
            ]
            for user in users:
                conn.execute(
                    """
                    INSERT INTO users(
                        openid, real_name, gender, phone_masked, phone_ciphertext,
                        phone_bound, phone_verified_at, id_card_masked, id_card_hash,
                        id_card_encrypted, role, realname_status, auth_status,
                        roster_match_status, village_group, created_at, updated_at
                    )
                    VALUES(?, ?, ?, ?, ?, 1, ?, ?, ?, ?, 'villager', 'verified',
                           'verified', 'matched', ?, ?, ?)
                    """,
                    (
                        user[0],
                        user[1],
                        user[2],
                        user[3],
                        user[4],
                        now_text,
                        user[5],
                        user[6],
                        user[7],
                        user[8],
                        now_text,
                        now_text,
                    ),
                )
            unverified_user_id = conn.execute(
                """
                INSERT INTO users(openid, nickname, created_at, updated_at)
                VALUES('openid-unverified', '未实名测试用户', ?, ?)
                """,
                (now_text, now_text),
            ).lastrowid
            conn.execute(
                "INSERT INTO auth_sessions(user_id, token, created_at, expires_at) VALUES(1, ?, ?, ?)",
                (services._hash_session_token("villager-a-token"), now_text, expires_at),
            )
            conn.execute(
                "INSERT INTO auth_sessions(user_id, token, created_at, expires_at) VALUES(2, ?, ?, ?)",
                (services._hash_session_token("villager-b-token"), now_text, expires_at),
            )
            conn.execute(
                "INSERT INTO auth_sessions(user_id, token, created_at, expires_at) VALUES(?, ?, ?, ?)",
                (
                    unverified_user_id,
                    services._hash_session_token("unverified-villager-token"),
                    now_text,
                    expires_at,
                ),
            )
            conn.execute(
                """
                INSERT INTO villager_requests(
                    user_id, request_type, title, content, contact_phone, reporter_name,
                    reporter_real_name, reporter_id_card_hash, created_at, updated_at
                )
                VALUES(1, 'service', 'A的事项', '本人事项', '10000138888', '王五', '王五', ?, ?, ?)
                """,
                (services._id_hash("00000019900101123X"), now_text, now_text),
            )
            conn.execute(
                """
                INSERT INTO villager_requests(
                    user_id, request_type, title, content, contact_phone, reporter_name,
                    reporter_real_name, reporter_id_card_hash, created_at, updated_at
                )
                VALUES(2, 'service', 'B的事项', '他人事项', '10100139999', '李四', '李四', ?, ?, ?)
                """,
                (services._id_hash("000004670401001"), now_text, now_text),
            )
            service_id = conn.execute(
                """
                INSERT INTO gov_services(service_key, title, created_at, updated_at)
                VALUES('certificate-test', '开具证明测试', ?, ?)
                """,
                (now_text, now_text),
            ).lastrowid
            conn.execute(
                """
                INSERT INTO gov_service_applications(
                    service_id, user_id, service_key, service_title, applicant_name,
                    applicant_phone, extra_fields, created_at, updated_at
                )
                VALUES(?, 1, 'certificate-test', '开具证明测试', '王五',
                       '10000138888', ?, ?, ?)
                """,
                (service_id, '{"id_card":"00000019900101123X"}', now_text, now_text),
            )
            conn.execute(
                """
                INSERT INTO gov_service_applications(
                    service_id, user_id, service_key, service_title, applicant_name,
                    applicant_phone, extra_fields, created_at, updated_at
                )
                VALUES(?, 2, 'certificate-test', '开具证明测试', '李四',
                       '10100139999', ?, ?, ?)
                """,
                (service_id, '{"id_card":"000004670401001"}', now_text, now_text),
            )
            authorization_id = conn.execute(
                """
                INSERT INTO cadre_authorizations(
                    real_name, office_title, role_level, can_export_full_id, status,
                    created_at, updated_at
                )
                VALUES('测试干部', '村委工作人员', 'cadre', 0, 'active', ?, ?)
                """,
                (now_text, now_text),
            ).lastrowid
            admin_user_id = conn.execute(
                """
                INSERT INTO admin_users(
                    login_name, display_name, real_name, cadre_authorization_id,
                    role_key, password_hash, status, mfa_enabled, created_at, updated_at
                )
                VALUES('test-cadre', '测试干部', '测试干部', ?, 'cadre',
                       'unused-in-session-test', 'active', 1, ?, ?)
                """,
                (authorization_id, now_text, now_text),
            ).lastrowid
            conn.execute(
                """
                INSERT INTO admin_sessions(
                    admin_user_id, token_hash, created_at, expires_at, last_seen_at
                )
                VALUES(?, ?, ?, ?, ?)
                """,
                (
                    admin_user_id,
                    admin_auth._hash_token("admin-token"),
                    now_text,
                    expires_at,
                    now_text,
                ),
            )
            conn.commit()
        finally:
            conn.close()

    def _decryption_patches(self) -> ExitStack:
        stack = ExitStack()
        id_cards = {
            "id-a": "00000019900101123X",
            "id-b": "000004670401001",
        }
        phones = {
            "phone-a": "10000138888",
            "phone-b": "10100139999",
        }
        stack.enter_context(
            patch.object(services, "decrypt_id_card_number", side_effect=lambda value: id_cards[value])
        )
        stack.enter_context(
            patch.object(admin_services, "_safe_decrypt_id_card", side_effect=lambda value: id_cards[value])
        )
        stack.enter_context(
            patch.object(admin_services, "_safe_decrypt_phone", side_effect=lambda value: phones[value])
        )
        return stack

    def test_villager_auth_me_is_bound_to_current_session(self) -> None:
        with self._decryption_patches():
            user = services.get_user_by_token("villager-a-token")
        self.assertEqual(user["id"], 1)
        self.assertEqual(user["real_name"], "王五")
        self.assertEqual(user["id_card_number"], "00000019900101123X")

    def test_unverified_villager_cannot_submit_proposal(self) -> None:
        with self.assertRaisesRegex(ValueError, "verified_villager_required"):
            services.create_proposal(
                "unverified-villager-token",
                {"content": "未实名账号不应能够提交建议"},
            )

    def test_villager_cannot_read_another_villager_request_by_id(self) -> None:
        own = services.get_villager_request_detail("villager-a-token", 1)
        self.assertEqual(own["title"], "A的事项")
        with self.assertRaisesRegex(ValueError, "request_not_found"):
            services.get_villager_request_detail("villager-a-token", 2)

    def test_villager_request_list_only_contains_own_records(self) -> None:
        rows = services.list_villager_requests("villager-a-token")
        self.assertEqual([row["title"] for row in rows], ["A的事项"])

    def test_villager_gov_records_only_contain_own_applications(self) -> None:
        rows = services.list_gov_applications("villager-a-token")
        self.assertEqual([row["user_id"] for row in rows], [1])
        with self.assertRaisesRegex(ValueError, "gov_application_not_found"):
            services.get_gov_application_detail("villager-a-token", 2)

    def test_unauthenticated_user_cannot_read_villager_profile(self) -> None:
        with self.assertRaisesRegex(ValueError, "invalid_token"):
            services.get_user_by_token("")

    def test_unauthenticated_admin_endpoint_returns_401(self) -> None:
        with self.assertRaises(HTTPException) as raised:
            admin_routes.api_admin_list_villager_identities(None)
        self.assertEqual(raised.exception.status_code, 401)

    def test_private_ledger_materials_require_cadre_session(self) -> None:
        with self.assertRaises(HTTPException) as raised:
            main.api_list_ledgers(None, None, None)
        self.assertEqual(raised.exception.status_code, 401)

    def test_villager_token_cannot_call_admin_identity_service(self) -> None:
        with self.assertRaisesRegex(ValueError, "invalid_admin_token"):
            admin_services.list_admin_villager_identities("villager-a-token")

    def test_authenticated_cadre_list_and_detail_payload_only_have_masked_ids(self) -> None:
        with self._decryption_patches():
            rows = admin_services.list_admin_villager_identities("admin-token")
        by_name = {row["real_name"]: row for row in rows}
        self.assertNotIn("id_card_number", by_name["王五"])
        self.assertNotIn("id_card_number", by_name["李四"])
        self.assertEqual(by_name["王五"]["id_card_masked"], "000000************")
        self.assertEqual(by_name["李四"]["id_card_masked"], "000004*********")
        self.assertNotIn("birth_date", by_name["王五"])
        self.assertNotIn("birth_date", by_name["李四"])

    def test_authenticated_cadre_business_detail_masks_submitted_id_card(self) -> None:
        row = admin_services.get_admin_gov_application_detail("admin-token", 1)
        self.assertEqual(row["extra_fields"]["id_card"], "000000************")
        self.assertEqual(row["applicant_phone"], "100****8888")

    def test_authenticated_cadre_export_contains_masked_identity_only(self) -> None:
        with self._decryption_patches():
            content = admin_services.export_admin_villager_identities_xlsx("admin-token")
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook["村民脱敏名册"]
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            [
                "姓名",
                "性别",
                "身份证号（脱敏）",
                "手机号（脱敏）",
                "村组",
                "年龄",
                "实名状态",
                "名单状态",
                "最后更新时间",
            ],
        )
        exported = {sheet.cell(row=index, column=1).value: sheet.cell(row=index, column=3) for index in range(2, 4)}
        self.assertEqual(exported["王五"].value, "000000************")
        self.assertEqual(exported["李四"].value, "000004*********")
        self.assertEqual(exported["王五"].data_type, "s")
        self.assertEqual(exported["李四"].data_type, "s")
        self.assertEqual(exported["王五"].number_format, "@")
        self.assertEqual(exported["李四"].number_format, "@")
        self.assertEqual({sheet.cell(row=index, column=7).value for index in range(2, 4)}, {"已完成"})
        self.assertEqual({sheet.cell(row=index, column=8).value for index in range(2, 4)}, {"已匹配"})
        self.assertTrue(all(isinstance(sheet.cell(row=index, column=9).value, datetime) for index in range(2, 4)))
        self.assertEqual(sheet.cell(row=2, column=9).number_format, "yyyy-mm-dd hh:mm")
        self.assertEqual(sheet.column_dimensions["C"].width, 24)
        self.assertEqual(sheet.column_dimensions["I"].width, 20)

    def test_public_response_models_do_not_expose_identity_fields(self) -> None:
        for model in (schemas.NoticeOut, schemas.LocalProductOut, schemas.VillageOpenProfileOut):
            self.assertNotIn("id_card_number", model.model_fields)
            self.assertNotIn("id_card_encrypted", model.model_fields)

    def test_privacy_policy_matches_authenticated_cadre_workflow(self) -> None:
        policy = services.get_privacy_policy()["content"]
        self.assertIn("干部 Web 管理后台仅显示身份证号码前 6 位", policy)
        self.assertIn("个人详情和导出文件不展示出生日期", policy)
        self.assertIn("普通后台接口及导出文件不得返回完整号码", policy)


if __name__ == "__main__":
    unittest.main()
