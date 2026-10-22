from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
import base64
import csv
import hashlib
import hmac
import io
import json
import logging
import os
from pathlib import Path
import re
import secrets
import sqlite3
import subprocess
import time
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote, unquote, urlencode, urlparse
from urllib.request import Request as UrlRequest, urlopen

from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerifyMismatchError

from .cloudbase_auth import CloudBaseIdentity, parse_identity_token
from .db import DBIntegrityError, get_connection
from .runtime_config import (
    get_app_env,
    get_database_backend,
    get_phone_aes_key_hex,
    get_public_api_base_url,
    is_cloudbase_auth_enabled,
    load_runtime_env,
)
from .storage import build_download_url, build_file_sha256, resolve_local_file_path, store_upload_bytes
from .secret_protection import keyed_lookup_hash, protect_text

load_runtime_env()

PRIVACY_POLICY_VERSION = os.getenv("PRIVACY_POLICY_VERSION", "2026-07-11-v1")
PRIVACY_POLICY_UPDATED_AT = "2026-07-11T00:00:00+08:00"
PRIVACY_POLICY_TITLE = "得胜服务通隐私政策"
PRIVACY_POLICY_CONTENT = """
个人信息处理者：平江县上塔市镇得胜村村民委员会
更新日期：2026年7月11日
生效日期：2026年7月11日

一、适用范围
本政策适用于“得胜服务通”微信小程序村民端，以及为完成村务服务而与其连接的后台系统。村干部仅通过 Web 管理后台处理其职责范围内的业务信息。本政策不适用于您主动跳转后由其他主体独立运营的第三方小程序或服务；相关处理规则请以第三方页面公示内容为准。

二、我们如何收集和使用个人信息
您可以在未登录状态下浏览首页通知、公开信息和本地好物。登录时，我们通过微信平台获取用于识别账号的 openid；微信在符合条件时可能同时提供 unionid。您主动提供微信昵称或头像时，系统也会保存相应信息。绑定手机号时，我们通过微信官方手机号能力取得手机号码，用于村委会联系、报名确认和办理通知。
我们仅在提供登录、实名认证、事项办理、报名、村务联系、进度反馈、通知提醒、安全防护和必要留档时处理个人信息。我们不出售个人信息，不将个人信息用于商业广告、用户画像或商业营销。

三、实名认证信息
实名认证会收集姓名和身份证号码，用于匹配村级名单、确认服务对象、关联本人办理记录，并形成实名状态、村组等核验结果。系统可能根据身份证号码校验并形成性别、出生日期和年龄等身份信息。手机号仅用于联系和办理通知，不作为村级名单匹配的唯一条件。
完整身份证号码属于敏感个人信息。村民登录后，仅可查看和处理本人的个人信息及办理数据。干部 Web 管理后台仅显示身份证号码前 6 位，其余位数全部以星号隐藏；个人详情和导出文件不展示出生日期，手机号按中间 4 位隐藏。普通后台接口及导出文件不得返回完整号码。未登录用户、公开页面、其他村民及普通后台管理员均不得取得完整身份号码。拒绝实名认证不影响浏览公开内容，但可能无法使用需要确认身份的办理、材料提交、报名或反馈功能。

四、事项办理及上传材料
开具证明会根据证明类型处理申请人姓名、联系电话、证明类型、补充说明和必要的身份字段；涉及关系证明等事项时，还可能处理相关人员的姓名、性别、身份证号码、关系类型或从业说明。系统保存申请时间、办理状态、干部回复、结果材料和办理记录。
低保申请会接收您主动提交的申请书、病历诊断书、出入院记录、费用清单、户口簿、银行卡及可能随户口簿提交的残疾证明等材料；特困认定会接收申请书、户口簿、身份证和银行卡等材料。材料中可能包含家庭关系、民政救助、健康、残疾、医疗费用和金融账户图片等敏感个人信息，请仅提交办理所必需的内容。
得胜服务通只负责申请材料收集、情况登记、联系核实和进度反馈，不进行系统自动审批，不以线上状态或结果代替法定审核、部门审批或线下办理程序。

五、随手拍、就业培训、志愿服务和意见建议
随手拍及一般诉求会处理您填写的标题、现场情况或诉求内容、联系电话、主动上传的图片，以及与实名账号关联的姓名、村组、微信标识、提交时间和办理结果。
就业培训和志愿服务报名会处理姓名、手机号、微信标识、所报名项目和报名时间，用于核验报名、防止重复提交、联系参与人和统计报名人数。
意见建议会处理建议内容、提交时间、微信标识，以及实名用户的姓名，用于村委会内部核实和联系。对外公开相关内容时，不公开提交人的真实身份、手机号、微信标识或其他个人信息。

六、相机、相册及其他权限说明
只有在您主动拍照、选择图片、上传办理材料或保存结果图片时，才会调用相机、相册选择或保存到相册能力。拒绝相关权限不会影响不需要图片的浏览、查询和其他功能。
拨打村委会服务电话只调用系统拨号能力，不读取通讯录、通话记录或短信内容。打开部分便民服务时，小程序可能将服务名称或小程序短链接写入剪贴板，便于您在微信中搜索或打开；系统不读取剪贴板内容。
当前代码未申请或调用位置、通讯录、麦克风、短信读取等权限，也未接入地图定位能力。后续如新增权限，我们会在实际调用前说明目的并依法征得授权。

七、微信订阅消息说明
只有您主动点击开启通知并在微信授权界面同意后，我们才记录订阅状态，并通过微信服务通知发送村务通知提醒。您可以在微信中取消订阅；取消后不再接收相应服务通知，但仍可在“办理记录”中查看事项进度和结果，也可主动进入小程序查看通知。

八、运行日志和系统安全信息
为保障账号、接口和后台安全，系统会记录必要的登录时间、操作类型、业务对象、处理状态和操作时间。接口防滥用功能会按请求 IP 地址生成限流记录；Web 管理后台还会记录管理员登录 IP、浏览器或客户端标识以及审计操作。当前小程序没有主动读取设备型号、通讯录等设备信息的业务代码。
这些信息仅用于身份会话、安全防护、故障排查、责任追溯和防止恶意提交，不用于广告投放或商业分析。

九、敏感个人信息处理说明
身份证号码、身份证或户口簿图片、银行卡图片、医疗诊断和住院材料、费用清单、残疾证明、低保或特困申请信息，以及可能反映家庭关系和经济状况的材料，均可能属于敏感个人信息。我们仅在您主动申请相应服务且确有必要时处理，并通过访问权限、页面脱敏、传输保护、私有存储和操作记录限制使用范围。
如材料包含其他家庭成员或相关人员的信息，请确认您有权提交，并避免上传与当前事项无关的内容。

十、个人信息保存地点和保存期限
个人信息保存在中华人民共和国境内。当前生产系统部署于腾讯云环境，使用云托管、MySQL 数据库和上海地域的腾讯云对象存储保存业务数据及上传文件。
我们在完成村务服务、办理留档、安全审计及法律法规要求所必需的期限内保存信息。事项办理记录、民政救助材料、村务档案或依法需要留存的信息，可能无法在申请删除后立即清除；超过必要期限且无法律或档案留存要求时，我们将删除、匿名化或停止继续使用。接口限流记录仅作短期安全防护使用。

十一、委托处理、对外提供和信息公开
村委会授权的工作人员仅在履行职责所需范围内访问和处理信息。我们委托云基础设施服务商提供系统托管、数据库和对象存储，并通过微信平台完成登录、手机号绑定、订阅消息和小程序能力调用。受托方仅在提供相应服务所需范围内处理信息。
除取得您的单独同意、履行法定义务、响应有权机关依法提出的要求，或保护人身财产安全等法律允许的情形外，我们不向其他单位或个人提供您的个人信息。村务公开、通知公开和建议内容公开不得包含申请人、提交人或相关人员的身份证号码、手机号、微信标识、家庭材料等个人信息。

十二、个人信息安全保护措施
我们采取与信息敏感程度相适应的保护措施，包括验证村民本人身份和干部后台登录状态、控制材料访问范围、在公开页面及非本人村民页面不显示完整身份信息、使用加密传输、将办理材料置于受控存储、保留必要操作记录并限制无关人员访问。发生可能影响您权益的安全事件时，我们将依法采取处置措施并按要求告知或报告。

十三、用户查询、更正、删除、撤回授权等权利
您可以通过得胜服务通首页公布的村委会服务电话，申请查询或更正本人信息、删除依法可以删除的信息、解除实名认证关联，或撤回相关授权。为保护账号安全，村委会可能先核验申请人身份。
您可以在微信系统设置中管理相机、相册和订阅消息等授权。撤回某项授权后，我们将不再通过该权限处理新的信息，但不影响撤回前基于授权已经完成的处理。法律法规、村务档案或事项办理要求必须保留的信息，在必要期限内可能无法立即删除；我们会向您说明原因。

十四、未成年人个人信息保护
本服务面向得胜村服务对象，不以未成年人为商业目标。未成年人使用涉及实名认证、材料上传或事项办理的功能时，应由监护人阅读本政策并给予同意或指导。我们仅为具体村务服务处理必要信息；监护人可以通过首页公布的村委会服务电话申请查询、更正或依法删除未成年人信息。

十五、隐私政策更新
当功能、信息处理目的、权限调用、保存方式或第三方服务发生实质变化时，我们会更新本政策，并通过小程序页面提示或其他适当方式告知。需要重新取得同意的，我们将在相关功能使用前再次征得同意。

十六、联系方式
个人信息处理者为平江县上塔市镇得胜村村民委员会。如对本政策、个人信息处理或权利申请有疑问，请拨打得胜服务通首页公布的村委会服务电话。村委会将在核实身份和申请内容后依法处理。

十七、第三方个人信息处理清单
微信平台能力；服务提供方：腾讯；使用目的：微信登录、手机号绑定、订阅消息、小程序跳转及相机相册等微信能力调用；处理的信息种类：openid、可能提供的 unionid、手机号、订阅授权状态、消息发送结果，以及您主动选择的图片或授权结果；使用场景：登录、绑定联系方式、开启通知、拍照或选图、打开第三方小程序。
腾讯云基础设施；服务提供方：腾讯云；使用目的：后端云托管、数据库存储、上传文件存储和受控下载；处理的信息种类：本政策所列账号、实名、办理、报名、建议、日志和上传材料；使用场景：小程序接口访问、数据保存、后台办理和材料查看。
湘税社保；服务提供方：该服务在微信页面公示的运营主体；使用目的：提供医保社保服务快捷入口；处理的信息种类：得胜服务通未发现自动向其传输本项目个人信息的代码，您进入后提交的信息由该服务按其规则处理；使用场景：首页主动点击快捷入口。
湖南智慧人社；服务提供方：该服务在微信页面公示的运营主体；使用目的：提供人社服务快捷入口；处理的信息种类：得胜服务通未发现自动向其传输本项目个人信息的代码，您进入后提交的信息由该服务按其规则处理；使用场景：首页主动点击快捷入口。
农村土地承包合同网签；服务提供方：该服务页面公示的运营主体；使用目的：提供土地承包服务搜索入口；处理的信息种类：得胜服务通仅将服务名称写入剪贴板，未发现自动传输本项目个人信息的代码；使用场景：首页主动点击后按提示搜索。
湘易办；服务提供方：该服务在微信页面公示的运营主体；使用目的：提供政务服务快捷入口；处理的信息种类：得胜服务通未发现自动向其传输本项目个人信息的代码，您进入后提交的信息由该服务按其规则处理；使用场景：首页主动点击快捷入口。
当前项目未接入短信发送、地图定位、第三方数据统计 SDK 或第三方错误监控服务。
""".strip()

VALID_CONSENT_TYPES = {"privacy_policy", "sensitive_identity"}
APP_ENV = get_app_env()
PUBLIC_API_BASE_URL = get_public_api_base_url()
UPLOAD_URL_SIGN_SECRET = str(os.getenv("UPLOAD_URL_SIGN_SECRET", "") or "").strip()
WECHAT_MINI_APP_ID = str(os.getenv("WECHAT_MINI_APP_ID", "") or "").strip()
WECHAT_MINI_APP_SECRET = str(os.getenv("WECHAT_MINI_APP_SECRET", "") or "").strip()
ALLOW_MOCK_WECHAT_LOGIN = os.getenv("ALLOW_MOCK_WECHAT_LOGIN", "0").strip() == "1"
MOCK_WECHAT_OPENID = str(os.getenv("MOCK_WECHAT_OPENID", "") or "").strip()
MOCK_WECHAT_UNIONID = str(os.getenv("MOCK_WECHAT_UNIONID", "") or "").strip()
USE_CLOUDBASE_IDENTITY = is_cloudbase_auth_enabled()
USE_WECHAT_OPENAPI_SERVICE = APP_ENV == "production" and USE_CLOUDBASE_IDENTITY
WECHAT_CODE2SESSION_URL = "https://api.weixin.qq.com/sns/jscode2session"
WECHAT_ACCESS_TOKEN_URL = "https://api.weixin.qq.com/cgi-bin/token"
WECHAT_STABLE_ACCESS_TOKEN_URL = "https://api.weixin.qq.com/cgi-bin/stable_token"
WECHAT_GET_USER_PHONE_URL = "https://api.weixin.qq.com/wxa/business/getuserphonenumber"
WECHAT_SUBSCRIBE_MESSAGE_SEND_URL = "https://api.weixin.qq.com/cgi-bin/message/subscribe/send"
CLOUDBASE_ACCESS_TOKEN_FILE = "/.tencentcloudbase/wx/cloudbase_access_token"
WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_ID = str(
    os.getenv("WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_ID", "")
    or ""
).strip()
WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_TITLE = str(
    os.getenv("WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_TITLE", "资讯更新提醒") or "资讯更新提醒"
).strip() or "资讯更新提醒"
WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_FIELDS_JSON = str(
    os.getenv(
        "WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_FIELDS_JSON",
        '{"thing6":"notice_title","thing8":"notice_digest","date4":"published_at","thing9":"notice_type","thing5":"notice_tip"}',
    )
    or '{"thing6":"notice_title","thing8":"notice_digest","date4":"published_at","thing9":"notice_type","thing5":"notice_tip"}'
).strip()
WECHAT_NOTICE_SUBSCRIBE_PAGE = str(
    os.getenv("WECHAT_NOTICE_SUBSCRIBE_PAGE", "pages/notice-detail/notice-detail?id={notice_id}")
    or "pages/notice-detail/notice-detail?id={notice_id}"
).strip()
WECHAT_NOTICE_SUBSCRIBE_LANG = str(os.getenv("WECHAT_NOTICE_SUBSCRIBE_LANG", "zh_CN") or "zh_CN").strip() or "zh_CN"
WECHAT_NOTICE_SUBSCRIBE_MINIPROGRAM_STATE = str(
    os.getenv(
        "WECHAT_NOTICE_SUBSCRIBE_MINIPROGRAM_STATE",
        "formal" if APP_ENV == "production" else "developer",
    )
    or ("formal" if APP_ENV == "production" else "developer")
).strip() or ("formal" if APP_ENV == "production" else "developer")
WECHAT_NOTICE_SUBSCRIBE_PROMPT_TITLE = str(
    os.getenv("WECHAT_NOTICE_SUBSCRIBE_PROMPT_TITLE", "开启村务通知提醒") or "开启村务通知提醒"
).strip() or "开启村务通知提醒"
WECHAT_NOTICE_SUBSCRIBE_PROMPT_CONTENT = str(
    os.getenv(
        "WECHAT_NOTICE_SUBSCRIBE_PROMPT_CONTENT",
        "开启后，村委发布通知、政策通知和重要公告时，微信会通过服务通知提醒您。",
    )
    or "开启后，村委发布通知、政策通知和重要公告时，微信会通过服务通知提醒您。"
).strip()
DEFAULT_CADRE_MANAGER_NAMES = [
    name
    for name in dict.fromkeys(
        [
            os.getenv("DEFAULT_CADRE_MANAGER_NAME", "示例人员1").strip() or "示例人员1",
            "示例人员1",
            "示例人员2",
            "示例人员3",
            "示例人员4",
            "示例人员5",
            "示例人员6",
        ]
    )
    if name
]
DEFAULT_CADRE_OFFICE_TITLES = {
    "示例人员3": "村书记",
    "示例人员4": "村副书记",
    "示例人员2": "村委委员",
    "示例人员1": "村委委员",
    "示例人员5": "镇政府",
    "示例人员6": "镇政府",
}
DEFAULT_VILLAGE_NAME = "得胜村"
GENERIC_GOV_CONTACT_NAME = "得胜村部"
GENERIC_GOV_CONTACT_PHONE = "<CONFIGURE_CONTACT_PHONE>"
LEGACY_GOV_CONTACT_NAMES = {
    "得胜村便民服务窗口",
    "示例人员7",
    "示例人员8",
    "示例人员9",
    "示例人员10",
    "示例人员11",
    "示例人员12",
}
LEGACY_GOV_CONTACT_PHONES = {
    "<CONFIGURE_CONTACT_PHONE>",
    "<CONFIGURE_CONTACT_PHONE>",
    "<CONFIGURE_CONTACT_PHONE>",
    "<CONFIGURE_CONTACT_PHONE>",
    "<CONFIGURE_CONTACT_PHONE>",
    "<CONFIGURE_CONTACT_PHONE>",
    "<CONFIGURE_CONTACT_PHONE>",
}
LEGACY_EMPLOYMENT_JOB_KEYS = {"kitchen_helper", "warehouse_packer"}
ID_CARD_AES_KEY_HEX = os.getenv("ID_CARD_AES_KEY_HEX", "").strip().lower()
ID_CARD_DEV_KEY_SEED = os.getenv("ID_CARD_DEV_KEY_SEED", "").strip()
ID_CARD_ENCRYPTION_VERSION = 1
ID_CARD_ENCRYPTION_ALGO = "aes-256-cbc+hmac-sha256"
PHONE_AES_KEY_HEX = get_phone_aes_key_hex()
PHONE_ENCRYPTION_VERSION = 1
PHONE_ENCRYPTION_ALGO = "aes-256-cbc+hmac-sha256"
_ID_CARD_LOOKUP_PURPOSE = "identity.id_card"
_PHONE_LOOKUP_PURPOSE = "identity.phone"
DATABASE_BACKEND = get_database_backend()

DEFAULT_VILLAGE_METRICS = [
    {"metric_key": "groups", "metric_label": "村组", "metric_value": "9"},
    {"metric_key": "families", "metric_label": "家庭", "metric_value": "538"},
    {"metric_key": "population", "metric_label": "人口", "metric_value": "1744"},
    {"metric_key": "members", "metric_label": "党员", "metric_value": "42"},
]
DEFAULT_VILLAGE_INTRO = (
    "得胜村位于湖南省平江县上塔市镇，下辖叶家组、如家组、何家组、胡墩组、新屋组、"
    "徐家组、鲁家组、下西组、花园组。村级治理以通知公开、议事协同、志愿服务和数字公开为主线，"
    "持续提升基层服务效率。"
)

VALID_GOV_DETAIL_TYPES = {"external", "certificate", "info_apply", "employment_training"}
UPLOAD_SCENE_FROM_URL_PATTERN = re.compile(r"/uploads/([a-z0-9_-]{1,32})/")
UPLOAD_ACCESS_URL_PATH_PATTERN = re.compile(r"/api/v1/uploads/access/([a-z0-9_-]{1,32})/([^/?#]+)")
SCENES_NOTICE = {"notice"}
SCENES_LEDGER = {"ledger"}
SCENES_REQUEST = {"request"}
SCENES_GOV_APPLY = {"gov_apply"}
SCENES_REQUEST_REPLY = {"request_reply"}
SCENES_GOV_REPLY = {"gov_reply"}
SCENES_OPEN = {"open"}
SCENES_GOV_DETAIL = {"gov", "gov_service", "gov_detail"}
GENDER_FIELD_ALIASES = {"gender", "sex", "性别"}
HIDDEN_GOV_SERVICE_KEYS = {
    "xiangyiban",
    "yuebanyuehao",
    "house_build",  # 隐藏-待恢复：村民端建房入口隐藏，后台同步停用新增来源
    "hukou_transfer",  # 隐藏-待恢复：村民端户口迁移入口隐藏，后台同步停用新增来源
}
CADRE_STEPUP_SESSION_HOURS = 12
CADRE_STEPUP_LOCK_MINUTES = 30
CADRE_STEPUP_MAX_FAILURES = 5
CADRE_ROLE_LEVEL_NONE = "none"
CADRE_ROLE_LEVEL_CADRE = "cadre"
CADRE_ROLE_LEVEL_SUPER_ADMIN = "super_admin"
CADRE_OPERATION_PASSWORD_HASHER = PasswordHasher(
    time_cost=3,
    memory_cost=65536,
    parallelism=2,
    hash_len=32,
    salt_len=16,
)
PRIVATE_UPLOAD_SCENES = (
    SCENES_REQUEST
    | SCENES_GOV_APPLY
    | SCENES_REQUEST_REPLY
    | SCENES_GOV_REPLY
    | SCENES_LEDGER
)
PUBLIC_UPLOAD_SCENES = (
    SCENES_NOTICE
    | SCENES_OPEN
    | SCENES_GOV_DETAIL
)
CADRE_UPLOAD_SCENES = (
    SCENES_NOTICE
    | SCENES_REQUEST_REPLY
    | SCENES_GOV_REPLY
    | SCENES_OPEN
    | SCENES_GOV_DETAIL
    | SCENES_LEDGER
)
VILLAGER_UPLOAD_SCENES = SCENES_REQUEST | SCENES_GOV_APPLY
GOV_APPLICATION_SUPPORTING_PHOTO_FIELD = "supporting_photo_urls"
CHINA_TIMEZONE = timezone(timedelta(hours=8))
CERTIFICATE_DAILY_APPLICATION_LIMIT = 5
CERTIFICATE_DUPLICATE_FIELD_KEYS = (
    "certificate_key",
    "certificate_title",
    "name",
    "gender",
    "id_card",
    "related_name",
    "related_gender",
    "related_id_card",
    "relation_type",
    "job_desc",
)
PRIVATE_UPLOAD_URL_TTL_SECONDS = 10 * 60
PUBLIC_UPLOAD_URL_TTL_SECONDS = 12 * 60 * 60
RATE_LIMIT_RETENTION_HOURS = 48
_WECHAT_ACCESS_TOKEN_CACHE: dict[str, float | str] = {
    "access_token": "",
    "expires_at": 0.0,
}
NOTICE_SUBSCRIPTION_STATUS_ENABLED_SET = {"accept", "acceptwithaudio", "acceptwithalert"}
NOTICE_SUBSCRIPTION_STATUS_SET = {
    "unknown",
    "accept",
    "acceptwithaudio",
    "acceptwithalert",
    "reject",
    "ban",
    "filter",
    "delivery_rejected",
}
NOTICE_TEMPLATE_VALUE_KEYS = {
    "notice_title",
    "notice_digest",
    "notice_tip",
    "publisher",
    "published_at",
    "notice_type",
    "village_name",
}
logger = logging.getLogger(__name__)

DEFAULT_GOV_SERVICES: list[dict[str, Any]] = [
    {
        "service_key": "certificate",
        "title": "我要开证明",
        "subtitle": "开具村级证明",
        "category": "government",
        "icon_emoji": "🧾",
        "theme_color": "#e8f3f1",
        "contact_name": GENERIC_GOV_CONTACT_NAME,
        "contact_phone": GENERIC_GOV_CONTACT_PHONE,
        "intro": "村民可在线选择证明类型并提交申请，村委核验后通知领取。",
        "detail_type": "certificate",
        "action_label": "开证明",
        "detail_payload": {
            "certificates": [
                {"key": "relationship", "title": "关系证明"},
                {"key": "single_child_loss", "title": "《独生子女父母光荣证》遗失证明"},
                {"key": "no_children", "title": "无子女证明"},
                {"key": "electric_name", "title": "电表更名证明"},
                {"key": "charger_install", "title": "个人充电桩安装同意书"},
                {"key": "kinship_2nd", "title": "亲属关系证明（第二继承人）"},
                {"key": "kinship_1st", "title": "亲属关系证明（第一继承人）"},
                {"key": "marriage_status", "title": "婚育情况证明"},
                {"key": "employment", "title": "从业证明"},
                {"key": "foundling_case", "title": "捡拾弃婴（儿童）情况证明"},
                {"key": "electric_split", "title": "电表分户证明"},
                {"key": "electric_new", "title": "电表上户证明"},
                {"key": "spouse_relation", "title": "夫妻关系证明"},
                {"key": "family_difficulty", "title": "家庭困难证明"},
                {"key": "household", "title": "住所（经营场所）证明"},
            ],
            "template_title": "证 明",
            "template_hint": "兹有我村村民，信息属实，特此证明。",
        },
    },
]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_role(role: str) -> str:
    return "cadre" if (role or "").strip().lower() == "cadre" else "villager"


def normalize_cadre_role_level(role_level: str) -> str:
    text = str(role_level or "").strip().lower()
    if text == CADRE_ROLE_LEVEL_SUPER_ADMIN:
        return CADRE_ROLE_LEVEL_SUPER_ADMIN
    if text == CADRE_ROLE_LEVEL_CADRE:
        return CADRE_ROLE_LEVEL_CADRE
    return CADRE_ROLE_LEVEL_NONE


def _clean_single_line_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _clip_text(value: Any, limit: int) -> str:
    text = _clean_single_line_text(value)
    if not text:
        return ""
    return text[:limit]


def _format_notice_published_at(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        dt = datetime.fromisoformat(text)
    except ValueError:
        return text
    local_dt = dt.astimezone(CHINA_TIMEZONE)
    return f"{local_dt.year}年{local_dt.month}月{local_dt.day}日 {local_dt:%H:%M}"


def _normalize_notice_subscription_status(status_text: Any) -> str:
    text = _clean_single_line_text(status_text).lower()
    if text in NOTICE_SUBSCRIPTION_STATUS_SET:
        return text
    return "unknown"


def _notice_subscription_is_enabled(status_text: Any) -> bool:
    return _normalize_notice_subscription_status(status_text) in NOTICE_SUBSCRIPTION_STATUS_ENABLED_SET


def _get_notice_subscription_template_fields() -> dict[str, str]:
    raw_text = WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_FIELDS_JSON
    if not raw_text:
        return {}
    try:
        payload = json.loads(raw_text)
    except json.JSONDecodeError:
        logger.warning("notice_subscription_template_fields_invalid_json")
        return {}
    if not isinstance(payload, dict):
        return {}
    mapping: dict[str, str] = {}
    for raw_key, raw_value in payload.items():
        key = _clean_single_line_text(raw_key).replace(".DATA", "")
        value = _clean_single_line_text(raw_value)
        if not key or not value or value not in NOTICE_TEMPLATE_VALUE_KEYS:
            continue
        mapping[key] = value
    return mapping


def _notice_subscription_config() -> dict[str, Any]:
    template_fields = _get_notice_subscription_template_fields()
    available = bool(WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_ID and template_fields)
    return {
        "available": available,
        "template_id": WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_ID if available else "",
        "template_title": WECHAT_NOTICE_SUBSCRIBE_TEMPLATE_TITLE,
        "template_fields": template_fields,
        "page": WECHAT_NOTICE_SUBSCRIBE_PAGE,
        "lang": WECHAT_NOTICE_SUBSCRIBE_LANG,
        "miniprogram_state": WECHAT_NOTICE_SUBSCRIBE_MINIPROGRAM_STATE,
        "prompt_title": WECHAT_NOTICE_SUBSCRIBE_PROMPT_TITLE,
        "prompt_content": WECHAT_NOTICE_SUBSCRIBE_PROMPT_CONTENT,
    }


def _build_notice_template_values(notice: dict[str, Any]) -> dict[str, str]:
    title_text = _clean_single_line_text(notice.get("title") or "")
    content_text = _clean_single_line_text(notice.get("content") or "")
    if "政策" in title_text or "政策" in content_text:
        notice_type = "政策通知"
    elif bool(notice.get("is_home_top")):
        notice_type = "置顶通知"
    else:
        notice_type = "村务通知"
    return {
        "notice_title": _clip_text(title_text or "村务通知", 20),
        "notice_digest": _clip_text(content_text or title_text or "请进入小程序查看详情。", 20),
        "notice_tip": _clip_text(f"{DEFAULT_VILLAGE_NAME}村委会发布，请进入小程序查看详情", 20),
        "publisher": _clip_text(notice.get("publisher") or f"{DEFAULT_VILLAGE_NAME}村委会", 20),
        "published_at": _format_notice_published_at(notice.get("created_at")),
        "notice_type": notice_type,
        "village_name": _clip_text(DEFAULT_VILLAGE_NAME, 20),
    }


def _sanitize_notice_template_field(field_key: str, raw_value: str) -> str:
    key = _clean_single_line_text(field_key).replace(".DATA", "")
    value = _clean_single_line_text(raw_value)
    if not value:
        return ""
    field_type_match = re.match(r"^([a-z_]+)", key)
    field_type = str(field_type_match.group(1) if field_type_match else "").lower()
    if field_type == "thing":
        return value[:20]
    if field_type == "phrase":
        return value[:5]
    if field_type == "name":
        return value[:10]
    if field_type == "character_string":
        return value[:32]
    if field_type == "phone_number":
        return value[:17]
    if field_type in {"date", "time"}:
        return value[:32]
    return value[:32]


def _render_notice_template_data(notice: dict[str, Any], template_fields: dict[str, str]) -> dict[str, dict[str, str]]:
    logical_values = _build_notice_template_values(notice)
    payload: dict[str, dict[str, str]] = {}
    for field_key, logical_key in template_fields.items():
        value = _sanitize_notice_template_field(field_key, logical_values.get(logical_key, ""))
        if not value:
            continue
        payload[field_key] = {"value": value}
    return payload


def _render_notice_subscription_page(notice_id: int) -> str:
    template = _clean_single_line_text(WECHAT_NOTICE_SUBSCRIBE_PAGE)
    if not template:
        return ""
    try:
        rendered = template.format(notice_id=int(notice_id or 0))
    except (IndexError, KeyError, ValueError):
        rendered = template
    return rendered[:256]


def _upsert_notice_subscription_conn(
    conn: sqlite3.Connection,
    *,
    user_id: int,
    openid: str,
    template_id: str,
    subscription_status: str,
    is_enabled: bool,
    scene: str,
    result_json: dict[str, Any],
    last_error_code: str = "",
    last_error_message: str = "",
) -> sqlite3.Row:
    now = now_iso()
    subscribed_at = now if is_enabled else ""
    unsubscribed_at = now if not is_enabled and subscription_status in {"reject", "ban", "filter", "delivery_rejected"} else ""
    existing = conn.execute(
        "SELECT id, last_subscribed_at, last_unsubscribed_at FROM notice_subscriptions WHERE user_id = ? LIMIT 1",
        (int(user_id or 0),),
    ).fetchone()
    if existing is None:
        conn.execute(
            """
            INSERT INTO notice_subscriptions(
                user_id, openid, template_id, subscription_status, is_enabled,
                last_request_scene, last_request_result_json, last_error_code, last_error_message,
                last_requested_at, last_subscribed_at, last_unsubscribed_at, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(user_id or 0),
                openid,
                template_id,
                subscription_status,
                1 if is_enabled else 0,
                scene,
                _json_text(result_json),
                last_error_code[:40],
                last_error_message[:255],
                now,
                subscribed_at,
                unsubscribed_at,
                now,
                now,
            ),
        )
    else:
        conn.execute(
            """
            UPDATE notice_subscriptions
            SET
                openid = ?,
                template_id = ?,
                subscription_status = ?,
                is_enabled = ?,
                last_request_scene = ?,
                last_request_result_json = ?,
                last_error_code = ?,
                last_error_message = ?,
                last_requested_at = ?,
                last_subscribed_at = CASE
                    WHEN ? = 1 THEN ?
                    ELSE last_subscribed_at
                END,
                last_unsubscribed_at = CASE
                    WHEN ? = 1 THEN last_unsubscribed_at
                    WHEN ? = '' THEN last_unsubscribed_at
                    ELSE ?
                END,
                updated_at = ?
            WHERE user_id = ?
            """,
            (
                openid,
                template_id,
                subscription_status,
                1 if is_enabled else 0,
                scene,
                _json_text(result_json),
                last_error_code[:40],
                last_error_message[:255],
                now,
                1 if is_enabled else 0,
                subscribed_at or str(existing["last_subscribed_at"] or ""),
                1 if is_enabled else 0,
                unsubscribed_at,
                unsubscribed_at or str(existing["last_unsubscribed_at"] or ""),
                now,
                int(user_id or 0),
            ),
        )
    row = conn.execute(
        "SELECT * FROM notice_subscriptions WHERE user_id = ? LIMIT 1",
        (int(user_id or 0),),
    ).fetchone()
    if row is None:
        raise ValueError("request_failed")
    return row


def _cadre_can_export_full_id(role_level: str, can_export_full_id: Any) -> bool:
    return False


def mask_phone(phone: str) -> str:
    if len(phone) >= 11 and phone.isdigit():
        return f"{phone[:3]}****{phone[-4:]}"
    return phone


def normalize_phone_number(phone: str) -> str:
    return re.sub(r"\D", "", str(phone or "")).strip()[:11]


def _phone_last4(phone: str) -> str:
    text = normalize_phone_number(phone)
    return text[-4:] if len(text) >= 4 else text


def _phone_hash(phone: str) -> str:
    return keyed_lookup_hash(
        normalize_phone_number(phone),
        purpose=_PHONE_LOOKUP_PURPOSE,
    )


def _get_phone_key_hex() -> str:
    key_hex = PHONE_AES_KEY_HEX or _get_id_card_key_hex()
    if re.fullmatch(r"[0-9a-f]{64}", key_hex or "") is None:
        raise ValueError("invalid_phone_aes_key")
    return key_hex


def mask_id_card(id_card_number: str) -> str:
    text = _normalize_id_card(id_card_number)
    if len(text) <= 6:
        return "*" * len(text)
    return f"{text[:6]}{'*' * (len(text) - 6)}"


def normalize_gender(gender: str) -> str:
    text = str(gender or "").strip()
    if not text:
        return "未知"
    mapping = {
        "男": "男",
        "male": "男",
        "m": "男",
        "1": "男",
        "女": "女",
        "female": "女",
        "f": "女",
        "2": "女",
        "未知": "未知",
        "其他": "未知",
        "unknown": "未知",
        "other": "未知",
        "0": "未知",
    }
    lowered = text.lower()
    if lowered in mapping:
        return mapping[lowered]
    if text in mapping:
        return mapping[text]
    return "未知"


def _is_gender_field_name(field_name: str) -> bool:
    text = str(field_name or "").strip().lower()
    if not text:
        return False
    compact = re.sub(r"[\s_-]+", "", text)
    return compact in GENDER_FIELD_ALIASES


def _is_phone_field_name(field_name: str) -> bool:
    text = str(field_name or "").strip().lower()
    if not text:
        return False
    compact = re.sub(r"[\s_-]+", "", text)
    return compact in {
        "phone",
        "mobile",
        "contactphone",
        "contactmobile",
        "applicantphone",
        "reporterphone",
        "手机号",
        "联系电话",
        "联系方式",
    }


def _is_id_card_field_name(field_name: str) -> bool:
    text = str(field_name or "").strip().lower()
    if not text:
        return False
    compact = re.sub(r"[\s_-]+", "", text)
    return compact in {
        "idcard",
        "idcardnumber",
        "idno",
        "identitynumber",
        "spouseidcard",
        "relatedidcard",
        "身份证号",
        "身份证号码",
    }


def _mask_sensitive_value(field_name: str, value: Any) -> Any:
    if value is None:
        return value
    text = str(value or "").strip()
    if not text:
        return value
    if _is_phone_field_name(field_name):
        return mask_phone(normalize_phone_number(text) or text)
    if _is_id_card_field_name(field_name):
        return mask_id_card(text)
    return value


def normalize_photo_url(url: str) -> str:
    text = str(url or "").strip()
    if not text:
        return ""
    if not (text.startswith("http://") or text.startswith("https://")):
        return ""
    parsed = urlparse(text)
    if not parsed.scheme or not parsed.netloc:
        return ""
    origin = f"{parsed.scheme}://{parsed.netloc}"
    upload_ref = _extract_upload_ref_from_path(parsed.path or "")
    if upload_ref is not None:
        return _canonical_upload_url(
            scene=upload_ref["scene"],
            filename=upload_ref["filename"],
            origin=origin,
        )
    return text[:500]


def _extract_upload_ref_from_path(path: str) -> dict[str, str] | None:
    text = str(path or "").strip()
    if not text:
        return None
    match = UPLOAD_ACCESS_URL_PATH_PATTERN.search(text)
    if match is None:
        match = UPLOAD_SCENE_FROM_URL_PATTERN.search(text)
        if match is None:
            return None
        path_parts = [part for part in text.split("/") if part]
        if len(path_parts) < 3:
            return None
        raw_filename = path_parts[-1]
        raw_scene = match.group(1)
    else:
        raw_scene = match.group(1)
        raw_filename = match.group(2)
    scene = str(raw_scene or "").strip().lower()
    filename = str(unquote(str(raw_filename or "")) or "").strip()
    if not scene or not filename:
        return None
    if not re.fullmatch(r"[A-Za-z0-9_.-]{1,160}", filename):
        return None
    return {
        "scene": scene,
        "filename": filename,
    }


def _extract_upload_ref_from_url(url: str) -> tuple[str, dict[str, str]] | None:
    text = str(url or "").strip()
    if not text:
        return None
    parsed = urlparse(text)
    if not parsed.scheme or not parsed.netloc:
        return None
    upload_ref = _extract_upload_ref_from_path(parsed.path or "")
    if upload_ref is None:
        return None
    origin = f"{parsed.scheme}://{parsed.netloc}"
    return origin, upload_ref


def _resolve_public_url_origin(origin: str = "") -> str:
    text = str(origin or "").strip().rstrip("/")
    if PUBLIC_API_BASE_URL:
        return PUBLIC_API_BASE_URL
    if text:
        return text
    return "http://127.0.0.1:8000"


def _canonical_upload_url(*, scene: str, filename: str, origin: str = "") -> str:
    safe_scene = quote(str(scene or "").strip().lower())
    safe_filename = quote(str(filename or "").strip())
    return f"{_resolve_public_url_origin(origin)}/api/v1/uploads/access/{safe_scene}/{safe_filename}"


def _get_upload_url_sign_secret() -> str:
    if UPLOAD_URL_SIGN_SECRET:
        return UPLOAD_URL_SIGN_SECRET
    if APP_ENV == "production" or not ID_CARD_DEV_KEY_SEED:
        raise RuntimeError("UPLOAD_URL_SIGN_SECRET 未配置。")
    return hashlib.sha256(f"{ID_CARD_DEV_KEY_SEED}:upload-url-sign".encode("utf-8")).hexdigest()


def _get_upload_access_ttl_seconds(scene: str) -> int:
    return PRIVATE_UPLOAD_URL_TTL_SECONDS if scene in PRIVATE_UPLOAD_SCENES else PUBLIC_UPLOAD_URL_TTL_SECONDS


def _build_upload_url_signature(scene: str, filename: str, expires_at: int) -> str:
    payload = f"{scene}:{filename}:{int(expires_at)}"
    return hmac.new(
        _get_upload_url_sign_secret().encode("utf-8"),
        payload.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()


def build_upload_access_url(url: str, *, expires_in_seconds: int | None = None) -> str:
    upload_ref = _extract_upload_ref_from_url(url)
    if upload_ref is None:
        return normalize_photo_url(url)
    origin, ref = upload_ref
    max_ttl_seconds = _get_upload_access_ttl_seconds(ref["scene"])
    ttl_seconds = int(expires_in_seconds or max_ttl_seconds)
    ttl_seconds = max(60, min(ttl_seconds, max_ttl_seconds))
    expires_at = int(datetime.now(timezone.utc).timestamp()) + ttl_seconds
    signature = _build_upload_url_signature(ref["scene"], ref["filename"], expires_at)
    base_url = _canonical_upload_url(
        scene=ref["scene"],
        filename=ref["filename"],
        origin=origin,
    )
    return f"{base_url}?exp={expires_at}&sig={signature}"


def verify_upload_access_signature(scene: str, filename: str, exp: int | str, sig: str) -> None:
    scene_text = str(scene or "").strip().lower()
    filename_text = str(filename or "").strip()
    signature = str(sig or "").strip().lower()
    if not scene_text or not filename_text:
        raise ValueError("invalid_upload_signature")
    if not re.fullmatch(r"[A-Za-z0-9_.-]{1,160}", filename_text):
        raise ValueError("invalid_upload_signature")
    try:
        expires_at = int(str(exp or "").strip())
    except (TypeError, ValueError) as exc:
        raise ValueError("invalid_upload_signature") from exc
    now_ts = int(datetime.now(timezone.utc).timestamp())
    if expires_at < now_ts:
        raise ValueError("upload_url_expired")
    max_expires_at = now_ts + _get_upload_access_ttl_seconds(scene_text)
    if expires_at > max_expires_at:
        raise ValueError("invalid_upload_signature")
    expected = _build_upload_url_signature(scene_text, filename_text, expires_at)
    if not hmac.compare_digest(expected, signature):
        raise ValueError("invalid_upload_signature")


def _stored_file_upsert_sql() -> str:
    if DATABASE_BACKEND == "mysql":
        return """
        INSERT INTO stored_files(
            scene, filename, storage_provider, bucket_name, object_key, content_type,
            size_bytes, sha256_hex, uploaded_by_user_id, is_private, source_path, status, created_at, updated_at
        )
        VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?)
        ON DUPLICATE KEY UPDATE
            storage_provider = VALUES(storage_provider),
            bucket_name = VALUES(bucket_name),
            object_key = VALUES(object_key),
            content_type = VALUES(content_type),
            size_bytes = VALUES(size_bytes),
            sha256_hex = VALUES(sha256_hex),
            uploaded_by_user_id = VALUES(uploaded_by_user_id),
            is_private = VALUES(is_private),
            source_path = VALUES(source_path),
            status = 'active',
            updated_at = VALUES(updated_at)
        """
    return """
    INSERT INTO stored_files(
        scene, filename, storage_provider, bucket_name, object_key, content_type,
        size_bytes, sha256_hex, uploaded_by_user_id, is_private, source_path, status, created_at, updated_at
    )
    VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?)
    ON CONFLICT(scene, filename)
    DO UPDATE SET
        storage_provider = excluded.storage_provider,
        bucket_name = excluded.bucket_name,
        object_key = excluded.object_key,
        content_type = excluded.content_type,
        size_bytes = excluded.size_bytes,
        sha256_hex = excluded.sha256_hex,
        uploaded_by_user_id = excluded.uploaded_by_user_id,
        is_private = excluded.is_private,
        source_path = excluded.source_path,
        status = 'active',
        updated_at = excluded.updated_at
    """


def store_uploaded_file_record(
    *,
    scene: str,
    filename: str,
    content_type: str,
    file_bytes: bytes,
    uploaded_by_token: str | None = None,
) -> dict[str, Any]:
    scene_text = str(scene or "").strip().lower()
    filename_text = str(filename or "").strip()
    conn = get_connection()
    try:
        uploaded_by_user_id = None
        if uploaded_by_token:
            user = _get_user_by_token_conn_optional(conn, uploaded_by_token)
            if user is not None:
                uploaded_by_user_id = int(user["id"] or 0) or None
        storage_payload = store_upload_bytes(
            scene=scene_text,
            filename=filename_text,
            content_type=content_type,
            file_bytes=file_bytes,
        )
        now = now_iso()
        is_private = 1 if scene_text in PRIVATE_UPLOAD_SCENES else 0
        conn.execute(
            _stored_file_upsert_sql(),
            (
                scene_text,
                filename_text,
                storage_payload["storage_provider"],
                storage_payload["bucket_name"],
                storage_payload["object_key"],
                content_type,
                len(file_bytes),
                build_file_sha256(file_bytes),
                uploaded_by_user_id,
                is_private,
                storage_payload["source_path"],
                now,
                now,
            ),
        )
        row = conn.execute(
            "SELECT * FROM stored_files WHERE scene = ? AND filename = ? LIMIT 1",
            (scene_text, filename_text),
        ).fetchone()
        conn.commit()
        return dict(row) if row is not None else {}
    finally:
        conn.close()


def get_stored_file_record(scene: str, filename: str) -> dict[str, Any] | None:
    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT *
            FROM stored_files
            WHERE scene = ? AND filename = ? AND status = 'active'
            LIMIT 1
            """,
            (str(scene or "").strip().lower(), str(filename or "").strip()),
        ).fetchone()
        return dict(row) if row is not None else None
    finally:
        conn.close()


def _validate_private_upload_ownership_conn(
    conn: sqlite3.Connection,
    urls: list[str],
    user_id: int,
) -> None:
    for url in urls:
        upload_ref = _extract_upload_ref_from_url(url)
        if upload_ref is None:
            continue
        _, ref = upload_ref
        scene_text = str(ref["scene"] or "").strip().lower()
        if scene_text not in PRIVATE_UPLOAD_SCENES:
            continue
        row = conn.execute(
            """
            SELECT uploaded_by_user_id, is_private
            FROM stored_files
            WHERE scene = ? AND filename = ? AND status = 'active'
            LIMIT 1
            """,
            (scene_text, str(ref["filename"] or "").strip()),
        ).fetchone()
        if (
            row is None
            or int(row["is_private"] or 0) != 1
            or int(row["uploaded_by_user_id"] or 0) != int(user_id or 0)
        ):
            raise ValueError("upload_access_denied")


def _present_photo_url(url: str) -> str:
    text = normalize_photo_url(url)
    if not text:
        return ""
    upload_ref = _extract_upload_ref_from_url(text)
    if upload_ref is None:
        return text
    _, ref = upload_ref
    record = get_stored_file_record(ref["scene"], ref["filename"])
    if record is not None:
        try:
            signed_url = build_download_url(
                record=record,
                expires_in_seconds=_get_upload_access_ttl_seconds(ref["scene"]),
            )
        except FileNotFoundError:
            signed_url = ""
        if signed_url:
            return signed_url
    return build_upload_access_url(text)


def _extract_upload_scene_from_url(url: str) -> str:
    ref = _extract_upload_ref_from_url(url)
    if ref is None:
        return ""
    _, upload_ref = ref
    return str(upload_ref["scene"] or "").strip().lower()


def _is_photo_url_scene_allowed(url: str, allowed_scenes: set[str], *, allow_external: bool) -> bool:
    if not allowed_scenes:
        return True
    scene = _extract_upload_scene_from_url(url)
    if scene:
        return scene in allowed_scenes
    return allow_external


def is_valid_phone(phone: str) -> bool:
    return re.fullmatch(r"1\d{10}", normalize_phone_number(phone)) is not None


def _parse_iso_datetime(value: str) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def consume_rate_limit(scope: str, bucket_key: str, limit: int, window_seconds: int) -> None:
    scope_text = str(scope or "").strip().lower()
    bucket_text = str(bucket_key or "").strip()
    if not scope_text or not bucket_text:
        raise ValueError("rate_limit_exceeded")
    safe_limit = max(1, int(limit))
    safe_window_seconds = max(1, int(window_seconds))
    now_dt = datetime.now(timezone.utc)
    now = now_dt.isoformat()
    window_epoch = int(now_dt.timestamp() // safe_window_seconds) * safe_window_seconds
    window_start = datetime.fromtimestamp(window_epoch, timezone.utc).isoformat()
    retention_before = (now_dt - timedelta(hours=RATE_LIMIT_RETENTION_HOURS)).isoformat()

    conn = get_connection()
    try:
        conn.execute(
            "DELETE FROM request_rate_limits WHERE updated_at < ?",
            (retention_before,),
        )
        row = conn.execute(
            """
            SELECT id, request_count
            FROM request_rate_limits
            WHERE scope = ? AND bucket_key = ? AND window_start = ?
            LIMIT 1
            """,
            (scope_text, bucket_text, window_start),
        ).fetchone()
        if row is None:
            conn.execute(
                """
                INSERT INTO request_rate_limits(
                    scope, bucket_key, window_start, request_count, created_at, updated_at
                )
                VALUES(?, ?, ?, 1, ?, ?)
                """,
                (scope_text, bucket_text, window_start, now, now),
            )
            conn.commit()
            return
        request_count = int(row["request_count"] or 0)
        if request_count >= safe_limit:
            conn.commit()
            raise ValueError("rate_limit_exceeded")
        conn.execute(
            """
            UPDATE request_rate_limits
            SET request_count = ?, updated_at = ?
            WHERE id = ?
            """,
            (request_count + 1, now, row["id"]),
        )
        conn.commit()
    finally:
        conn.close()

def is_valid_id_card(id_card_number: str) -> bool:
    text = _normalize_id_card(id_card_number)
    if re.fullmatch(r"\d{15}", text) is None and re.fullmatch(r"\d{17}[0-9X]", text) is None:
        return False
    birth_date_text = _extract_birth_date_from_id_card(text)
    if not birth_date_text:
        return False
    try:
        birth_date = datetime.strptime(birth_date_text, "%Y-%m-%d").date()
    except ValueError:
        return False
    today = datetime.now(CHINA_TIMEZONE).date()
    if birth_date > today or birth_date < date(today.year - 130, 1, 1):
        return False
    if len(text) == 15:
        return True
    weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
    checksum_map = "10X98765432"
    checksum_index = sum(int(text[index]) * weights[index] for index in range(17)) % 11
    return text[-1] == checksum_map[checksum_index]


def _extract_birth_date_from_id_card(id_card_number: str) -> str:
    text = _normalize_id_card(id_card_number)
    if re.fullmatch(r"\d{17}[0-9X]", text):
        raw = text[6:14]
    elif re.fullmatch(r"\d{15}", text):
        raw = f"19{text[6:12]}"
    else:
        return ""
    try:
        birth_dt = datetime.strptime(raw, "%Y%m%d").date()
    except ValueError:
        return ""
    return birth_dt.isoformat()


def _calculate_age_from_birth_date(birth_date: str) -> int | None:
    text = str(birth_date or "").strip()
    if not text:
        return None
    try:
        birth_dt = datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None
    today = datetime.now(timezone(timedelta(hours=8))).date()
    age = today.year - birth_dt.year
    if (today.month, today.day) < (birth_dt.month, birth_dt.day):
        age -= 1
    return age if 0 <= age <= 130 else None


def _gender_from_id_card(id_card_number: str) -> str:
    text = _normalize_id_card(id_card_number)
    if not is_valid_id_card(text):
        return ""
    if len(text) == 18:
        gender_digit = int(text[16])
    elif len(text) == 15:
        gender_digit = int(text[14])
    else:
        return ""
    return "女" if gender_digit % 2 == 0 else "男"


def _resolve_gender(stored_gender: str, id_card_number: str) -> str:
    normalized = normalize_gender(stored_gender)
    if normalized in {"男", "女"}:
        return normalized
    return _gender_from_id_card(id_card_number) or normalized


def calc_auth_status(phone_bound: bool, realname_status: str, face_status: str) -> str:
    _ = face_status
    if phone_bound and realname_status == "verified":
        return "verified"
    if phone_bound or realname_status != "unsubmitted":
        return "partial"
    return "unverified"


def _normalize_id_card(id_card_number: str) -> str:
    return (id_card_number or "").strip().replace(" ", "").upper()


def _id_hash(id_card_number: str) -> str:
    return keyed_lookup_hash(
        _normalize_id_card(id_card_number),
        purpose=_ID_CARD_LOOKUP_PURPOSE,
    )


def _get_id_card_key_hex() -> str:
    key_hex = ID_CARD_AES_KEY_HEX
    if not key_hex:
        if APP_ENV == "production":
            raise ValueError("id_card_aes_key_not_configured")
        if not ID_CARD_DEV_KEY_SEED:
            raise ValueError("id_card_aes_key_not_configured")
        key_hex = hashlib.sha256(ID_CARD_DEV_KEY_SEED.encode("utf-8")).hexdigest()
    if re.fullmatch(r"[0-9a-f]{64}", key_hex or "") is None:
        raise ValueError("invalid_id_card_aes_key")
    return key_hex


def _run_openssl_aes256_cbc(
    *,
    decrypt: bool,
    key_hex: str,
    iv_bytes: bytes,
    payload: bytes,
) -> bytes:
    command = [
        "openssl",
        "enc",
        "-aes-256-cbc",
        "-d" if decrypt else "-e",
        "-K",
        key_hex,
        "-iv",
        iv_bytes.hex(),
        "-nosalt",
    ]
    try:
        process = subprocess.run(
            command,
            input=payload,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
        )
    except OSError as exc:
        raise ValueError("id_card_decrypt_failed" if decrypt else "id_card_encrypt_failed") from exc
    if process.returncode != 0:
        raise ValueError("id_card_decrypt_failed" if decrypt else "id_card_encrypt_failed")
    return process.stdout


def encrypt_id_card_number(id_card_number: str) -> str:
    plain_text = _normalize_id_card(id_card_number)
    if not plain_text:
        raise ValueError("invalid_id_card")
    key_hex = _get_id_card_key_hex()
    key_bytes = bytes.fromhex(key_hex)
    iv_bytes = secrets.token_bytes(16)
    cipher_bytes = _run_openssl_aes256_cbc(
        decrypt=False,
        key_hex=key_hex,
        iv_bytes=iv_bytes,
        payload=plain_text.encode("utf-8"),
    )
    mac_hex = hmac.new(key_bytes, iv_bytes + cipher_bytes, hashlib.sha256).hexdigest()
    return _json_text(
        {
            "v": ID_CARD_ENCRYPTION_VERSION,
            "alg": ID_CARD_ENCRYPTION_ALGO,
            "iv": base64.b64encode(iv_bytes).decode("ascii"),
            "ct": base64.b64encode(cipher_bytes).decode("ascii"),
            "mac": mac_hex,
        }
    )


def decrypt_id_card_number(encrypted_payload: str) -> str:
    payload_text = str(encrypted_payload or "").strip()
    if not payload_text:
        return ""
    try:
        payload_obj = json.loads(payload_text)
    except json.JSONDecodeError as exc:
        raise ValueError("id_card_decrypt_failed") from exc
    if not isinstance(payload_obj, dict):
        raise ValueError("id_card_decrypt_failed")
    if int(payload_obj.get("v", 0) or 0) != ID_CARD_ENCRYPTION_VERSION:
        raise ValueError("id_card_decrypt_failed")
    if str(payload_obj.get("alg") or "") != ID_CARD_ENCRYPTION_ALGO:
        raise ValueError("id_card_decrypt_failed")

    try:
        iv_bytes = base64.b64decode(str(payload_obj.get("iv") or ""), validate=True)
        cipher_bytes = base64.b64decode(str(payload_obj.get("ct") or ""), validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("id_card_decrypt_failed") from exc
    if len(iv_bytes) != 16 or not cipher_bytes:
        raise ValueError("id_card_decrypt_failed")

    mac_hex = str(payload_obj.get("mac") or "").lower()
    if re.fullmatch(r"[0-9a-f]{64}", mac_hex) is None:
        raise ValueError("id_card_decrypt_failed")

    key_hex = _get_id_card_key_hex()
    key_bytes = bytes.fromhex(key_hex)
    expected_mac = hmac.new(key_bytes, iv_bytes + cipher_bytes, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(expected_mac, mac_hex):
        raise ValueError("id_card_decrypt_failed")

    plain_bytes = _run_openssl_aes256_cbc(
        decrypt=True,
        key_hex=key_hex,
        iv_bytes=iv_bytes,
        payload=cipher_bytes,
    )
    try:
        plain_text = plain_bytes.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("id_card_decrypt_failed") from exc
    return _normalize_id_card(plain_text)


def encrypt_phone_number(phone: str) -> str:
    plain_text = normalize_phone_number(phone)
    if not plain_text:
        return ""
    key_hex = _get_phone_key_hex()
    key_bytes = bytes.fromhex(key_hex)
    iv_bytes = secrets.token_bytes(16)
    cipher_bytes = _run_openssl_aes256_cbc(
        decrypt=False,
        key_hex=key_hex,
        iv_bytes=iv_bytes,
        payload=plain_text.encode("utf-8"),
    )
    mac_hex = hmac.new(key_bytes, iv_bytes + cipher_bytes, hashlib.sha256).hexdigest()
    return _json_text(
        {
            "v": PHONE_ENCRYPTION_VERSION,
            "alg": PHONE_ENCRYPTION_ALGO,
            "iv": base64.b64encode(iv_bytes).decode("ascii"),
            "ct": base64.b64encode(cipher_bytes).decode("ascii"),
            "mac": mac_hex,
        }
    )


def decrypt_phone_number(encrypted_payload: str) -> str:
    payload_text = str(encrypted_payload or "").strip()
    if not payload_text:
        return ""
    try:
        payload_obj = json.loads(payload_text)
    except json.JSONDecodeError as exc:
        raise ValueError("phone_decrypt_failed") from exc
    if not isinstance(payload_obj, dict):
        raise ValueError("phone_decrypt_failed")
    if int(payload_obj.get("v", 0) or 0) != PHONE_ENCRYPTION_VERSION:
        raise ValueError("phone_decrypt_failed")
    if str(payload_obj.get("alg") or "") != PHONE_ENCRYPTION_ALGO:
        raise ValueError("phone_decrypt_failed")
    try:
        iv_bytes = base64.b64decode(str(payload_obj.get("iv") or ""), validate=True)
        cipher_bytes = base64.b64decode(str(payload_obj.get("ct") or ""), validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("phone_decrypt_failed") from exc
    if len(iv_bytes) != 16 or not cipher_bytes:
        raise ValueError("phone_decrypt_failed")
    mac_hex = str(payload_obj.get("mac") or "").lower()
    if re.fullmatch(r"[0-9a-f]{64}", mac_hex) is None:
        raise ValueError("phone_decrypt_failed")
    key_hex = _get_phone_key_hex()
    key_bytes = bytes.fromhex(key_hex)
    expected_mac = hmac.new(key_bytes, iv_bytes + cipher_bytes, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(expected_mac, mac_hex):
        raise ValueError("phone_decrypt_failed")
    plain_bytes = _run_openssl_aes256_cbc(
        decrypt=True,
        key_hex=key_hex,
        iv_bytes=iv_bytes,
        payload=cipher_bytes,
    )
    try:
        return normalize_phone_number(plain_bytes.decode("utf-8"))
    except UnicodeDecodeError as exc:
        raise ValueError("phone_decrypt_failed") from exc


def _phone_fields(phone: str) -> dict[str, str]:
    plain_text = normalize_phone_number(phone)
    if not plain_text:
        return {
            "phone": "",
            "phone_ciphertext": "",
            "phone_hash": "",
            "phone_masked": "",
            "phone_last4": "",
        }
    masked = mask_phone(plain_text)
    return {
        "phone": masked,
        "phone_ciphertext": encrypt_phone_number(plain_text),
        "phone_hash": _phone_hash(plain_text),
        "phone_masked": masked,
        "phone_last4": _phone_last4(plain_text),
    }


def _protected_business_phone_fields(phone: str, *, purpose: str) -> dict[str, str]:
    plain_text = normalize_phone_number(phone)
    if not plain_text:
        return {
            "ciphertext": "",
            "lookup_hash": "",
            "masked": "",
            "legacy_token": "",
        }
    lookup_hash = keyed_lookup_hash(plain_text, purpose=purpose)
    return {
        "ciphertext": encrypt_phone_number(plain_text),
        "lookup_hash": lookup_hash,
        "masked": mask_phone(plain_text),
        "legacy_token": f"protected:{lookup_hash[:16]}",
    }


def _protect_gov_application_extra_fields(
    extra_fields: dict[str, Any],
) -> tuple[dict[str, Any], str, str]:
    public_fields = _normalize_json_object(extra_fields)
    sensitive_fields: dict[str, Any] = {}
    applicant_id_card_hash = ""
    for key, value in list(public_fields.items()):
        if not (_is_phone_field_name(key) or _is_id_card_field_name(key)):
            continue
        text = str(value or "").strip()
        if not text:
            continue
        if _is_phone_field_name(key) and not is_valid_phone(text):
            continue
        if _is_id_card_field_name(key) and not is_valid_id_card(text):
            continue
        sensitive_fields[key] = value
        public_fields[key] = _mask_sensitive_value(key, value)
        if not applicant_id_card_hash and _is_id_card_field_name(key):
            normalized_id = _normalize_id_card(text)
            if normalized_id:
                applicant_id_card_hash = _id_hash(normalized_id)
    encrypted_fields = (
        protect_text(_json_text(sensitive_fields), purpose="gov_application_extra")
        if sensitive_fields
        else ""
    )
    return public_fields, encrypted_fields, applicant_id_card_hash

def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _villager_registry_upsert_sql() -> str:
    if DATABASE_BACKEND == "mysql":
        return """
        INSERT INTO villager_registry(
            user_id, village_name, group_name, real_name, gender, phone,
            phone_ciphertext, phone_hash, phone_masked, phone_last4,
            id_card_masked, id_card_last4, id_card_hash, id_card_encrypted, face_photo_url,
            realname_submitted_at, face_verified_at, auth_status, created_at, updated_at
        )
        VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON DUPLICATE KEY UPDATE
            village_name = VALUES(village_name),
            group_name = VALUES(group_name),
            real_name = VALUES(real_name),
            gender = VALUES(gender),
            phone = VALUES(phone),
            phone_ciphertext = VALUES(phone_ciphertext),
            phone_hash = VALUES(phone_hash),
            phone_masked = VALUES(phone_masked),
            phone_last4 = VALUES(phone_last4),
            id_card_masked = VALUES(id_card_masked),
            id_card_last4 = VALUES(id_card_last4),
            id_card_hash = VALUES(id_card_hash),
            id_card_encrypted = VALUES(id_card_encrypted),
            face_photo_url = VALUES(face_photo_url),
            realname_submitted_at = VALUES(realname_submitted_at),
            face_verified_at = VALUES(face_verified_at),
            auth_status = VALUES(auth_status),
            updated_at = VALUES(updated_at)
        """
    return """
    INSERT INTO villager_registry(
        user_id, village_name, group_name, real_name, gender, phone,
        phone_ciphertext, phone_hash, phone_masked, phone_last4,
        id_card_masked, id_card_last4, id_card_hash, id_card_encrypted, face_photo_url,
        realname_submitted_at, face_verified_at, auth_status, created_at, updated_at
    )
    VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ON CONFLICT(user_id)
    DO UPDATE SET
        village_name = excluded.village_name,
        group_name = excluded.group_name,
        real_name = excluded.real_name,
        gender = excluded.gender,
        phone = excluded.phone,
        phone_ciphertext = excluded.phone_ciphertext,
        phone_hash = excluded.phone_hash,
        phone_masked = excluded.phone_masked,
        phone_last4 = excluded.phone_last4,
        id_card_masked = excluded.id_card_masked,
        id_card_last4 = excluded.id_card_last4,
        id_card_hash = excluded.id_card_hash,
        id_card_encrypted = excluded.id_card_encrypted,
        face_photo_url = excluded.face_photo_url,
        realname_submitted_at = excluded.realname_submitted_at,
        face_verified_at = excluded.face_verified_at,
        auth_status = excluded.auth_status,
        updated_at = excluded.updated_at
    """


def _cadre_stepup_secret_upsert_sql() -> str:
    if DATABASE_BACKEND == "mysql":
        return """
        INSERT INTO cadre_stepup_secrets(
            user_id, password_hash, password_algorithm, failed_count, locked_until, created_at, updated_at
        )
        VALUES(?, ?, 'argon2id', 0, '', ?, ?)
        ON DUPLICATE KEY UPDATE
            password_hash = VALUES(password_hash),
            password_algorithm = VALUES(password_algorithm),
            failed_count = 0,
            locked_until = '',
            updated_at = VALUES(updated_at)
        """
    return """
    INSERT INTO cadre_stepup_secrets(
        user_id, password_hash, password_algorithm, failed_count, locked_until, created_at, updated_at
    )
    VALUES(?, ?, 'argon2id', 0, '', ?, ?)
    ON CONFLICT(user_id)
    DO UPDATE SET
        password_hash = excluded.password_hash,
        password_algorithm = excluded.password_algorithm,
        failed_count = 0,
        locked_until = '',
        updated_at = excluded.updated_at
    """


def _privacy_consent_upsert_sql() -> str:
    if DATABASE_BACKEND == "mysql":
        return """
        INSERT INTO privacy_consents(
            user_id, policy_version, consent_type, status, created_at, updated_at, withdrawn_at
        )
        VALUES(?, ?, ?, 'agreed', ?, ?, '')
        ON DUPLICATE KEY UPDATE
            status = 'agreed',
            updated_at = VALUES(updated_at),
            withdrawn_at = ''
        """
    return """
    INSERT INTO privacy_consents(
        user_id, policy_version, consent_type, status, created_at, updated_at, withdrawn_at
    )
    VALUES(?, ?, ?, 'agreed', ?, ?, '')
    ON CONFLICT(user_id, policy_version, consent_type)
    DO UPDATE SET status = 'agreed', updated_at = excluded.updated_at, withdrawn_at = ''
    """


def _normalize_photo_urls(
    value: Any,
    *,
    allowed_scenes: set[str] | None = None,
    allow_external: bool = True,
    strict_scene: bool = False,
    invalid_scene_error: str = "invalid_photo_scene",
) -> list[str]:
    if not isinstance(value, list):
        return []
    result: list[str] = []
    for item in value:
        text = normalize_photo_url(item)
        if not text:
            continue
        if allowed_scenes is not None and not _is_photo_url_scene_allowed(
            text,
            allowed_scenes,
            allow_external=allow_external,
        ):
            if strict_scene:
                raise ValueError(invalid_scene_error)
            continue
        if text not in result:
            result.append(text)
        if len(result) >= 9:
            break
    return result


def _parse_photo_urls_field(
    raw: Any,
    *,
    allowed_scenes: set[str] | None = None,
    allow_external: bool = True,
) -> list[str]:
    if isinstance(raw, list):
        return [
            _present_photo_url(item)
            for item in _normalize_photo_urls(
                raw,
                allowed_scenes=allowed_scenes,
                allow_external=allow_external,
            )
        ]
    text = str(raw or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return []
    return [
        _present_photo_url(item)
        for item in _normalize_photo_urls(
            parsed,
            allowed_scenes=allowed_scenes,
            allow_external=allow_external,
        )
    ]


def _normalize_json_object(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {}
    result: dict[str, Any] = {}
    for key, raw in value.items():
        text_key = str(key or "").strip()
        if not text_key:
            continue
        if isinstance(raw, (str, int, float, bool)) or raw is None:
            result[text_key[:60]] = raw
            continue
        if isinstance(raw, list):
            result[text_key[:60]] = raw[:50]
            continue
        if isinstance(raw, dict):
            result[text_key[:60]] = raw
            continue
        result[text_key[:60]] = str(raw)
    return result


def _parse_json_object(raw: Any) -> dict[str, Any]:
    if isinstance(raw, dict):
        return _normalize_json_object(raw)
    text = str(raw or "").strip()
    if not text:
        return {}
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        return {}
    return _normalize_json_object(value)


def _mask_sensitive_json_object(raw: Any) -> dict[str, Any]:
    payload = _parse_json_object(raw)
    return {
        key: _mask_sensitive_value(key, value)
        for key, value in payload.items()
    }


def _sanitize_gov_detail_payload(raw: Any, *, strict_scene: bool = False) -> dict[str, Any]:
    payload = _parse_json_object(raw)
    if "image_urls" not in payload:
        return payload
    payload["image_urls"] = [
        _present_photo_url(item)
        for item in _normalize_photo_urls(
            payload.get("image_urls"),
            allowed_scenes=SCENES_GOV_DETAIL,
            allow_external=True,
            strict_scene=strict_scene,
        )
    ]
    return payload


def _resolve_user_phone_number(row: sqlite3.Row | dict[str, Any]) -> str:
    data = dict(row)
    encrypted_phone = str(data.get("phone_ciphertext") or "").strip()
    if encrypted_phone:
        try:
            phone_number = decrypt_phone_number(encrypted_phone)
            if is_valid_phone(phone_number):
                return phone_number
        except ValueError:
            pass
    phone_number = normalize_phone_number(str(data.get("phone") or ""))
    return phone_number if is_valid_phone(phone_number) else ""


def _resolve_user_phone_masked(row: sqlite3.Row | dict[str, Any]) -> str:
    data = dict(row)
    stored_masked = str(data.get("phone_masked") or "").strip()
    if stored_masked:
        return stored_masked
    phone_number = _resolve_user_phone_number(data)
    if phone_number:
        return mask_phone(phone_number)
    return mask_phone(str(data.get("phone") or ""))


def auth_user_to_dict(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    data = dict(row)
    phone_number = _resolve_user_phone_number(data)
    phone_masked = _resolve_user_phone_masked(data)
    phone_verified_at = str(data.get("phone_verified_at") or "").strip()
    phone_bound = bool(data.get("phone_bound", 0)) and bool(phone_verified_at)
    cadre_eligible = bool(data.get("cadre_eligible", False))
    auth_status = calc_auth_status(
        phone_bound,
        str(data.get("realname_status") or "unsubmitted"),
        str(data.get("face_status") or "unsubmitted"),
    )
    id_card_number = ""
    encrypted_id_card = str(data.get("id_card_encrypted") or "").strip()
    if encrypted_id_card:
        try:
            id_card_number = decrypt_id_card_number(encrypted_id_card)
        except ValueError:
            id_card_number = ""
    birth_date = _extract_birth_date_from_id_card(id_card_number)
    age = _calculate_age_from_birth_date(birth_date)
    gender = _resolve_gender(data.get("gender", ""), id_card_number)
    return {
        "id": data["id"],
        "nickname": data.get("nickname", ""),
        "avatar_url": data.get("avatar_url", ""),
        "phone": phone_number if phone_bound and phone_number else phone_masked,
        "phone_masked": phone_masked,
        "phone_bound": phone_bound,
        "phone_verified_at": phone_verified_at,
        "real_name": data.get("real_name", ""),
        "gender": gender,
        "id_card_masked": mask_id_card(id_card_number or str(data.get("id_card_masked") or "")),
        "id_card_number": id_card_number,
        "birth_date": birth_date,
        "age": age,
        "role": "cadre" if cadre_eligible else "villager",
        "cadre_eligible": cadre_eligible,
        "cadre_role_level": normalize_cadre_role_level(data.get("cadre_role_level", "")),
        "cadre_password_set": bool(data.get("cadre_password_set", False)),
        "realname_status": data.get("realname_status", "unsubmitted"),
        "auth_status": auth_status,
        "roster_match_status": data.get("roster_match_status", "unmatched"),
        "village_group": data.get("village_group", ""),
        "created_at": data["created_at"],
        "updated_at": data["updated_at"],
    }


def _is_phone_verified_user(row: sqlite3.Row | dict[str, Any]) -> bool:
    data = dict(row)
    return bool(data.get("phone_bound", 0)) and bool(str(data.get("phone_verified_at") or "").strip())


def _is_verified_identity_user(row: sqlite3.Row | dict[str, Any]) -> bool:
    data = dict(row)
    return (
        _is_phone_verified_user(data)
        and str(data.get("auth_status") or "") == "verified"
        and str(data.get("realname_status") or "") == "verified"
        and bool(str(data.get("id_card_hash") or "").strip())
    )


def _get_cadre_authorization_for_user_conn(
    conn: sqlite3.Connection,
    user: sqlite3.Row | dict[str, Any],
) -> sqlite3.Row | None:
    data = dict(user)
    user_id = int(data.get("id") or 0)
    openid = str(data.get("openid") or "").strip()
    unionid = str(data.get("unionid") or "").strip()
    real_name = str(data.get("real_name") or "").strip()
    candidates: list[sqlite3.Row] = []
    if user_id:
        row = conn.execute(
            """
            SELECT *
            FROM cadre_authorizations
            WHERE user_id = ? AND status = 'active'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if row is not None:
            candidates.append(row)
    if unionid:
        row = conn.execute(
            """
            SELECT *
            FROM cadre_authorizations
            WHERE unionid = ? AND status = 'active'
            LIMIT 1
            """,
            (unionid,),
        ).fetchone()
        if row is not None:
            candidates.append(row)
    if openid:
        row = conn.execute(
            """
            SELECT *
            FROM cadre_authorizations
            WHERE openid = ? AND status = 'active'
            LIMIT 1
            """,
            (openid,),
        ).fetchone()
        if row is not None:
            candidates.append(row)
    if real_name:
        row = conn.execute(
            """
            SELECT *
            FROM cadre_authorizations
            WHERE real_name = ? AND status = 'active'
            LIMIT 1
            """,
            (real_name,),
        ).fetchone()
        if row is not None:
            candidates.append(row)
    for row in candidates:
        auth_user_id = int(row["user_id"] or 0)
        auth_openid = str(row["openid"] or "").strip()
        auth_unionid = str(row["unionid"] or "").strip()
        if auth_user_id and auth_user_id != user_id:
            continue
        if auth_openid and openid and auth_openid != openid:
            continue
        if auth_unionid and unionid and auth_unionid != unionid:
            continue
        if real_name and str(row["real_name"] or "").strip() != real_name:
            continue
        return row
    return None


def _sync_user_cadre_state_conn(conn: sqlite3.Connection, user_id: int) -> sqlite3.Row | None:
    user = conn.execute("SELECT * FROM users WHERE id = ? LIMIT 1", (user_id,)).fetchone()
    if user is None:
        return None
    auth_row = None
    if _is_verified_identity_user(user):
        auth_row = _get_cadre_authorization_for_user_conn(conn, user)
        if auth_row is not None:
            now = now_iso()
            openid = str(user["openid"] or "").strip()
            unionid = str(user["unionid"] or "").strip()
            conn.execute(
                """
                UPDATE cadre_authorizations
                SET
                    user_id = ?,
                    openid = CASE WHEN openid IS NULL OR openid = '' THEN ? ELSE openid END,
                    unionid = CASE
                        WHEN ? = '' THEN unionid
                        WHEN unionid IS NULL OR unionid = '' THEN ?
                        ELSE unionid
                    END,
                    last_bound_at = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (user_id, openid, unionid, unionid, now, now, auth_row["id"]),
            )
            auth_row = conn.execute(
                "SELECT * FROM cadre_authorizations WHERE id = ? LIMIT 1",
                (auth_row["id"],),
            ).fetchone()
    if normalize_role(str(user["role"] or "")) != "villager":
        now = now_iso()
        conn.execute(
            "UPDATE users SET role = 'villager', role_request = 'villager', updated_at = ? WHERE id = ?",
            (now, user_id),
        )
        user = conn.execute("SELECT * FROM users WHERE id = ? LIMIT 1", (user_id,)).fetchone()
    return auth_row


def _cadre_password_is_set_conn(conn: sqlite3.Connection, user_id: int) -> bool:
    row = conn.execute(
        """
        SELECT id
        FROM cadre_stepup_secrets
        WHERE user_id = ?
        LIMIT 1
        """,
        (user_id,),
    ).fetchone()
    return row is not None


def _build_auth_user_conn(conn: sqlite3.Connection, user_id: int) -> dict[str, Any]:
    _sync_user_cadre_state_conn(conn, user_id)
    row = conn.execute("SELECT * FROM users WHERE id = ? LIMIT 1", (user_id,)).fetchone()
    if row is None:
        raise ValueError("invalid_token")
    auth_row = _get_cadre_authorization_for_user_conn(conn, row)
    data = dict(row)
    data["cadre_eligible"] = auth_row is not None
    data["cadre_role_level"] = (
        normalize_cadre_role_level(auth_row["role_level"]) if auth_row is not None else CADRE_ROLE_LEVEL_NONE
    )
    data["cadre_password_set"] = (
        _cadre_password_is_set_conn(conn, user_id) if auth_row is not None else False
    )
    return auth_user_to_dict(data)


def _sync_villager_registry_conn(conn: sqlite3.Connection, user_id: int) -> None:
    row = conn.execute(
        """
        SELECT
            id,
            role,
            real_name,
            gender,
            phone,
            phone_ciphertext,
            phone_hash,
            phone_masked,
            phone_last4,
            id_card_masked,
            id_card_last4,
            id_card_hash,
            id_card_encrypted,
            face_photo_url,
            village_group,
            realname_status,
            face_status,
            auth_status,
            updated_at,
            created_at,
            face_verified_at
        FROM users
        WHERE id = ?
        LIMIT 1
        """,
        (user_id,),
    ).fetchone()
    if row is None:
        return
    if normalize_role(row["role"]) != "villager":
        return
    now = now_iso()
    id_card_number = ""
    encrypted_id_card = str(row["id_card_encrypted"] or "").strip()
    if encrypted_id_card:
        try:
            id_card_number = decrypt_id_card_number(encrypted_id_card)
        except ValueError:
            id_card_number = ""
    gender = _resolve_gender(str(row["gender"] or ""), id_card_number)
    realname_submitted_at = (
        row["updated_at"] if str(row["realname_status"] or "") != "unsubmitted" else ""
    )
    conn.execute(
        _villager_registry_upsert_sql(),
        (
            user_id,
            DEFAULT_VILLAGE_NAME,
            str(row["village_group"] or ""),
            str(row["real_name"] or ""),
            gender,
            str(row["phone"] or ""),
            str(row["phone_ciphertext"] or ""),
            str(row["phone_hash"] or ""),
            str(row["phone_masked"] or ""),
            str(row["phone_last4"] or ""),
            str(row["id_card_masked"] or ""),
            str(row["id_card_last4"] or ""),
            str(row["id_card_hash"] or ""),
            str(row["id_card_encrypted"] or ""),
            str(row["face_photo_url"] or ""),
            str(realname_submitted_at or ""),
            str(row["face_verified_at"] or ""),
            str(row["auth_status"] or "unverified"),
            str(row["created_at"] or now),
            now,
        ),
    )


def backfill_missing_identity_gender() -> int:
    conn = get_connection()
    updated_count = 0
    try:
        rows = conn.execute(
            """
            SELECT id, gender, id_card_encrypted
            FROM users
            WHERE role = 'villager'
              AND realname_status = 'verified'
              AND id_card_encrypted <> ''
            """
        ).fetchall()
        for row in rows:
            stored_gender = normalize_gender(str(row["gender"] or ""))
            if stored_gender in {"男", "女"}:
                continue
            encrypted_id_card = str(row["id_card_encrypted"] or "").strip()
            if not encrypted_id_card:
                continue
            try:
                id_card_number = decrypt_id_card_number(encrypted_id_card)
            except ValueError:
                continue
            gender = _gender_from_id_card(id_card_number)
            if not gender:
                continue
            user_id = int(row["id"] or 0)
            conn.execute("UPDATE users SET gender = ? WHERE id = ?", (gender, user_id))
            conn.execute("UPDATE villager_registry SET gender = ? WHERE user_id = ?", (gender, user_id))
            updated_count += 1
        conn.commit()
        return updated_count
    finally:
        conn.close()


def export_latest_villager_registry_csv(token: str, cadre_token: str | None = None) -> bytes:
    conn = get_connection()
    try:
        cadre_user, _, _ = _require_cadre_stepup_user_conn(conn, token, cadre_token)
        rows = conn.execute(
            """
            SELECT
                village_name,
                group_name,
                real_name,
                gender,
                phone,
                id_card_masked,
                auth_status,
                realname_submitted_at,
                updated_at
            FROM villager_registry
            ORDER BY
                CASE
                    WHEN group_name = '叶家组' THEN 1
                    WHEN group_name = '如家组' THEN 2
                    WHEN group_name = '何家组' THEN 3
                    WHEN group_name = '胡墩组' THEN 4
                    WHEN group_name = '新屋组' THEN 5
                    WHEN group_name = '徐家组' THEN 6
                    WHEN group_name = '鲁家组' THEN 7
                    WHEN group_name = '下西组' THEN 8
                    WHEN group_name = '花园组' THEN 9
                    ELSE 99
                END ASC,
                real_name ASC,
                id ASC
            """
        ).fetchall()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(
            [
                "村名",
                "组别",
                "姓名",
                "性别",
                "联系电话",
                "身份证号(脱敏)",
                "认证状态",
                "实名提交时间",
                "更新时间",
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    str(row["village_name"] or DEFAULT_VILLAGE_NAME),
                    str(row["group_name"] or ""),
                    str(row["real_name"] or ""),
                    str(row["gender"] or ""),
                    mask_phone(normalize_phone_number(str(row["phone"] or "")) or str(row["phone"] or "")),
                    mask_id_card(str(row["id_card_masked"] or "")),
                    str(row["auth_status"] or ""),
                    str(row["realname_submitted_at"] or ""),
                    str(row["updated_at"] or ""),
                ]
            )
        _audit_log_conn(
            conn,
            int(cadre_user["id"] or 0),
            "villager_registry_export",
            "villager_registry",
            "latest",
            {
                "count": len(rows),
                "full_id_exported": False,
            },
        )
        conn.commit()
        return output.getvalue().encode("utf-8-sig")
    finally:
        conn.close()


def hash_cadre_operation_password(password: str) -> str:
    return CADRE_OPERATION_PASSWORD_HASHER.hash(password)


def verify_cadre_operation_password(password: str, password_hash: str) -> bool:
    try:
        return bool(CADRE_OPERATION_PASSWORD_HASHER.verify(password_hash, password))
    except (VerifyMismatchError, InvalidHashError):
        return False


def _audit_log_conn(
    conn: sqlite3.Connection,
    actor_user_id: int | None,
    action: str,
    target_type: str,
    target_id: str,
    detail: dict[str, Any],
) -> None:
    conn.execute(
        """
        INSERT INTO audit_logs(actor_user_id, action, target_type, target_id, detail_json, created_at)
        VALUES(?, ?, ?, ?, ?, ?)
        """,
        (
            actor_user_id,
            action,
            target_type,
            target_id,
            _json_text(detail),
            now_iso(),
        ),
    )


def ensure_default_cadre_authorizations_conn(conn: sqlite3.Connection) -> None:
    now = now_iso()
    for index, manager_name in enumerate(DEFAULT_CADRE_MANAGER_NAMES):
        role_level = CADRE_ROLE_LEVEL_SUPER_ADMIN if index == 0 else CADRE_ROLE_LEVEL_CADRE
        can_export_full_id = 0
        office_title = str(DEFAULT_CADRE_OFFICE_TITLES.get(manager_name) or "").strip()
        row = conn.execute(
            """
            SELECT id
            FROM cadre_authorizations
            WHERE real_name = ?
            LIMIT 1
            """,
            (manager_name,),
        ).fetchone()
        if row is None:
            conn.execute(
                """
                INSERT INTO cadre_authorizations(
                    real_name, office_title, role_level, can_export_full_id, status, created_at, updated_at
                )
                VALUES(?, ?, ?, ?, 'active', ?, ?)
                """,
                (manager_name, office_title, role_level, can_export_full_id, now, now),
            )
            continue
        conn.execute(
            """
            UPDATE cadre_authorizations
            SET
                office_title = ?,
                role_level = ?,
                can_export_full_id = ?,
                status = 'active',
                updated_at = ?
            WHERE id = ?
            """,
            (office_title, role_level, can_export_full_id, now, row["id"]),
        )
    conn.execute(
        "UPDATE cadre_managers SET status = 'inactive', updated_at = ?",
        (now,),
    )


def ensure_default_cadre_authorizations() -> None:
    conn = get_connection()
    try:
        ensure_default_cadre_authorizations_conn(conn)
        conn.commit()
    finally:
        conn.close()


def ensure_default_cadre_managers() -> None:
    ensure_default_cadre_authorizations()


def backfill_phone_security() -> None:
    conn = get_connection()
    try:
        user_rows = conn.execute(
            """
            SELECT
                id,
                phone,
                phone_ciphertext,
                phone_hash,
                phone_masked,
                phone_last4,
                phone_bound,
                phone_verified_at,
                realname_status,
                face_status,
                auth_status
            FROM users
            WHERE phone != '' OR phone_bound != 0 OR auth_status != 'unverified'
            """
        ).fetchall()
        for row in user_rows:
            phone_text = normalize_phone_number(str(row["phone"] or ""))
            if len(phone_text) == 11:
                phone_fields = _phone_fields(phone_text)
                if not (
                    str(row["phone"] or "") == phone_fields["phone"]
                    and str(row["phone_ciphertext"] or "") == phone_fields["phone_ciphertext"]
                    and str(row["phone_hash"] or "") == phone_fields["phone_hash"]
                    and str(row["phone_masked"] or "") == phone_fields["phone_masked"]
                    and str(row["phone_last4"] or "") == phone_fields["phone_last4"]
                ):
                    conn.execute(
                        """
                        UPDATE users
                        SET phone = ?, phone_ciphertext = ?, phone_hash = ?, phone_masked = ?, phone_last4 = ?
                        WHERE id = ?
                        """,
                        (
                            phone_fields["phone"],
                            phone_fields["phone_ciphertext"],
                            phone_fields["phone_hash"],
                            phone_fields["phone_masked"],
                            phone_fields["phone_last4"],
                            row["id"],
                        ),
                    )
            expected_auth_status = calc_auth_status(
                _is_phone_verified_user(row),
                str(row["realname_status"] or "unsubmitted"),
                str(row["face_status"] or "unsubmitted"),
            )
            if str(row["auth_status"] or "") != expected_auth_status:
                conn.execute(
                    "UPDATE users SET auth_status = ? WHERE id = ?",
                    (expected_auth_status, row["id"]),
                )

        registry_rows = conn.execute(
            """
            SELECT id, phone, phone_ciphertext, phone_hash, phone_masked, phone_last4
            FROM villager_registry
            WHERE phone != ''
            """
        ).fetchall()
        for row in registry_rows:
            phone_text = normalize_phone_number(str(row["phone"] or ""))
            if len(phone_text) != 11:
                continue
            phone_fields = _phone_fields(phone_text)
            if (
                str(row["phone"] or "") == phone_fields["phone"]
                and str(row["phone_ciphertext"] or "") == phone_fields["phone_ciphertext"]
                and str(row["phone_hash"] or "") == phone_fields["phone_hash"]
                and str(row["phone_masked"] or "") == phone_fields["phone_masked"]
                and str(row["phone_last4"] or "") == phone_fields["phone_last4"]
            ):
                continue
            conn.execute(
                """
                UPDATE villager_registry
                SET phone = ?, phone_ciphertext = ?, phone_hash = ?, phone_masked = ?, phone_last4 = ?
                WHERE id = ?
                """,
                (
                    phone_fields["phone"],
                    phone_fields["phone_ciphertext"],
                    phone_fields["phone_hash"],
                    phone_fields["phone_masked"],
                    phone_fields["phone_last4"],
                    row["id"],
                ),
            )

        conn.commit()
    finally:
        conn.close()


def _hash_session_token(token: str) -> str:
    return hashlib.sha256(str(token or "").encode("utf-8")).hexdigest()


def _find_user_by_wechat_identity_conn(
    conn: sqlite3.Connection,
    *,
    openid: str,
    unionid: str = "",
) -> sqlite3.Row | dict[str, Any] | None:
    unionid_text = str(unionid or "").strip()
    openid_text = str(openid or "").strip()
    if unionid_text:
        row = conn.execute("SELECT * FROM users WHERE unionid = ? LIMIT 1", (unionid_text,)).fetchone()
        if row is not None:
            return row
    if openid_text:
        return conn.execute("SELECT * FROM users WHERE openid = ? LIMIT 1", (openid_text,)).fetchone()
    return None


def _ensure_user_by_wechat_identity_conn(
    conn: sqlite3.Connection,
    *,
    openid: str,
    unionid: str = "",
    nickname: str = "",
    avatar_url: str = "",
) -> sqlite3.Row | dict[str, Any]:
    openid_text = str(openid or "").strip()
    unionid_text = str(unionid or "").strip()
    if not openid_text:
        raise ValueError("invalid_token")

    now = now_iso()
    row = _find_user_by_wechat_identity_conn(conn, openid=openid_text, unionid=unionid_text)
    if row is None:
        cursor = conn.execute(
            """
            INSERT INTO users(
                openid, unionid, nickname, avatar_url, role, role_request,
                realname_status, face_status, auth_status, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, 'villager', 'villager', 'unsubmitted', 'unsubmitted', 'unverified', ?, ?)
            """,
            (openid_text, unionid_text, nickname, avatar_url, now, now),
        )
        user_id = cursor.lastrowid
    else:
        user_id = int(row["id"] or 0)
        effective_nickname = nickname or str(row["nickname"] or "")
        effective_avatar = avatar_url or str(row["avatar_url"] or "")
        conn.execute(
            """
            UPDATE users
            SET
                openid = ?,
                unionid = CASE WHEN ? = '' THEN unionid ELSE ? END,
                nickname = ?,
                avatar_url = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (openid_text, unionid_text, unionid_text, effective_nickname, effective_avatar, now, user_id),
        )
    _sync_user_cadre_state_conn(conn, user_id)
    _sync_villager_registry_conn(conn, user_id)
    ensured = conn.execute("SELECT * FROM users WHERE id = ? LIMIT 1", (user_id,)).fetchone()
    if ensured is None:
        raise ValueError("invalid_token")
    return ensured


def _get_user_by_token_conn(conn: sqlite3.Connection, token: str) -> sqlite3.Row:
    token_text = str(token or "").strip()
    if not token_text:
        raise ValueError("invalid_token")
    cloudbase_identity = parse_identity_token(token_text)
    if USE_CLOUDBASE_IDENTITY and cloudbase_identity is not None:
        row = _ensure_user_by_wechat_identity_conn(
            conn,
            openid=cloudbase_identity.openid,
            unionid=cloudbase_identity.unionid,
        )
        data = dict(row)
        data["session_id"] = 0
        data["session_token"] = ""
        data["expires_at"] = ""
        return data
    token_hash = _hash_session_token(token_text)
    row = conn.execute(
        """
        SELECT
            s.id AS session_id,
            s.token AS session_token,
            u.id,
            u.openid,
            u.unionid,
            u.nickname,
            u.avatar_url,
            u.phone,
            u.phone_ciphertext,
            u.phone_hash,
            u.phone_masked,
            u.phone_last4,
            u.phone_bound,
            u.phone_verified_at,
            u.real_name,
            u.gender,
            u.id_card_masked,
            u.id_card_last4,
            u.id_card_hash,
            u.id_card_encrypted,
            u.face_photo_url,
            u.face_verified_at,
            u.role,
            u.role_request,
            u.realname_status,
            u.face_status,
            u.auth_status,
            u.roster_match_status,
            u.village_group,
            u.created_at,
            u.updated_at,
            s.expires_at
        FROM auth_sessions s
        JOIN users u ON u.id = s.user_id
        WHERE s.token IN (?, ?)
        ORDER BY CASE WHEN s.token = ? THEN 0 ELSE 1 END ASC
        LIMIT 1
        """,
        (token_hash, token_text, token_hash),
    ).fetchone()
    if row is None:
        raise ValueError("invalid_token")

    if str(row["session_token"] or "") != token_hash:
        conn.execute(
            "UPDATE auth_sessions SET token = ? WHERE id = ?",
            (token_hash, row["session_id"]),
        )

    expires_at = row["expires_at"]
    try:
        expires_dt = datetime.fromisoformat(expires_at)
    except ValueError as exc:
        raise ValueError("invalid_token") from exc
    if expires_dt < datetime.now(timezone.utc):
        conn.execute("DELETE FROM auth_sessions WHERE id = ?", (row["session_id"],))
        conn.commit()
        raise ValueError("token_expired")
    return row


def _require_cadre_user_conn(conn: sqlite3.Connection, token: str) -> sqlite3.Row:
    user, _ = _require_cadre_authorized_user_conn(conn, token)
    return user


def _require_villager_user_conn(conn: sqlite3.Connection, token: str) -> sqlite3.Row:
    user = _get_user_by_token_conn(conn, token)
    return user


def _require_verified_villager_user_conn(conn: sqlite3.Connection, token: str) -> sqlite3.Row:
    user = _require_villager_user_conn(conn, token)
    if not _is_verified_identity_user(user):
        raise ValueError("verified_villager_required")
    return user


def _get_user_by_token_conn_optional(conn: sqlite3.Connection, token: str | None) -> sqlite3.Row | None:
    if not str(token or "").strip():
        return None
    try:
        return _get_user_by_token_conn(conn, str(token or "").strip())
    except ValueError:
        return None


def get_user_by_token(token: str) -> dict[str, Any]:
    conn = get_connection()
    try:
        row = _get_user_by_token_conn(conn, token)
        data = _build_auth_user_conn(conn, int(row["id"] or 0))
        conn.commit()
        return data
    finally:
        conn.close()


def resolve_role_by_token(token: str | None) -> str | None:
    if not token:
        return None
    try:
        user = get_user_by_token(token)
    except ValueError:
        return None
    return "cadre" if user.get("cadre_eligible") else "villager"


def _get_cadre_authorization_context_conn(
    conn: sqlite3.Connection,
    user: sqlite3.Row | dict[str, Any],
) -> dict[str, Any] | None:
    if not _is_verified_identity_user(user):
        return None
    auth_row = _sync_user_cadre_state_conn(conn, int(dict(user).get("id") or 0))
    if auth_row is None:
        return None
    return {
        "authorization_id": int(auth_row["id"] or 0),
        "role_level": normalize_cadre_role_level(auth_row["role_level"]),
        "can_export_full_id": _cadre_can_export_full_id(
            auth_row["role_level"],
            auth_row["can_export_full_id"],
        ),
    }


def _require_cadre_authorized_user_conn(
    conn: sqlite3.Connection,
    token: str,
) -> tuple[sqlite3.Row, dict[str, Any]]:
    user = _get_user_by_token_conn(conn, token)
    access = _get_cadre_authorization_context_conn(conn, user)
    if access is None:
        raise ValueError("cadre_authorization_required")
    return user, access


def _get_cadre_stepup_secret_conn(
    conn: sqlite3.Connection,
    user_id: int,
) -> sqlite3.Row | None:
    return conn.execute(
        """
        SELECT *
        FROM cadre_stepup_secrets
        WHERE user_id = ?
        LIMIT 1
        """,
        (user_id,),
    ).fetchone()


def _create_cadre_stepup_session_conn(conn: sqlite3.Connection, user_id: int) -> tuple[str, str]:
    now = now_iso()
    token = secrets.token_urlsafe(32)
    token_hash = _hash_session_token(token)
    expires_at = (datetime.now(timezone.utc) + timedelta(hours=CADRE_STEPUP_SESSION_HOURS)).isoformat()
    conn.execute(
        """
        UPDATE cadre_stepup_sessions
        SET revoked_at = ?
        WHERE user_id = ? AND revoked_at = ''
        """,
        (now, user_id),
    )
    conn.execute(
        """
        INSERT INTO cadre_stepup_sessions(user_id, token, created_at, expires_at, last_used_at, revoked_at)
        VALUES(?, ?, ?, ?, ?, '')
        """,
        (user_id, token_hash, now, expires_at, now),
    )
    return token, expires_at


def _verify_cadre_stepup_session_conn(
    conn: sqlite3.Connection,
    user_id: int,
    cadre_token: str | None,
) -> sqlite3.Row:
    token_text = str(cadre_token or "").strip()
    if not token_text:
        raise ValueError("cadre_stepup_required")
    token_hash = _hash_session_token(token_text)
    row = conn.execute(
        """
        SELECT id, user_id, token, expires_at, revoked_at
        FROM cadre_stepup_sessions
        WHERE user_id = ? AND token IN (?, ?) AND revoked_at = ''
        ORDER BY CASE WHEN token = ? THEN 0 ELSE 1 END ASC
        LIMIT 1
        """,
        (user_id, token_hash, token_text, token_hash),
    ).fetchone()
    if row is None:
        raise ValueError("cadre_stepup_required")
    expires_dt = _parse_iso_datetime(str(row["expires_at"] or ""))
    if expires_dt is None or expires_dt <= datetime.now(timezone.utc):
        conn.execute(
            "UPDATE cadre_stepup_sessions SET revoked_at = ? WHERE id = ?",
            (now_iso(), row["id"]),
        )
        raise ValueError("cadre_stepup_required")
    if str(row["token"] or "") != token_hash:
        conn.execute(
            "UPDATE cadre_stepup_sessions SET token = ?, last_used_at = ? WHERE id = ?",
            (token_hash, now_iso(), row["id"]),
        )
    else:
        conn.execute(
            "UPDATE cadre_stepup_sessions SET last_used_at = ? WHERE id = ?",
            (now_iso(), row["id"]),
        )
    return row


def _require_cadre_stepup_user_conn(
    conn: sqlite3.Connection,
    token: str,
    cadre_token: str | None,
) -> tuple[sqlite3.Row, dict[str, Any], sqlite3.Row]:
    user, access = _require_cadre_authorized_user_conn(conn, token)
    session_row = _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
    return user, access, session_row


def get_cadre_access_status(token: str, cadre_token: str | None = None) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        access = _get_cadre_authorization_context_conn(conn, user)
        if access is None:
            return {
                "cadre_eligible": False,
                "cadre_role_level": CADRE_ROLE_LEVEL_NONE,
                "password_set": False,
                "session_active": False,
                "session_expires_at": "",
                "locked_until": "",
            }
        secret_row = _get_cadre_stepup_secret_conn(conn, int(user["id"] or 0))
        session_active = False
        session_expires_at = ""
        try:
            session_row = _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
            session_active = True
            session_expires_at = str(session_row["expires_at"] or "")
        except ValueError:
            session_active = False
        locked_until = str(secret_row["locked_until"] or "") if secret_row is not None else ""
        conn.commit()
        return {
            "cadre_eligible": True,
            "cadre_role_level": access["role_level"],
            "password_set": secret_row is not None,
            "session_active": session_active,
            "session_expires_at": session_expires_at,
            "locked_until": locked_until,
        }
    finally:
        conn.close()


def set_cadre_operation_password(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    password = str(payload.get("password") or "").strip()
    if len(password) < 8 or password.isdigit():
        raise ValueError("weak_cadre_operation_password")
    conn = get_connection()
    try:
        user, access = _require_cadre_authorized_user_conn(conn, token)
        now = now_iso()
        conn.execute(
            _cadre_stepup_secret_upsert_sql(),
            (user["id"], hash_cadre_operation_password(password), now, now),
        )
        _audit_log_conn(
            conn,
            int(user["id"] or 0),
            "cadre_operation_password_set",
            "cadre_access",
            str(user["id"]),
            {"role_level": access["role_level"]},
        )
        conn.commit()
        return get_cadre_access_status(token)
    finally:
        conn.close()


def verify_cadre_access_password(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    password = str(payload.get("password") or "").strip()
    if not password:
        raise ValueError("invalid_cadre_operation_password")
    conn = get_connection()
    try:
        user, access = _require_cadre_authorized_user_conn(conn, token)
        secret_row = _get_cadre_stepup_secret_conn(conn, int(user["id"] or 0))
        if secret_row is None:
            raise ValueError("cadre_operation_password_not_set")
        locked_until = str(secret_row["locked_until"] or "")
        locked_until_dt = _parse_iso_datetime(locked_until)
        if locked_until_dt is not None and locked_until_dt > datetime.now(timezone.utc):
            raise ValueError("cadre_access_locked")
        if not verify_cadre_operation_password(password, str(secret_row["password_hash"] or "")):
            now = now_iso()
            failed_count = int(secret_row["failed_count"] or 0) + 1
            if failed_count >= CADRE_STEPUP_MAX_FAILURES:
                next_locked_until = (
                    datetime.now(timezone.utc) + timedelta(minutes=CADRE_STEPUP_LOCK_MINUTES)
                ).isoformat()
                conn.execute(
                    """
                    UPDATE cadre_stepup_secrets
                    SET failed_count = ?, locked_until = ?, updated_at = ?
                    WHERE user_id = ?
                    """,
                    (failed_count, next_locked_until, now, user["id"]),
                )
                _audit_log_conn(
                    conn,
                    int(user["id"] or 0),
                    "cadre_operation_password_locked",
                    "cadre_access",
                    str(user["id"]),
                    {"role_level": access["role_level"], "failed_count": failed_count},
                )
                conn.commit()
                raise ValueError("cadre_access_locked")
            conn.execute(
                """
                UPDATE cadre_stepup_secrets
                SET failed_count = ?, updated_at = ?
                WHERE user_id = ?
                """,
                (failed_count, now, user["id"]),
            )
            _audit_log_conn(
                conn,
                int(user["id"] or 0),
                "cadre_operation_password_failed",
                "cadre_access",
                str(user["id"]),
                {"role_level": access["role_level"], "failed_count": failed_count},
            )
            conn.commit()
            raise ValueError("invalid_cadre_operation_password")
        now = now_iso()
        conn.execute(
            """
            UPDATE cadre_stepup_secrets
            SET failed_count = 0, locked_until = '', updated_at = ?
            WHERE user_id = ?
            """,
            (now, user["id"]),
        )
        step_token, expires_at = _create_cadre_stepup_session_conn(conn, int(user["id"] or 0))
        _audit_log_conn(
            conn,
            int(user["id"] or 0),
            "cadre_operation_password_verified",
            "cadre_access",
            str(user["id"]),
            {"role_level": access["role_level"]},
        )
        conn.commit()
        return {
            "token": step_token,
            "expires_at": expires_at,
            "status": "ok",
        }
    finally:
        conn.close()


def logout_cadre_access(token: str, cadre_token: str | None) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        if str(cadre_token or "").strip():
            token_hash = _hash_session_token(str(cadre_token or "").strip())
            conn.execute(
                """
                UPDATE cadre_stepup_sessions
                SET revoked_at = ?
                WHERE user_id = ? AND token IN (?, ?) AND revoked_at = ''
                """,
                (now_iso(), user["id"], token_hash, str(cadre_token or "").strip()),
            )
        conn.commit()
        return {"status": "ok", "message": "已退出干部工作台"}
    finally:
        conn.close()


def validate_upload_scene_access(token: str, scene: str, cadre_token: str | None = None) -> None:
    scene_text = str(scene or "").strip().lower()
    conn = get_connection()
    try:
        if scene_text in (SCENES_REQUEST | SCENES_GOV_APPLY):
            _require_verified_villager_user_conn(conn, token)
            return
        if scene_text in CADRE_UPLOAD_SCENES:
            _require_cadre_stepup_user_conn(conn, token, cadre_token)
            conn.commit()
            return
        raise ValueError("invalid_upload_scene")
    finally:
        conn.close()


def _require_consent_conn(
    conn: sqlite3.Connection,
    user_id: int,
    consent_type: str,
    error_code: str,
) -> None:
    row = conn.execute(
        """
        SELECT id FROM privacy_consents
        WHERE user_id = ?
          AND policy_version = ?
          AND consent_type = ?
          AND status = 'agreed'
        LIMIT 1
        """,
        (user_id, PRIVACY_POLICY_VERSION, consent_type),
    ).fetchone()
    if row is None:
        raise ValueError(error_code)


def _list_privacy_consents_conn(conn: sqlite3.Connection, user_id: int) -> list[dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT consent_type, status, policy_version, updated_at
        FROM privacy_consents
        WHERE user_id = ? AND policy_version = ?
        ORDER BY consent_type ASC
        """,
        (user_id, PRIVACY_POLICY_VERSION),
    ).fetchall()
    return [dict(row) for row in rows]


def get_privacy_policy() -> dict[str, Any]:
    return {
        "version": PRIVACY_POLICY_VERSION,
        "title": PRIVACY_POLICY_TITLE,
        "content": PRIVACY_POLICY_CONTENT,
        "updated_at": PRIVACY_POLICY_UPDATED_AT,
        "required_consents": ["privacy_policy", "sensitive_identity"],
    }


def list_privacy_consents(token: str) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        return {"items": _list_privacy_consents_conn(conn, user["id"])}
    finally:
        conn.close()


def submit_privacy_consent(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    policy_version = (payload.get("policy_version") or "").strip()
    consent_types = payload.get("consent_types") or []
    if policy_version != PRIVACY_POLICY_VERSION:
        raise ValueError("invalid_policy_version")
    if not isinstance(consent_types, list) or not consent_types:
        raise ValueError("invalid_consent_type")

    normalized_types = []
    for consent_type in consent_types:
        text = (consent_type or "").strip()
        if text not in VALID_CONSENT_TYPES:
            raise ValueError("invalid_consent_type")
        if text not in normalized_types:
            normalized_types.append(text)

    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        now = now_iso()
        for consent_type in normalized_types:
            conn.execute(
                _privacy_consent_upsert_sql(),
                (user["id"], policy_version, consent_type, now, now),
            )

        _audit_log_conn(
            conn,
            user["id"],
            "privacy_consent_submit",
            "privacy",
            str(user["id"]),
            {"consent_types": normalized_types, "policy_version": policy_version},
        )
        items = _list_privacy_consents_conn(conn, user["id"])
        conn.commit()
        return {"items": items}
    finally:
        conn.close()


def withdraw_privacy_consent(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    reason = (payload.get("reason") or "").strip()
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        now = now_iso()
        conn.execute(
            """
            UPDATE privacy_consents
            SET status = 'withdrawn', withdrawn_at = ?, updated_at = ?
            WHERE user_id = ? AND status = 'agreed'
            """,
            (now, now, user["id"]),
        )
        conn.execute("DELETE FROM auth_sessions WHERE user_id = ?", (user["id"],))
        conn.execute(
            """
            UPDATE users
            SET
                phone = '',
                phone_ciphertext = '',
                phone_hash = '',
                phone_masked = '',
                phone_last4 = '',
                phone_bound = 0,
                phone_verified_at = '',
                real_name = '',
                gender = '',
                id_card_masked = '',
                id_card_last4 = '',
                id_card_hash = '',
                id_card_encrypted = '',
                face_photo_url = '',
                face_verified_at = '',
                role = 'villager',
                role_request = 'villager',
                realname_status = 'unsubmitted',
                face_status = 'unsubmitted',
                auth_status = 'unverified',
                roster_match_status = 'unmatched',
                village_group = '',
                updated_at = ?
            WHERE id = ?
            """,
            (now, user["id"]),
        )
        conn.execute("DELETE FROM villager_registry WHERE user_id = ?", (user["id"],))
        _audit_log_conn(
            conn,
            user["id"],
            "privacy_consent_withdraw",
            "privacy",
            str(user["id"]),
            {"reason": reason},
        )
        conn.commit()
        return {"status": "withdrawn", "message": "已撤回授权并清空认证信息"}
    finally:
        conn.close()


def _normalize_header(text: str) -> str:
    return re.sub(r"[\s_（）()：:]", "", (text or "").strip().lower())


_NAME_HEADERS = {
    "姓名",
    "村民姓名",
    "name",
}
_ID_HEADERS = {
    "身份证号",
    "身份证",
    "证件号",
    "idcard",
    "idcardnumber",
    "idnumber",
}
_GROUP_HEADERS = {
    "组别",
    "小组",
    "村民组",
    "网格",
    "group",
    "groupname",
}


def _pick_field(raw_row: dict[str, Any], aliases: set[str]) -> str:
    for key, value in raw_row.items():
        if _normalize_header(str(key)) in aliases:
            text = str(value or "").strip()
            if text:
                return text
    return ""


def _parse_csv_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    last_exc: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            text = file_bytes.decode(encoding)
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise ValueError("empty_roster_file")
            return [dict(row) for row in reader]
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
            continue
    raise ValueError("invalid_roster_file") from last_exc


def _parse_xlsx_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("xlsx_not_supported") from exc

    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("empty_roster_file")
    headers = [str(item or "").strip() for item in rows[0]]
    if not any(headers):
        raise ValueError("empty_roster_file")

    result: list[dict[str, Any]] = []
    for values in rows[1:]:
        row: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = values[idx] if idx < len(values) else ""
            row[header] = str(value or "").strip()
        if row:
            result.append(row)
    return result


def _map_roster_row(raw_row: dict[str, Any]) -> dict[str, str] | None:
    full_name = _pick_field(raw_row, _NAME_HEADERS)
    id_card_number = _normalize_id_card(_pick_field(raw_row, _ID_HEADERS))
    group_name = _pick_field(raw_row, _GROUP_HEADERS)
    if not full_name or not id_card_number or not is_valid_id_card(id_card_number):
        return None
    id_hash = _id_hash(id_card_number)
    return {
        "full_name": full_name,
        "id_last4": "",
        "id_hash": id_hash,
        "group_name": group_name,
    }


def import_villager_roster(
    token: str,
    file_name: str,
    file_bytes: bytes,
) -> dict[str, Any]:
    ext = (file_name.rsplit(".", 1)[-1] if "." in file_name else "").lower()
    if ext not in {"csv", "xlsx"}:
        raise ValueError("unsupported_file_type")

    if ext == "csv":
        raw_rows = _parse_csv_rows(file_bytes)
    else:
        raw_rows = _parse_xlsx_rows(file_bytes)

    conn = get_connection()
    try:
        user = _require_cadre_user_conn(conn, token)

        entries: list[dict[str, str]] = []
        skipped_count = 0
        de_dup_set: set[tuple[str, str, str, str]] = set()
        for raw_row in raw_rows:
            row = _map_roster_row(raw_row)
            if row is None:
                skipped_count += 1
                continue
            key = (
                row["full_name"],
                row["id_last4"],
                row["id_hash"],
                row["group_name"],
            )
            if key in de_dup_set:
                skipped_count += 1
                continue
            de_dup_set.add(key)
            entries.append(row)

        if not entries:
            raise ValueError("empty_roster_file")

        now = now_iso()
        imported_by = user["real_name"] or user["nickname"] or user["openid"]
        cursor = conn.execute(
            """
            INSERT INTO roster_import_batches(filename, imported_by, imported_count, skipped_count, created_at)
            VALUES(?, ?, 0, 0, ?)
            """,
            (file_name, imported_by, now),
        )
        batch_id = cursor.lastrowid

        conn.execute("UPDATE villager_roster SET status = 'inactive' WHERE status = 'active'")
        conn.executemany(
            """
            INSERT INTO villager_roster(
                batch_id, full_name, id_last4, id_hash, group_name, status, created_at
            )
            VALUES(?, ?, ?, ?, ?, 'active', ?)
            """,
            [
                (
                    batch_id,
                    row["full_name"],
                    row["id_last4"],
                    row["id_hash"],
                    row["group_name"],
                    now,
                )
                for row in entries
            ],
        )
        conn.execute(
            """
            UPDATE roster_import_batches
            SET imported_count = ?, skipped_count = ?
            WHERE id = ?
            """,
            (len(entries), skipped_count, batch_id),
        )
        _audit_log_conn(
            conn,
            user["id"],
            "roster_import",
            "roster_batch",
            str(batch_id),
            {
                "filename": file_name,
                "imported_count": len(entries),
                "skipped_count": skipped_count,
            },
        )
        conn.commit()
        return {
            "batch_id": batch_id,
            "filename": file_name,
            "imported_count": len(entries),
            "skipped_count": skipped_count,
            "created_at": now,
        }
    finally:
        conn.close()


def list_roster_import_batches(token: str, limit: int = 20) -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        _require_cadre_user_conn(conn, token)
        rows = conn.execute(
            """
            SELECT id, filename, imported_by, imported_count, skipped_count, created_at
            FROM roster_import_batches
            ORDER BY id DESC
            LIMIT ?
            """,
            (max(1, min(limit, 100)),),
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def _find_roster_match_conn(
    conn: sqlite3.Connection,
    real_name: str,
    id_hash: str,
) -> sqlite3.Row | None:
    return conn.execute(
        """
        SELECT id, group_name
        FROM villager_roster
        WHERE status = 'active'
          AND full_name = ?
          AND id_hash = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        (real_name, id_hash),
    ).fetchone()


def _require_active_roster_conn(conn: sqlite3.Connection | MySQLConnectionAdapter) -> None:
    row = conn.execute(
        "SELECT COUNT(1) AS cnt FROM villager_roster WHERE status = 'active'"
    ).fetchone()
    active_count = int(row["cnt"] or 0) if row is not None else 0
    if active_count <= 0:
        raise ValueError("roster_not_ready")


def _find_existing_identity_binding_conn(
    conn: sqlite3.Connection | MySQLConnectionAdapter,
    id_hash: str,
    *,
    exclude_user_id: int = 0,
) -> sqlite3.Row | dict[str, Any] | None:
    if not str(id_hash or "").strip():
        return None
    return conn.execute(
        """
        SELECT id, real_name, phone_masked, auth_status, updated_at
        FROM users
        WHERE id_card_hash = ?
          AND id <> ?
          AND realname_status = 'verified'
        ORDER BY id DESC
        LIMIT 1
        """,
        (id_hash, int(exclude_user_id or 0)),
    ).fetchone()


def _create_token_conn(conn: sqlite3.Connection, user_id: int) -> tuple[str, str]:
    now = now_iso()
    token = secrets.token_urlsafe(32)
    token_hash = _hash_session_token(token)
    expires_at = (datetime.now(timezone.utc) + timedelta(days=7)).isoformat()
    conn.execute("DELETE FROM auth_sessions WHERE expires_at <= ?", (now,))
    conn.execute("DELETE FROM auth_sessions WHERE user_id = ?", (user_id,))
    conn.execute(
        "INSERT INTO auth_sessions(user_id, token, created_at, expires_at) VALUES(?, ?, ?, ?)",
        (user_id, token_hash, now, expires_at),
    )
    return token, expires_at


def _build_mock_wechat_identity(code: str) -> dict[str, str]:
    if MOCK_WECHAT_OPENID and APP_ENV != "production":
        return {
            "openid": MOCK_WECHAT_OPENID,
            "unionid": MOCK_WECHAT_UNIONID,
        }
    return {
        "openid": f"mock_{hashlib.sha256(code.encode('utf-8')).hexdigest()[:24]}",
        "unionid": "",
    }


def _exchange_wechat_code_for_identity(code: str) -> dict[str, str]:
    if ALLOW_MOCK_WECHAT_LOGIN:
        if APP_ENV == "production":
            raise ValueError("wechat_login_unconfigured")
        return _build_mock_wechat_identity(code)
    if not WECHAT_MINI_APP_ID or not WECHAT_MINI_APP_SECRET:
        raise ValueError("wechat_login_unconfigured")
    query = urlencode(
        {
            "appid": WECHAT_MINI_APP_ID,
            "secret": WECHAT_MINI_APP_SECRET,
            "js_code": code,
            "grant_type": "authorization_code",
        }
    )
    request_url = f"{WECHAT_CODE2SESSION_URL}?{query}"
    try:
        with urlopen(request_url, timeout=8) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise ValueError("wechat_service_unavailable") from exc
    if not isinstance(payload, dict):
        raise ValueError("wechat_service_unavailable")
    errcode = int(payload.get("errcode") or 0)
    if errcode:
        if errcode in {40029, 40163}:
            raise ValueError("invalid_wechat_code")
        raise ValueError("wechat_service_unavailable")
    openid = str(payload.get("openid") or "").strip()
    if not openid:
        raise ValueError("wechat_service_unavailable")
    return {
        "openid": openid,
        "unionid": str(payload.get("unionid") or "").strip(),
    }


def _request_json(
    *,
    url: str,
    payload: dict[str, Any] | None = None,
    headers: dict[str, str] | None = None,
    timeout: int = 8,
) -> dict[str, Any]:
    request_headers = dict(headers or {})
    body = None
    method = "GET"
    if payload is not None:
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        request_headers.setdefault("Content-Type", "application/json; charset=utf-8")
        method = "POST"
    request_obj = UrlRequest(url, data=body, headers=request_headers, method=method)
    raw_bytes = b""
    try:
        with urlopen(request_obj, timeout=timeout) as response:
            raw_bytes = response.read()
    except HTTPError as exc:
        raw_bytes = exc.read() if hasattr(exc, "read") else b""
        if not raw_bytes:
            raise ValueError("upstream_service_unavailable") from exc
    except (URLError, TimeoutError) as exc:
        raise ValueError("upstream_service_unavailable") from exc
    try:
        data = json.loads(raw_bytes.decode("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise ValueError("upstream_service_unavailable") from exc
    if not isinstance(data, dict):
        raise ValueError("upstream_service_unavailable")
    return data


def _get_wechat_access_token(
    *,
    unconfigured_error: str = "wechat_phone_unconfigured",
    unavailable_error: str = "wechat_phone_unavailable",
) -> str:
    cached_token = str(_WECHAT_ACCESS_TOKEN_CACHE.get("access_token") or "").strip()
    expires_at = float(_WECHAT_ACCESS_TOKEN_CACHE.get("expires_at") or 0.0)
    if cached_token and time.time() + 60 < expires_at:
        return cached_token
    if not WECHAT_MINI_APP_ID or not WECHAT_MINI_APP_SECRET:
        raise ValueError(unconfigured_error)
    payload = _request_json(
        url=WECHAT_STABLE_ACCESS_TOKEN_URL,
        payload={
            "grant_type": "client_credential",
            "appid": WECHAT_MINI_APP_ID,
            "secret": WECHAT_MINI_APP_SECRET,
            "force_refresh": False,
        },
        timeout=8,
    )
    errcode = int(payload.get("errcode") or 0)
    if errcode:
        logger.warning(
            "wechat_access_token_failed: errcode=%s errmsg=%s appid=%s",
            errcode,
            str(payload.get("errmsg") or "").strip(),
            WECHAT_MINI_APP_ID,
        )
        raise ValueError(unavailable_error)
    access_token = str(payload.get("access_token") or "").strip()
    expires_in = int(payload.get("expires_in") or 0)
    if not access_token or expires_in <= 0:
        logger.warning(
            "wechat_access_token_invalid_payload: appid=%s payload_keys=%s expires_in=%s",
            WECHAT_MINI_APP_ID,
            sorted(str(key) for key in payload.keys()),
            expires_in,
        )
        raise ValueError(unavailable_error)
    _WECHAT_ACCESS_TOKEN_CACHE["access_token"] = access_token
    _WECHAT_ACCESS_TOKEN_CACHE["expires_at"] = time.time() + max(expires_in - 120, 300)
    return access_token


def _read_cloudbase_access_token() -> str:
    try:
        token = Path(CLOUDBASE_ACCESS_TOKEN_FILE).read_text(encoding="utf-8").strip()
    except OSError:
        return ""
    return token


def _as_cloudbase_openapi_url(url: str) -> str:
    text = str(url or "").strip()
    if text.startswith("https://api.weixin.qq.com/"):
        return "http://api.weixin.qq.com/" + text.split("https://api.weixin.qq.com/", 1)[1]
    return text


def _request_wechat_openapi_json(
    *,
    url: str,
    payload: dict[str, Any] | None = None,
    timeout: int = 8,
    unconfigured_error: str = "wechat_phone_unconfigured",
    unavailable_error: str = "wechat_phone_unavailable",
) -> dict[str, Any]:
    if USE_WECHAT_OPENAPI_SERVICE:
        try:
            response = _request_json(
                url=_as_cloudbase_openapi_url(url),
                payload=payload,
                timeout=timeout,
            )
            logger.info("wechat_openapi_service_used: url=%s", url)
            return response
        except ValueError as exc:
            logger.warning(
                "wechat_openapi_service_failed: url=%s detail=%s",
                url,
                str(exc),
            )
    token_param_name, token_value = _get_wechat_openapi_auth(
        unconfigured_error=unconfigured_error,
        unavailable_error=unavailable_error,
    )
    return _request_json(
        url=f"{url}?{token_param_name}={quote(token_value)}",
        payload=payload,
        timeout=timeout,
    )


def _get_wechat_openapi_auth(
    *,
    unconfigured_error: str = "wechat_phone_unconfigured",
    unavailable_error: str = "wechat_phone_unavailable",
) -> tuple[str, str]:
    cloudbase_access_token = _read_cloudbase_access_token()
    if cloudbase_access_token:
        return "cloudbase_access_token", cloudbase_access_token
    return "access_token", _get_wechat_access_token(
        unconfigured_error=unconfigured_error,
        unavailable_error=unavailable_error,
    )


def _exchange_wechat_phone_code(phone_code: str) -> str:
    code_text = str(phone_code or "").strip()
    if not code_text:
        raise ValueError("invalid_wechat_phone_code")
    try:
        payload = _request_wechat_openapi_json(
            url=WECHAT_GET_USER_PHONE_URL,
            payload={"code": code_text},
            timeout=8,
            unconfigured_error="wechat_phone_unconfigured",
            unavailable_error="wechat_phone_unavailable",
        )
    except ValueError as exc:
        detail = str(exc)
        if detail == "upstream_service_unavailable":
            logger.warning("wechat_phone_exchange_upstream_unavailable")
            raise ValueError("wechat_phone_unavailable") from exc
        raise
    errcode = int(payload.get("errcode") or 0)
    if errcode:
        logger.warning(
            "wechat_phone_exchange_failed: errcode=%s errmsg=%s",
            errcode,
            str(payload.get("errmsg") or "").strip(),
        )
        if errcode in {40029, 40163, 61013}:
            raise ValueError("invalid_wechat_phone_code")
        raise ValueError("wechat_phone_unavailable")
    phone_info = payload.get("phone_info") or {}
    if not isinstance(phone_info, dict):
        logger.warning(
            "wechat_phone_exchange_missing_phone_info: payload_keys=%s",
            sorted(str(key) for key in payload.keys()),
        )
        raise ValueError("wechat_phone_unavailable")
    phone = normalize_phone_number(
        phone_info.get("phoneNumber")
        or phone_info.get("purePhoneNumber")
        or ""
    )
    if not is_valid_phone(phone):
        logger.warning(
            "wechat_phone_exchange_invalid_phone: payload_keys=%s phone_length=%s",
            sorted(str(key) for key in payload.keys()),
            len(phone),
        )
        raise ValueError("wechat_phone_unavailable")
    return phone


def logout_session(token: str) -> dict[str, str]:
    conn = get_connection()
    try:
        row = _get_user_by_token_conn(conn, token)
        session_id = int(row["session_id"] or 0)
        if session_id:
            conn.execute("DELETE FROM auth_sessions WHERE id = ?", (session_id,))
        conn.execute(
            "UPDATE cadre_stepup_sessions SET revoked_at = ? WHERE user_id = ? AND revoked_at = ''",
            (now_iso(), row["id"]),
        )
        _audit_log_conn(
            conn,
            int(row["id"] or 0),
            "session_logout",
            "auth_session",
            str(session_id),
            {},
        )
        conn.commit()
        return {"status": "ok", "message": "已退出登录"}
    finally:
        conn.close()


def wechat_login(payload: dict[str, Any]) -> dict[str, Any]:
    nickname = (payload.get("nickname") or "").strip()
    avatar_url = (payload.get("avatar_url") or "").strip()
    cloudbase_identity = payload.get("cloudbase_identity")
    if isinstance(cloudbase_identity, dict) and str(cloudbase_identity.get("openid") or "").strip():
        identity = {
            "openid": str(cloudbase_identity.get("openid") or "").strip(),
            "unionid": str(cloudbase_identity.get("unionid") or "").strip(),
        }
    else:
        code = (payload.get("code") or "").strip()
        if not code:
            raise ValueError("invalid_wechat_code")
        identity = _exchange_wechat_code_for_identity(code)
    openid = identity["openid"]
    unionid = identity.get("unionid", "")

    conn = get_connection()
    try:
        row = _ensure_user_by_wechat_identity_conn(
            conn,
            openid=openid,
            unionid=unionid,
            nickname=nickname,
            avatar_url=avatar_url,
        )
        user_id = int(row["id"] or 0)
        token, expires_at = _create_token_conn(conn, user_id)
        _audit_log_conn(
            conn,
            user_id,
            "wechat_login",
            "user",
            str(user_id),
            {"login_type": "cloudbase" if cloudbase_identity else "wechat"},
        )
        user_data = _build_auth_user_conn(conn, user_id)
        conn.commit()
        return {
            "token": token,
            "expires_at": expires_at,
            "user": user_data,
        }
    finally:
        conn.close()


def cadre_login(payload: dict[str, Any]) -> dict[str, Any]:
    _ = payload
    raise ValueError("cadre_password_login_disabled")


def bind_phone(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    phone_code = str(payload.get("phone_code") or "").strip()
    if not phone_code:
        raise ValueError("phone_bind_requires_wechat_verify")
    phone = _exchange_wechat_phone_code(phone_code)
    verified_at = now_iso()

    conn = get_connection()
    try:
        row = _get_user_by_token_conn(conn, token)
        _require_consent_conn(conn, row["id"], "privacy_policy", "privacy_consent_required")
        auth_status = calc_auth_status(
            bool(verified_at),
            row["realname_status"],
            row["face_status"],
        )
        now = verified_at or now_iso()
        phone_fields = _phone_fields(phone)
        conn.execute(
            """
            UPDATE users
            SET
                phone = ?,
                phone_ciphertext = ?,
                phone_hash = ?,
                phone_masked = ?,
                phone_last4 = ?,
                phone_bound = 1,
                phone_verified_at = ?,
                auth_status = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                phone_fields["phone"],
                phone_fields["phone_ciphertext"],
                phone_fields["phone_hash"],
                phone_fields["phone_masked"],
                phone_fields["phone_last4"],
                verified_at,
                auth_status,
                now,
                row["id"],
            ),
        )
        _sync_villager_registry_conn(conn, row["id"])
        _audit_log_conn(
            conn,
            row["id"],
            "bind_phone",
            "user",
            str(row["id"]),
            {"phone_masked": mask_phone(phone), "method": "wechat_phone_component"},
        )
        user_data = _build_auth_user_conn(conn, int(row["id"] or 0))
        conn.commit()
        return user_data
    finally:
        conn.close()


def submit_realname(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    real_name = (payload.get("real_name") or "").strip()
    id_card_number = _normalize_id_card(payload.get("id_card_number") or "")
    if not real_name:
        raise ValueError("invalid_real_name")
    if not is_valid_id_card(id_card_number):
        raise ValueError("invalid_id_card")

    conn = get_connection()
    try:
        row = _get_user_by_token_conn(conn, token)
        _require_consent_conn(conn, row["id"], "privacy_policy", "privacy_consent_required")
        _require_consent_conn(
            conn,
            row["id"],
            "sensitive_identity",
            "sensitive_consent_required",
        )
        if not _is_phone_verified_user(row):
            raise ValueError("phone_bind_requires_wechat_verify")

        gender = _resolve_gender(row["gender"] or "", id_card_number)
        if len(gender) > 10:
            raise ValueError("invalid_gender")
        id_card_hash = _id_hash(id_card_number)
        id_card_encrypted = encrypt_id_card_number(id_card_number)
        _require_active_roster_conn(conn)
        match_row = _find_roster_match_conn(conn, real_name, id_card_hash)
        if match_row is None:
            _audit_log_conn(
                conn,
                row["id"],
                "identity_match_failed",
                "user",
                str(row["id"]),
                {"real_name": real_name, "id_card_hash": id_card_hash},
            )
            conn.commit()
            raise ValueError("identity_not_in_roster")
        conflict_row = _find_existing_identity_binding_conn(
            conn,
            id_card_hash,
            exclude_user_id=int(row["id"] or 0),
        )
        if conflict_row is not None:
            _audit_log_conn(
                conn,
                row["id"],
                "identity_bind_conflict",
                "user",
                str(row["id"]),
                {
                    "real_name": real_name,
                    "id_card_hash": id_card_hash,
                    "conflict_user_id": str(conflict_row["id"] or ""),
                },
            )
            conn.commit()
            raise ValueError("id_card_already_verified")

        roster_match_status = "matched"
        village_group = match_row["group_name"] or ""
        roster_match_id = str(match_row["id"])

        realname_status = "verified"
        face_status = "unsubmitted"
        auth_status = calc_auth_status(
            _is_phone_verified_user(row),
            realname_status,
            face_status,
        )
        now = now_iso()
        conn.execute(
            """
            UPDATE users
            SET
                real_name = ?,
                gender = ?,
                id_card_masked = ?,
                id_card_last4 = ?,
                id_card_hash = ?,
                id_card_encrypted = ?,
                role_request = 'villager',
                realname_status = ?,
                face_status = ?,
                auth_status = ?,
                roster_match_status = ?,
                village_group = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                real_name,
                gender,
                mask_id_card(id_card_number),
                "",
                id_card_hash,
                id_card_encrypted,
                realname_status,
                face_status,
                auth_status,
                roster_match_status,
                village_group,
                now,
                row["id"],
            ),
        )
        _sync_villager_registry_conn(conn, row["id"])
        _audit_log_conn(
            conn,
            row["id"],
            "submit_realname",
            "user",
            str(row["id"]),
            {
                "real_name": real_name,
                "roster_match_status": roster_match_status,
                "roster_match_id": roster_match_id,
                "village_group": village_group,
            },
        )
        user_data = _build_auth_user_conn(conn, int(row["id"] or 0))
        conn.commit()
        return user_data
    finally:
        conn.close()


def _gov_service_to_dict(
    row: sqlite3.Row | dict[str, Any],
    *,
    include_payload: bool = False,
) -> dict[str, Any]:
    item = dict(row)
    item["is_enabled"] = bool(item.get("is_enabled", 0))
    item["detail_payload"] = _sanitize_gov_detail_payload(item.get("detail_payload"))
    if include_payload:
        return item
    item.pop("contact_name", None)
    item.pop("contact_phone", None)
    item.pop("intro", None)
    item.pop("detail_payload", None)
    item.pop("updated_at", None)
    item.pop("created_at", None)
    return item


def _gov_application_to_dict(
    row: sqlite3.Row | dict[str, Any],
    *,
    mask_sensitive: bool = True,
) -> dict[str, Any]:
    # Sensitive values are never returned to either the villager or ordinary
    # administration APIs.  The argument remains for backward-compatible
    # internal callers, but no longer permits plaintext disclosure.
    _ = mask_sensitive
    item = dict(row)
    item["village_group"] = str(item.get("village_group") or "").strip()
    item["applicant_phone"] = str(item.get("applicant_phone_masked") or "").strip() or str(
        _mask_sensitive_value("applicant_phone", item.get("applicant_phone", "")) or ""
    )
    extra_fields = _parse_json_object(item.get("extra_fields"))
    item["supporting_photo_urls"] = _parse_photo_urls_field(
        extra_fields.pop(GOV_APPLICATION_SUPPORTING_PHOTO_FIELD, []),
        allowed_scenes=SCENES_GOV_APPLY,
        allow_external=False,
    )
    item["extra_fields"] = _mask_sensitive_json_object(extra_fields)
    item["reply_photo_urls"] = _parse_photo_urls_field(
        item.get("reply_photo_urls"),
        allowed_scenes=SCENES_GOV_REPLY,
        allow_external=False,
    )
    for field_name in (
        "applicant_phone_ciphertext",
        "applicant_phone_hash",
        "applicant_phone_masked",
        "sensitive_extra_fields_encrypted",
        "applicant_id_card_hash",
        "application_fingerprint_hash",
    ):
        item.pop(field_name, None)
    return item


def _build_verified_identity_claim(user: sqlite3.Row | dict[str, Any]) -> dict[str, str]:
    data = dict(user)
    real_name = str(data.get("real_name") or "").strip()
    id_hash = str(data.get("id_card_hash") or "").strip()
    realname_status = str(data.get("realname_status") or "").strip().lower()
    if realname_status != "verified" or not real_name or not id_hash:
        return {}
    return {"real_name": real_name, "id_hash": id_hash}


def _gov_application_matches_identity_claim(
    row: sqlite3.Row | dict[str, Any],
    identity_claim: dict[str, str],
) -> bool:
    if not identity_claim:
        return False
    item = dict(row)
    real_name = str(identity_claim.get("real_name") or "").strip()
    id_hash = str(identity_claim.get("id_hash") or "").strip()
    if not real_name or not id_hash:
        return False
    applicant_name = str(item.get("applicant_name") or "").strip()
    extra_fields = _parse_json_object(item.get("extra_fields"))
    extra_name = str(extra_fields.get("name") or applicant_name).strip()
    stored_id_hash = str(item.get("applicant_id_card_hash") or "").strip()
    if not stored_id_hash:
        legacy_id_card = _normalize_id_card(str(extra_fields.get("id_card") or ""))
        stored_id_hash = _id_hash(legacy_id_card) if legacy_id_card else ""
    if applicant_name != real_name or extra_name != real_name or not stored_id_hash:
        return False
    return stored_id_hash == id_hash


def _gov_application_visible_to_user(
    row: sqlite3.Row | dict[str, Any],
    user: sqlite3.Row | dict[str, Any],
    identity_claim: dict[str, str] | None = None,
) -> bool:
    item = dict(row)
    data = dict(user)
    if int(item.get("user_id") or 0) == int(data.get("id") or 0):
        return True
    return _gov_application_matches_identity_claim(row, identity_claim or _build_verified_identity_claim(user))


def _china_day_bounds_utc(base_time: datetime | None = None) -> tuple[str, str]:
    local_now = (base_time or datetime.now(timezone.utc)).astimezone(CHINA_TIMEZONE)
    start_local = local_now.replace(hour=0, minute=0, second=0, microsecond=0)
    end_local = start_local + timedelta(days=1)
    return start_local.astimezone(timezone.utc).isoformat(), end_local.astimezone(timezone.utc).isoformat()


def _normalize_certificate_fingerprint_value(key: str, value: Any) -> str:
    text = str(value or "").strip()
    if key in {"certificate_key"}:
        return text.lower()
    if key in {"id_card", "related_id_card"}:
        return text.upper()
    return text


def _build_certificate_application_fingerprint(
    *,
    service_id: int,
    selected_option: str,
    extra_fields: dict[str, Any],
) -> str:
    normalized: dict[str, Any] = {
        "service_id": int(service_id or 0),
        "selected_option": str(selected_option or "").strip(),
    }
    for key in CERTIFICATE_DUPLICATE_FIELD_KEYS:
        value = _normalize_certificate_fingerprint_value(key, extra_fields.get(key))
        if value:
            normalized[key] = value
    return json.dumps(normalized, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _count_today_certificate_applications_conn(conn: sqlite3.Connection, user_id: int) -> int:
    start_utc, end_utc = _china_day_bounds_utc()
    row = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM gov_service_applications
        WHERE user_id = ?
          AND service_key = 'certificate'
          AND created_at >= ?
          AND created_at < ?
        """,
        (int(user_id or 0), start_utc, end_utc),
    ).fetchone()
    return int((row["total"] if row is not None else 0) or 0)


def _count_today_certificate_applications_for_user_conn(
    conn: sqlite3.Connection,
    user: sqlite3.Row | dict[str, Any],
) -> int:
    data = dict(user)
    identity_claim = _build_verified_identity_claim(data)
    start_utc, end_utc = _china_day_bounds_utc()
    if identity_claim:
        rows = conn.execute(
            """
            SELECT id, user_id, applicant_name, extra_fields, applicant_id_card_hash
            FROM gov_service_applications
            WHERE service_key = 'certificate'
              AND created_at >= ?
              AND created_at < ?
              AND (user_id = ? OR applicant_name = ?)
            """,
            (
                start_utc,
                end_utc,
                int(data.get("id") or 0),
                identity_claim["real_name"],
            ),
        ).fetchall()
        return sum(
            1
            for row in rows
            if _gov_application_visible_to_user(row, data, identity_claim)
        )
    return _count_today_certificate_applications_conn(conn, int(data.get("id") or 0))


def _has_duplicate_certificate_application_conn(
    conn: sqlite3.Connection,
    *,
    user_id: int,
    service_id: int,
    selected_option: str,
    extra_fields: dict[str, Any],
) -> bool:
    target_fingerprint = _build_certificate_application_fingerprint(
        service_id=service_id,
        selected_option=selected_option,
        extra_fields=extra_fields,
    )
    target_fingerprint_hash = keyed_lookup_hash(
        target_fingerprint,
        purpose="gov_application_fingerprint",
    )
    rows = conn.execute(
        """
        SELECT selected_option, extra_fields, application_fingerprint_hash
        FROM gov_service_applications
        WHERE user_id = ?
          AND service_id = ?
          AND service_key = 'certificate'
          AND status IN ('pending', 'processing')
        ORDER BY id DESC
        """,
        (int(user_id or 0), int(service_id or 0)),
    ).fetchall()
    for row in rows:
        stored_fingerprint_hash = str(row["application_fingerprint_hash"] or "").strip()
        if stored_fingerprint_hash:
            if hmac.compare_digest(stored_fingerprint_hash, target_fingerprint_hash):
                return True
            continue
        existing_extra_fields = _parse_json_object(row["extra_fields"])
        existing_extra_fields.pop(GOV_APPLICATION_SUPPORTING_PHOTO_FIELD, None)
        if (
            _build_certificate_application_fingerprint(
                service_id=service_id,
                selected_option=row["selected_option"],
                extra_fields=existing_extra_fields,
            )
            == target_fingerprint
        ):
            return True
    return False


def _has_duplicate_certificate_application_for_user_conn(
    conn: sqlite3.Connection,
    *,
    user: sqlite3.Row | dict[str, Any],
    service_id: int,
    selected_option: str,
    extra_fields: dict[str, Any],
) -> bool:
    data = dict(user)
    identity_claim = _build_verified_identity_claim(data)
    target_fingerprint = _build_certificate_application_fingerprint(
        service_id=service_id,
        selected_option=selected_option,
        extra_fields=extra_fields,
    )
    target_fingerprint_hash = keyed_lookup_hash(
        target_fingerprint,
        purpose="gov_application_fingerprint",
    )
    if identity_claim:
        rows = conn.execute(
            """
            SELECT user_id, applicant_name, selected_option, extra_fields,
                   applicant_id_card_hash, application_fingerprint_hash
            FROM gov_service_applications
            WHERE service_id = ?
              AND service_key = 'certificate'
              AND status IN ('pending', 'processing')
              AND (user_id = ? OR applicant_name = ?)
            ORDER BY id DESC
            """,
            (
                int(service_id or 0),
                int(data.get("id") or 0),
                identity_claim["real_name"],
            ),
        ).fetchall()
        rows = [
            row for row in rows if _gov_application_visible_to_user(row, data, identity_claim)
        ]
    else:
        rows = conn.execute(
            """
            SELECT selected_option, extra_fields, application_fingerprint_hash
            FROM gov_service_applications
            WHERE user_id = ?
              AND service_id = ?
              AND service_key = 'certificate'
              AND status IN ('pending', 'processing')
            ORDER BY id DESC
            """,
            (int(data.get("id") or 0), int(service_id or 0)),
        ).fetchall()
    for row in rows:
        stored_fingerprint_hash = str(row["application_fingerprint_hash"] or "").strip()
        if stored_fingerprint_hash:
            if hmac.compare_digest(stored_fingerprint_hash, target_fingerprint_hash):
                return True
            continue
        existing_extra_fields = _parse_json_object(row["extra_fields"])
        existing_extra_fields.pop(GOV_APPLICATION_SUPPORTING_PHOTO_FIELD, None)
        if (
            _build_certificate_application_fingerprint(
                service_id=service_id,
                selected_option=row["selected_option"],
                extra_fields=existing_extra_fields,
            )
            == target_fingerprint
        ):
            return True
    return False


def _seed_default_gov_services_conn(conn: sqlite3.Connection) -> None:
    existing_rows = conn.execute(
        """
        SELECT id, service_key, detail_payload
        FROM gov_services
        ORDER BY sort_order ASC, id ASC
        """
    ).fetchall()
    existing_by_key = {str(row["service_key"]): row for row in existing_rows}
    now = now_iso()
    if not existing_rows:
        conn.executemany(
            """
            INSERT INTO gov_services(
                service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, action_label, detail_type,
                detail_payload, is_enabled, sort_order, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?)
            """,
            [
                (
                    item["service_key"],
                    item["title"],
                    item.get("subtitle", ""),
                    item.get("category", "government"),
                    item.get("icon_emoji", ""),
                    item.get("theme_color", ""),
                    item.get("contact_name", ""),
                    item.get("contact_phone", ""),
                    item.get("intro", ""),
                    item.get("action_label", "我要申请"),
                    item.get("detail_type", "info_apply"),
                    _json_text(_normalize_json_object(item.get("detail_payload"))),
                    index,
                    now,
                    now,
                )
                for index, item in enumerate(DEFAULT_GOV_SERVICES, start=1)
            ],
        )
        return

    # Existing DB compatibility:
    # 1) insert missing default service keys
    # 2) append new certificate options to existing "我要开证明" payload
    max_sort_row = conn.execute("SELECT COALESCE(MAX(sort_order), 0) AS max_sort FROM gov_services").fetchone()
    next_sort = int(max_sort_row["max_sort"]) + 1 if max_sort_row else 1
    missing_items: list[tuple[Any, ...]] = []
    for item in DEFAULT_GOV_SERVICES:
        if item["service_key"] in existing_by_key:
            continue
        missing_items.append(
            (
                item["service_key"],
                item["title"],
                item.get("subtitle", ""),
                item.get("category", "government"),
                item.get("icon_emoji", ""),
                item.get("theme_color", ""),
                item.get("contact_name", ""),
                item.get("contact_phone", ""),
                item.get("intro", ""),
                item.get("action_label", "我要申请"),
                item.get("detail_type", "info_apply"),
                _json_text(_normalize_json_object(item.get("detail_payload"))),
                next_sort,
                now,
                now,
            )
        )
        next_sort += 1
    if missing_items:
        conn.executemany(
            """
            INSERT INTO gov_services(
                service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, action_label, detail_type,
                detail_payload, is_enabled, sort_order, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?)
            """,
            missing_items,
        )

    cert_default = next(
        (item for item in DEFAULT_GOV_SERVICES if item.get("service_key") == "certificate"),
        None,
    )
    cert_row = existing_by_key.get("certificate")
    if cert_default and cert_row:
        default_payload = _normalize_json_object(cert_default.get("detail_payload"))
        default_certs = default_payload.get("certificates") or []
        current_payload = _parse_json_object(cert_row["detail_payload"])
        current_certs = current_payload.get("certificates") if isinstance(current_payload, dict) else []
        if not isinstance(current_certs, list):
            current_certs = []
        existing_cert_keys = {
            str(item.get("key", "")).strip()
            for item in current_certs
            if isinstance(item, dict) and str(item.get("key", "")).strip()
        }
        merged_certs = list(current_certs)
        for item in default_certs:
            key = str(item.get("key", "")).strip() if isinstance(item, dict) else ""
            if not key or key in existing_cert_keys:
                continue
            merged_certs.append(item)
            existing_cert_keys.add(key)
        if len(merged_certs) != len(current_certs):
            merged_payload = dict(current_payload) if isinstance(current_payload, dict) else {}
            merged_payload["certificates"] = merged_certs
            conn.execute(
                """
                UPDATE gov_services
                SET detail_payload = ?, updated_at = ?
                WHERE id = ?
                """,
                (_json_text(merged_payload), now, cert_row["id"]),
            )

    rows = conn.execute(
        """
        SELECT id, service_key, contact_name, contact_phone, detail_payload
        FROM gov_services
        """
    ).fetchall()
    for row in rows:
        service_key = str(row["service_key"] or "").strip()
        current_contact_name = str(row["contact_name"] or "").strip()
        current_contact_phone = re.sub(r"\D", "", str(row["contact_phone"] or "")).strip()
        next_contact_name = (
            GENERIC_GOV_CONTACT_NAME if current_contact_name in LEGACY_GOV_CONTACT_NAMES else current_contact_name
        )
        next_contact_phone = (
            GENERIC_GOV_CONTACT_PHONE if current_contact_phone in LEGACY_GOV_CONTACT_PHONES else current_contact_phone
        )
        current_payload = _parse_json_object(row["detail_payload"])
        next_payload = dict(current_payload) if isinstance(current_payload, dict) else {}
        payload_changed = False
        jobs = next_payload.get("jobs")
        if isinstance(jobs, list):
            normalized_jobs: list[Any] = []
            job_keys: set[str] = set()
            for item in jobs:
                if not isinstance(item, dict):
                    normalized_jobs.append(item)
                    continue
                job_item = dict(item)
                job_key = str(job_item.get("job_key") or "").strip()
                if job_key:
                    job_keys.add(job_key)
                job_contact_name = str(job_item.get("contact_name") or "").strip()
                job_contact_phone = re.sub(r"\D", "", str(job_item.get("contact_phone") or "")).strip()
                if job_contact_name in LEGACY_GOV_CONTACT_NAMES:
                    job_item["contact_name"] = GENERIC_GOV_CONTACT_NAME
                    payload_changed = True
                if job_contact_phone in LEGACY_GOV_CONTACT_PHONES:
                    job_item["contact_phone"] = GENERIC_GOV_CONTACT_PHONE
                    payload_changed = True
                normalized_jobs.append(job_item)
            if service_key == "employment_training" and job_keys and job_keys.issubset(LEGACY_EMPLOYMENT_JOB_KEYS):
                normalized_jobs = []
                payload_changed = True
            if payload_changed:
                next_payload["jobs"] = normalized_jobs
        if (
            next_contact_name == current_contact_name
            and next_contact_phone == current_contact_phone
            and not payload_changed
        ):
            continue
        conn.execute(
            """
            UPDATE gov_services
            SET contact_name = ?, contact_phone = ?, detail_payload = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                next_contact_name[:50],
                next_contact_phone[:20],
                _json_text(next_payload if payload_changed else current_payload),
                now,
                row["id"],
            ),
        )


def list_gov_services(token: str | None = None, cadre_token: str | None = None) -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn_optional(conn, token)
        _seed_default_gov_services_conn(conn)
        is_cadre = False
        if user is not None and str(cadre_token or "").strip():
            access = _get_cadre_authorization_context_conn(conn, user)
            if access is not None:
                _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
                is_cadre = True
        if is_cadre:
            rows = conn.execute(
                """
                SELECT
                    id, service_key, title, subtitle, category, icon_emoji, theme_color,
                    detail_type, action_label, is_enabled, sort_order, created_at, updated_at,
                    contact_name, contact_phone, intro, detail_payload
                FROM gov_services
                ORDER BY sort_order ASC, id ASC
                """
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT
                    id, service_key, title, subtitle, category, icon_emoji, theme_color,
                    detail_type, action_label, is_enabled, sort_order, created_at, updated_at,
                    contact_name, contact_phone, intro, detail_payload
                FROM gov_services
                WHERE is_enabled = 1
                ORDER BY sort_order ASC, id ASC
                """
            ).fetchall()
        conn.commit()
        return [
            _gov_service_to_dict(row, include_payload=False)
            for row in rows
            if str(row["service_key"] or "") not in HIDDEN_GOV_SERVICE_KEYS
        ]
    finally:
        conn.close()


def get_gov_service(
    token: str | None,
    service_id: int,
    cadre_token: str | None = None,
) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn_optional(conn, token)
        _seed_default_gov_services_conn(conn)
        is_cadre = False
        if user is not None and str(cadre_token or "").strip():
            access = _get_cadre_authorization_context_conn(conn, user)
            if access is not None:
                _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
                is_cadre = True
        params: tuple[Any, ...] = (service_id,)
        sql = """
            SELECT
                id, service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, detail_type, detail_payload,
                action_label, is_enabled, sort_order, updated_at, created_at
            FROM gov_services
            WHERE id = ?
        """
        if not is_cadre:
            sql += " AND is_enabled = 1"
        sql += " LIMIT 1"
        row = conn.execute(sql, params).fetchone()
        conn.commit()
        if row is None or str(row["service_key"] or "") in HIDDEN_GOV_SERVICE_KEYS:
            raise ValueError("gov_service_not_found")
        return _gov_service_to_dict(row, include_payload=True)
    finally:
        conn.close()


def update_gov_service(token: str, service_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    title = str(payload.get("title") or "").strip()
    subtitle = str(payload.get("subtitle") or "").strip()
    category = str(payload.get("category") or "government").strip() or "government"
    icon_emoji = str(payload.get("icon_emoji") or "").strip()
    theme_color = str(payload.get("theme_color") or "").strip()
    contact_name = str(payload.get("contact_name") or "").strip()
    contact_phone = re.sub(r"\D", "", str(payload.get("contact_phone") or "")).strip()[:20]
    intro = str(payload.get("intro") or "").strip()
    detail_type = str(payload.get("detail_type") or "info_apply").strip()
    action_label = str(payload.get("action_label") or "").strip() or "我要申请"
    sort_order = int(payload.get("sort_order") or 0)
    is_enabled = bool(payload.get("is_enabled", True))
    detail_payload = _sanitize_gov_detail_payload(payload.get("detail_payload"), strict_scene=True)

    if not title:
        raise ValueError("invalid_gov_service_payload")
    if detail_type not in VALID_GOV_DETAIL_TYPES:
        raise ValueError("invalid_gov_service_payload")
    if contact_phone and not re.fullmatch(r"\d{6,20}", contact_phone):
        raise ValueError("invalid_phone")

    conn = get_connection()
    try:
        cadre_user = _require_cadre_user_conn(conn, token)
        _seed_default_gov_services_conn(conn)
        row = conn.execute(
            "SELECT id FROM gov_services WHERE id = ? LIMIT 1",
            (service_id,),
        ).fetchone()
        if row is None:
            raise ValueError("gov_service_not_found")

        now = now_iso()
        conn.execute(
            """
            UPDATE gov_services
            SET
                title = ?,
                subtitle = ?,
                category = ?,
                icon_emoji = ?,
                theme_color = ?,
                contact_name = ?,
                contact_phone = ?,
                intro = ?,
                detail_type = ?,
                detail_payload = ?,
                action_label = ?,
                is_enabled = ?,
                sort_order = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                title[:60],
                subtitle[:120],
                category[:40],
                icon_emoji[:20],
                theme_color[:20],
                contact_name[:50],
                contact_phone[:20],
                intro[:2000],
                detail_type[:40],
                _json_text(detail_payload),
                action_label[:30],
                1 if is_enabled else 0,
                max(0, min(sort_order, 1000)),
                now,
                service_id,
            ),
        )
        updated = conn.execute(
            """
            SELECT
                id, service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, detail_type, detail_payload,
                action_label, is_enabled, sort_order, updated_at, created_at
            FROM gov_services
            WHERE id = ?
            LIMIT 1
            """,
            (service_id,),
        ).fetchone()
        _audit_log_conn(
            conn,
            cadre_user["id"],
            "gov_service_update",
            "gov_service",
            str(service_id),
            {"title": title, "is_enabled": is_enabled, "detail_type": detail_type},
        )
        conn.commit()
        return _gov_service_to_dict(updated, include_payload=True)
    finally:
        conn.close()


def create_gov_application(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    service_id = int(payload.get("service_id") or 0)
    applicant_name = str(payload.get("applicant_name") or "").strip()
    applicant_phone = normalize_phone_number(payload.get("applicant_phone") or "")
    selected_option = str(payload.get("selected_option") or "").strip()
    note = str(payload.get("note") or "").strip()
    extra_fields = _normalize_json_object(payload.get("extra_fields"))
    supporting_photo_urls = _normalize_photo_urls(
        payload.get(GOV_APPLICATION_SUPPORTING_PHOTO_FIELD),
        allowed_scenes=SCENES_GOV_APPLY,
        allow_external=False,
        strict_scene=True,
    )
    if not supporting_photo_urls:
        supporting_photo_urls = _normalize_photo_urls(
            extra_fields.get(GOV_APPLICATION_SUPPORTING_PHOTO_FIELD),
            allowed_scenes=SCENES_GOV_APPLY,
            allow_external=False,
            strict_scene=True,
        )
    if supporting_photo_urls:
        extra_fields[GOV_APPLICATION_SUPPORTING_PHOTO_FIELD] = supporting_photo_urls
    else:
        extra_fields.pop(GOV_APPLICATION_SUPPORTING_PHOTO_FIELD, None)
    for field_key in list(extra_fields.keys()):
        if _is_gender_field_name(field_key):
            extra_fields[field_key] = normalize_gender(extra_fields.get(field_key))

    if service_id <= 0:
        raise ValueError("invalid_gov_application_payload")

    conn = get_connection()
    try:
        user = _require_verified_villager_user_conn(conn, token)
        _validate_private_upload_ownership_conn(conn, supporting_photo_urls, int(user["id"] or 0))
        if not applicant_name:
            applicant_name = str(user["real_name"] or user["nickname"] or "").strip()
        if not applicant_phone:
            applicant_phone = _resolve_user_phone_number(user)
        if not applicant_name:
            raise ValueError("invalid_gov_application_payload")
        if applicant_phone and not is_valid_phone(applicant_phone):
            raise ValueError("invalid_phone")
        _seed_default_gov_services_conn(conn)
        service = conn.execute(
            """
            SELECT id, service_key, title, is_enabled
            FROM gov_services
            WHERE id = ?
            LIMIT 1
            """,
            (service_id,),
        ).fetchone()
        if service is None or int(service["is_enabled"] or 0) != 1:
            raise ValueError("gov_service_not_found")

        if str(service["service_key"] or "") == "certificate":
            if not str(extra_fields.get("name") or "").strip():
                extra_fields["name"] = applicant_name
            if normalize_gender(extra_fields.get("gender") or "") not in {"男", "女"}:
                encrypted_id_card = str(user["id_card_encrypted"] or "").strip()
                id_card_for_gender = ""
                if encrypted_id_card:
                    try:
                        id_card_for_gender = decrypt_id_card_number(encrypted_id_card)
                    except ValueError:
                        id_card_for_gender = ""
                extra_fields["gender"] = _resolve_gender(user["gender"] or "", id_card_for_gender) or "未知"
            if not str(extra_fields.get("id_card") or "").strip():
                encrypted_id_card = str(user["id_card_encrypted"] or "").strip()
                if encrypted_id_card:
                    extra_fields["id_card"] = decrypt_id_card_number(encrypted_id_card)
            if _has_duplicate_certificate_application_for_user_conn(
                conn,
                user=user,
                service_id=int(service["id"] or 0),
                selected_option=selected_option,
                extra_fields=extra_fields,
            ):
                raise ValueError("certificate_duplicate_submission")
            if _count_today_certificate_applications_for_user_conn(conn, user) >= CERTIFICATE_DAILY_APPLICATION_LIMIT:
                raise ValueError("certificate_daily_limit_reached")

        now = now_iso()
        duplicate_since = (datetime.now(timezone.utc) - timedelta(minutes=2)).isoformat()
        application_fingerprint_hash = ""
        if str(service["service_key"] or "") == "certificate":
            application_fingerprint_hash = keyed_lookup_hash(
                _build_certificate_application_fingerprint(
                    service_id=int(service["id"] or 0),
                    selected_option=selected_option,
                    extra_fields=extra_fields,
                ),
                purpose="gov_application_fingerprint",
            )
        stored_extra_fields, encrypted_extra_fields, applicant_id_card_hash = (
            _protect_gov_application_extra_fields(extra_fields)
        )
        serialized_extra_fields = _json_text(stored_extra_fields)
        applicant_phone_fields = _protected_business_phone_fields(
            applicant_phone,
            purpose="gov_application_phone",
        )
        duplicate_row = conn.execute(
            """
            SELECT id
            FROM gov_service_applications
            WHERE user_id = ?
              AND service_id = ?
              AND selected_option = ?
              AND note = ?
              AND extra_fields = ?
              AND created_at >= ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (
                user["id"],
                service["id"],
                selected_option,
                note,
                serialized_extra_fields,
                duplicate_since,
            ),
        ).fetchone()
        if duplicate_row is not None:
            raise ValueError("gov_application_duplicate_submission")
        cursor = conn.execute(
            """
            INSERT INTO gov_service_applications(
                service_id, user_id, service_key, service_title, applicant_name, applicant_phone,
                applicant_phone_ciphertext, applicant_phone_hash, applicant_phone_masked,
                selected_option, note, extra_fields, reply_photo_urls, status, cadre_reply, handled_by,
                sensitive_extra_fields_encrypted, applicant_id_card_hash, application_fingerprint_hash,
                created_at, updated_at, handled_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '[]', 'pending', '', '', ?, ?, ?, ?, ?, '')
            """,
            (
                service["id"],
                user["id"],
                service["service_key"],
                service["title"],
                applicant_name[:50],
                applicant_phone_fields["masked"][:20],
                applicant_phone_fields["ciphertext"],
                applicant_phone_fields["lookup_hash"],
                applicant_phone_fields["masked"][:20],
                selected_option[:120],
                note[:2000],
                serialized_extra_fields,
                encrypted_extra_fields,
                applicant_id_card_hash,
                application_fingerprint_hash,
                now,
                now,
            ),
        )
        application_id = cursor.lastrowid
        row = conn.execute(
            """
            SELECT
                id, service_id, service_key, service_title, user_id,
                applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                created_at, updated_at, handled_at
            FROM gov_service_applications
            WHERE id = ?
            LIMIT 1
            """,
            (application_id,),
        ).fetchone()
        _audit_log_conn(
            conn,
            user["id"],
            "gov_application_create",
            "gov_application",
            str(application_id),
            {
                "service_id": service_id,
                "service_key": service["service_key"],
                "selected_option": selected_option[:120],
                "supporting_photo_count": len(supporting_photo_urls),
            },
        )
        conn.commit()
        return _gov_application_to_dict(row)
    finally:
        conn.close()


def list_gov_applications(token: str, cadre_token: str | None = None) -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        access = _get_cadre_authorization_context_conn(conn, user)
        if access is not None and str(cadre_token or "").strip():
            _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
            rows = conn.execute(
                """
                SELECT
                    id, service_id, service_key, service_title, user_id,
                    applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                    extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                    created_at, updated_at, handled_at
                FROM gov_service_applications
                ORDER BY id DESC
                """
            ).fetchall()
            _audit_log_conn(
                conn,
                int(user["id"] or 0),
                "gov_application_list",
                "gov_application",
                "all",
                {"count": len(rows), "scope": "cadre"},
            )
            conn.commit()
        else:
            identity_claim = _build_verified_identity_claim(user)
            if identity_claim:
                rows = conn.execute(
                    """
                    SELECT
                        id, service_id, service_key, service_title, user_id,
                        applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                        extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                        created_at, updated_at, handled_at
                    FROM gov_service_applications
                    WHERE user_id = ? OR applicant_name = ?
                    ORDER BY id DESC
                    """,
                    (user["id"], identity_claim["real_name"]),
                ).fetchall()
                rows = [
                    row for row in rows if _gov_application_visible_to_user(row, user, identity_claim)
                ]
            else:
                rows = conn.execute(
                    """
                    SELECT
                        id, service_id, service_key, service_title, user_id,
                        applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                        extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                        created_at, updated_at, handled_at
                    FROM gov_service_applications
                    WHERE user_id = ?
                    ORDER BY id DESC
                    """,
                    (user["id"],),
                ).fetchall()
        return [_gov_application_to_dict(row) for row in rows]
    finally:
        conn.close()


def get_gov_application_detail(
    token: str,
    application_id: int,
    cadre_token: str | None = None,
) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        access = _get_cadre_authorization_context_conn(conn, user)
        query = """
            SELECT
                id, service_id, service_key, service_title, user_id,
                applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                created_at, updated_at, handled_at
            FROM gov_service_applications
            WHERE id = ?
        """
        params: list[Any] = [int(application_id or 0)]
        detail_scope = "self"
        if access is not None and str(cadre_token or "").strip():
            _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
            detail_scope = "cadre"
        else:
            detail_scope = "self"
        query += " LIMIT 1"
        row = conn.execute(query, tuple(params)).fetchone()
        if row is None or (
            detail_scope != "cadre"
            and not _gov_application_visible_to_user(row, user)
        ):
            raise ValueError("gov_application_not_found")
        if detail_scope == "cadre":
            _audit_log_conn(
                conn,
                int(user["id"] or 0),
                "gov_application_detail",
                "gov_application",
                str(application_id),
                {"scope": detail_scope},
            )
            conn.commit()
        return _gov_application_to_dict(row)
    finally:
        conn.close()


def update_gov_application(
    token: str,
    application_id: int,
    payload: dict[str, Any],
) -> dict[str, Any]:
    status = str(payload.get("status") or "").strip().lower()
    cadre_reply = str(payload.get("cadre_reply") or "").strip()
    reply_photo_urls = _normalize_photo_urls(
        payload.get("reply_photo_urls"),
        allowed_scenes=SCENES_GOV_REPLY,
        allow_external=False,
        strict_scene=True,
    )
    if status not in {"pending", "processing", "resolved", "rejected", "closed"}:
        raise ValueError("invalid_gov_application_status")

    conn = get_connection()
    try:
        cadre_user = _require_cadre_user_conn(conn, token)
        row = conn.execute(
            """
            SELECT id, handled_at, service_key
            FROM gov_service_applications
            WHERE id = ?
            LIMIT 1
            """,
            (application_id,),
        ).fetchone()
        if row is None:
            raise ValueError("gov_application_not_found")
        if status in {"processing", "resolved", "rejected", "closed"} and not cadre_reply and not reply_photo_urls:
            raise ValueError("gov_handle_requires_feedback")
        if str(row["service_key"] or "") == "certificate" and status in {"resolved", "closed"} and not reply_photo_urls:
            raise ValueError("certificate_reply_photo_required")

        now = now_iso()
        handled_by = cadre_user["real_name"] or cadre_user["nickname"] or "村委管理人"
        handled_at = (
            now
            if status in {"processing", "resolved", "rejected", "closed"} or bool(cadre_reply) or bool(reply_photo_urls)
            else (row["handled_at"] or "")
        )
        conn.execute(
            """
            UPDATE gov_service_applications
            SET
                status = ?,
                cadre_reply = ?,
                reply_photo_urls = ?,
                handled_by = ?,
                updated_at = ?,
                handled_at = ?
            WHERE id = ?
            """,
            (
                status,
                cadre_reply[:2000],
                _json_text(reply_photo_urls),
                handled_by[:50],
                now,
                handled_at,
                application_id,
            ),
        )
        updated = conn.execute(
            """
            SELECT
                id, service_id, service_key, service_title, user_id,
                applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                created_at, updated_at, handled_at
            FROM gov_service_applications
            WHERE id = ?
            LIMIT 1
            """,
            (application_id,),
        ).fetchone()
        _audit_log_conn(
            conn,
            cadre_user["id"],
            "gov_application_update",
            "gov_application",
            str(application_id),
            {
                "status": status,
                "reply_length": len(cadre_reply),
                "reply_photo_count": len(reply_photo_urls),
            },
        )
        conn.commit()
        return _gov_application_to_dict(updated)
    finally:
        conn.close()


def cancel_gov_application(token: str, application_id: int) -> dict[str, str]:
    conn = get_connection()
    try:
        user = _require_verified_villager_user_conn(conn, token)
        row = conn.execute(
            """
            SELECT
                id, service_id, service_key, service_title, user_id,
                applicant_name, applicant_phone, applicant_phone_masked, selected_option, note,
                extra_fields, applicant_id_card_hash, reply_photo_urls, status, cadre_reply, handled_by,
                created_at, updated_at, handled_at
            FROM gov_service_applications
            WHERE id = ?
            LIMIT 1
            """,
            (application_id,),
        ).fetchone()
        if row is None or not _gov_application_visible_to_user(row, user):
            raise ValueError("gov_application_not_found")
        status = str(row["status"] or "").strip().lower()
        if status not in {"pending", "processing"}:
            raise ValueError("gov_application_cancel_not_allowed")

        now = now_iso()
        result_text = "村民已主动取消本次申请。"
        conn.execute(
            """
            UPDATE gov_service_applications
            SET
                status = 'closed',
                cadre_reply = ?,
                handled_by = ?,
                updated_at = ?,
                handled_at = ?
            WHERE id = ?
            """,
            (result_text, "村民本人", now, now, application_id),
        )
        _audit_log_conn(
            conn,
            user["id"],
            "gov_application_cancel",
            "gov_application",
            str(application_id),
            {
                "service_key": str(row["service_key"] or ""),
                "service_title": str(row["selected_option"] or row["service_title"] or ""),
            },
        )
        conn.commit()
        return {
            "status": "ok",
            "message": "申请已取消",
        }
    finally:
        conn.close()


def _sanitize_metric_key(metric_key: str) -> str:
    text = re.sub(r"[^a-zA-Z0-9_]+", "_", (metric_key or "").strip().lower())
    text = text.strip("_")
    return text[:64]


def _seed_default_metrics_conn(conn: sqlite3.Connection) -> None:
    row = conn.execute("SELECT COUNT(1) AS cnt FROM village_metrics").fetchone()
    if row and row["cnt"] > 0:
        return
    now = now_iso()
    conn.executemany(
        """
        INSERT INTO village_metrics(
            metric_key, metric_label, metric_value, sort_order, updated_by, created_at, updated_at
        )
        VALUES(?, ?, ?, ?, '系统预置', ?, ?)
        """,
        [
            (
                item["metric_key"],
                item["metric_label"],
                item["metric_value"],
                index,
                now,
                now,
            )
            for index, item in enumerate(DEFAULT_VILLAGE_METRICS, start=1)
        ],
    )


def list_village_metrics() -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        _seed_default_metrics_conn(conn)
        rows = conn.execute(
            """
            SELECT metric_key, metric_label, metric_value, sort_order, updated_at
            FROM village_metrics
            ORDER BY sort_order ASC, id ASC
            """
        ).fetchall()
        conn.commit()
        items = []
        for row in rows:
            item = dict(row)
            if item.get("metric_key") == "groups":
                item["metric_label"] = "村组"
            items.append(item)
        return items
    finally:
        conn.close()


def _normalize_metric_number(value: Any, fallback: str = "") -> str:
    text = str(value or "").strip()
    match = re.search(r"\d+", text)
    if match:
        return match.group(0)
    return fallback


def _metric_value_for_profile(rows: list[dict[str, Any]], metric_key: str, label: str, fallback: str) -> str:
    for row in rows:
        if str(row.get("metric_key") or "").strip() == metric_key or str(row.get("metric_label") or "").strip() == label:
            return _normalize_metric_number(row.get("metric_value"), fallback)
    return fallback


def get_village_open_profile() -> dict[str, Any]:
    conn = get_connection()
    try:
        _seed_default_metrics_conn(conn)
        metric_rows = conn.execute(
            """
            SELECT metric_key, metric_label, metric_value, sort_order, updated_at
            FROM village_metrics
            ORDER BY sort_order ASC, id ASC
            """
        ).fetchall()
        metrics = [dict(row) for row in metric_rows]
        profile_row = conn.execute(
            """
            SELECT village_intro, village_photo_url
            FROM open_profile
            WHERE id = 1
            LIMIT 1
            """
        ).fetchone()
        village_intro = ""
        village_photo_url = ""
        if profile_row is not None:
            village_intro = str(profile_row["village_intro"] or "").strip()
            raw_photo_url = str(profile_row["village_photo_url"] or "").strip()
            if raw_photo_url:
                village_photo_url = _present_photo_url(raw_photo_url)
        if not village_intro:
            village_intro = DEFAULT_VILLAGE_INTRO
        conn.commit()
        return {
            "village_groups": _metric_value_for_profile(metrics, "groups", "村组", "9"),
            "households": _metric_value_for_profile(metrics, "families", "家庭", "538"),
            "population": _metric_value_for_profile(metrics, "population", "人口", "1744"),
            "party_members": _metric_value_for_profile(metrics, "members", "党员", "42"),
            "village_intro": village_intro,
            "village_photo_url": village_photo_url,
        }
    finally:
        conn.close()


def _normalize_open_profile_payload(payload: dict[str, Any]) -> dict[str, str]:
    village_intro = str(payload.get("village_intro") or "").strip()[:2000]
    village_photo_url = str(payload.get("village_photo_url") or "").strip()
    normalized_photo_url = ""
    if village_photo_url:
        urls = _normalize_photo_urls(
            [village_photo_url],
            allowed_scenes=SCENES_OPEN,
            allow_external=False,
            strict_scene=True,
        )
        normalized_photo_url = urls[0] if urls else ""
        if not normalized_photo_url:
            raise ValueError("invalid_photo_scene")
    return {
        "village_intro": village_intro,
        "village_photo_url": normalized_photo_url,
    }


def update_village_open_profile(updated_by: str, payload: dict[str, Any]) -> dict[str, Any]:
    normalized = _normalize_open_profile_payload(payload)
    conn = get_connection()
    try:
        now = now_iso()
        row = conn.execute("SELECT id FROM open_profile WHERE id = 1 LIMIT 1").fetchone()
        if row is None:
            conn.execute(
                """
                INSERT INTO open_profile(id, village_intro, village_photo_url, updated_by, created_at, updated_at)
                VALUES(1, ?, ?, ?, ?, ?)
                """,
                (
                    normalized["village_intro"],
                    normalized["village_photo_url"],
                    str(updated_by or "").strip()[:120],
                    now,
                    now,
                ),
            )
        else:
            conn.execute(
                """
                UPDATE open_profile
                SET village_intro = ?, village_photo_url = ?, updated_by = ?, updated_at = ?
                WHERE id = 1
                """,
                (
                    normalized["village_intro"],
                    normalized["village_photo_url"],
                    str(updated_by or "").strip()[:120],
                    now,
                ),
            )
        conn.commit()
    finally:
        conn.close()
    return get_village_open_profile()


def update_village_metrics(token: str, payload: dict[str, Any]) -> list[dict[str, Any]]:
    items = payload.get("items") or []
    if not isinstance(items, list) or not items:
        raise ValueError("invalid_metric_payload")

    conn = get_connection()
    try:
        user = _require_cadre_user_conn(conn, token)
        now = now_iso()
        updated_by = user["real_name"] or user["nickname"] or user["openid"]

        normalized: list[dict[str, Any]] = []
        used_keys: set[str] = set()
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                raise ValueError("invalid_metric_payload")
            metric_label = str(item.get("metric_label") or "").strip()
            metric_value = str(item.get("metric_value") or "").strip()
            metric_key = _sanitize_metric_key(str(item.get("metric_key") or ""))
            if not metric_key:
                metric_key = f"metric_{index}"
            while metric_key in used_keys:
                metric_key = f"{metric_key}_{index}"
            if not metric_label or not metric_value:
                raise ValueError("invalid_metric_payload")
            used_keys.add(metric_key)
            normalized.append(
                {
                    "metric_key": metric_key,
                    "metric_label": metric_label[:40],
                    "metric_value": metric_value[:40],
                    "sort_order": index,
                }
            )

        conn.execute("DELETE FROM village_metrics")
        conn.executemany(
            """
            INSERT INTO village_metrics(
                metric_key, metric_label, metric_value, sort_order, updated_by, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    item["metric_key"],
                    item["metric_label"],
                    item["metric_value"],
                    item["sort_order"],
                    updated_by,
                    now,
                    now,
                )
                for item in normalized
            ],
        )
        _audit_log_conn(
            conn,
            user["id"],
            "metrics_update",
            "village_metrics",
            "all",
            {
                "count": len(normalized),
                "updated_by": updated_by,
            },
        )
        rows = conn.execute(
            """
            SELECT metric_key, metric_label, metric_value, sort_order, updated_at
            FROM village_metrics
            ORDER BY sort_order ASC, id ASC
            """
        ).fetchall()
        conn.commit()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def _villager_request_to_dict(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item.pop("reporter_id_card_hash", None)
    item.pop("reporter_real_name", None)
    item["contact_phone"] = str(item.get("contact_phone_masked") or "").strip() or str(
        _mask_sensitive_value("contact_phone", item.get("contact_phone", "")) or ""
    )
    for field_name in ("contact_phone_ciphertext", "contact_phone_hash", "contact_phone_masked"):
        item.pop(field_name, None)
    item["photo_urls"] = _parse_photo_urls_field(
        item.get("photo_urls"),
        allowed_scenes=SCENES_REQUEST,
        allow_external=False,
    )
    item["reply_photo_urls"] = _parse_photo_urls_field(
        item.get("reply_photo_urls"),
        allowed_scenes=SCENES_REQUEST_REPLY,
        allow_external=False,
    )
    return item


def _villager_request_visible_to_user(
    row: sqlite3.Row | dict[str, Any],
    user: sqlite3.Row | dict[str, Any],
    identity_claim: dict[str, str] | None = None,
) -> bool:
    item = dict(row)
    data = dict(user)
    if int(item.get("user_id") or 0) == int(data.get("id") or 0):
        return True
    claim = identity_claim or _build_verified_identity_claim(data)
    if not claim:
        return False
    reporter_id_hash = str(item.get("reporter_id_card_hash") or "").strip()
    reporter_real_name = str(item.get("reporter_real_name") or item.get("reporter_name") or "").strip()
    return (
        bool(reporter_id_hash)
        and reporter_id_hash == claim["id_hash"]
        and reporter_real_name == claim["real_name"]
    )


def create_villager_request(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    request_type = str(payload.get("request_type") or "service").strip().lower()
    title = str(payload.get("title") or "").strip()
    content = str(payload.get("content") or "").strip()
    contact_phone = normalize_phone_number(payload.get("contact_phone") or "")
    photo_urls = _normalize_photo_urls(
        payload.get("photo_urls"),
        allowed_scenes=SCENES_REQUEST,
        allow_external=False,
        strict_scene=True,
    )
    if request_type not in {"complaint", "service", "dibao", "tekun"}:
        raise ValueError("invalid_request_type")
    if not title or not content:
        raise ValueError("invalid_request_payload")
    conn = get_connection()
    try:
        user = _require_verified_villager_user_conn(conn, token)
        _validate_private_upload_ownership_conn(conn, photo_urls, int(user["id"] or 0))
        if not contact_phone:
            contact_phone = _resolve_user_phone_number(user)
        if not contact_phone:
            raise ValueError("request_phone_required")
        if not is_valid_phone(contact_phone):
            raise ValueError("invalid_phone")
        reporter_real_name = str(user["real_name"] or "").strip()
        reporter_id_card_hash = str(user["id_card_hash"] or "").strip()
        reporter_name = reporter_real_name or user["nickname"] or user["openid"]
        protected_phone = _protected_business_phone_fields(
            contact_phone,
            purpose="villager_request_phone",
        )
        now = now_iso()
        duplicate_since = (datetime.now(timezone.utc) - timedelta(minutes=2)).isoformat()
        duplicate_row = conn.execute(
            """
            SELECT id
            FROM villager_requests
            WHERE user_id = ?
              AND request_type = ?
              AND title = ?
              AND content = ?
              AND contact_phone_hash = ?
              AND created_at >= ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (
                user["id"],
                request_type,
                title,
                content,
                protected_phone["lookup_hash"],
                duplicate_since,
            ),
        ).fetchone()
        if duplicate_row is not None:
            raise ValueError("request_duplicate_submission")
        cursor = conn.execute(
            """
            INSERT INTO villager_requests(
                user_id, request_type, title, content, contact_phone,
                contact_phone_ciphertext, contact_phone_hash, contact_phone_masked,
                reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '[]', 'pending', '', '', ?, ?, '')
            """,
            (
                user["id"],
                request_type,
                title,
                content,
                protected_phone["masked"],
                protected_phone["ciphertext"],
                protected_phone["lookup_hash"],
                protected_phone["masked"],
                reporter_name,
                reporter_real_name,
                reporter_id_card_hash,
                _json_text(photo_urls),
                now,
                now,
            ),
        )
        request_id = cursor.lastrowid
        row = conn.execute(
            """
            SELECT
                id, user_id, request_type, title, content, contact_phone, contact_phone_masked,
                reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
            FROM villager_requests
            WHERE id = ?
            """,
            (request_id,),
        ).fetchone()
        _audit_log_conn(
            conn,
            user["id"],
            "villager_request_create",
            "villager_request",
            str(request_id),
            {"request_type": request_type},
        )
        conn.commit()
        return _villager_request_to_dict(row)
    finally:
        conn.close()


def list_villager_requests(token: str, cadre_token: str | None = None) -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        access = _get_cadre_authorization_context_conn(conn, user)
        if access is not None and str(cadre_token or "").strip():
            _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
            rows = conn.execute(
                """
                SELECT
                    id, user_id, request_type, title, content, contact_phone,
                    reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                    reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
                FROM villager_requests
                ORDER BY id DESC
                """
            ).fetchall()
            _audit_log_conn(
                conn,
                int(user["id"] or 0),
                "villager_request_list",
                "villager_request",
                "all",
                {"count": len(rows), "scope": "cadre"},
            )
            conn.commit()
        else:
            identity_claim = _build_verified_identity_claim(user)
            if identity_claim:
                rows = conn.execute(
                    """
                    SELECT
                        id, user_id, request_type, title, content, contact_phone,
                        reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                        reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
                    FROM villager_requests
                    WHERE user_id = ? OR reporter_id_card_hash = ?
                    ORDER BY id DESC
                    """,
                    (user["id"], identity_claim["id_hash"]),
                ).fetchall()
                rows = [
                    row for row in rows if _villager_request_visible_to_user(row, user, identity_claim)
                ]
            else:
                rows = conn.execute(
                    """
                    SELECT
                        id, user_id, request_type, title, content, contact_phone,
                        reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                        reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
                    FROM villager_requests
                    WHERE user_id = ?
                    ORDER BY id DESC
                    """,
                    (user["id"],),
                ).fetchall()
        return [_villager_request_to_dict(row) for row in rows]
    finally:
        conn.close()


def get_villager_request_detail(
    token: str,
    request_id: int,
    cadre_token: str | None = None,
) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        access = _get_cadre_authorization_context_conn(conn, user)
        query = """
            SELECT
                id, user_id, request_type, title, content, contact_phone,
                reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
            FROM villager_requests
            WHERE id = ?
        """
        params: list[Any] = [int(request_id or 0)]
        detail_scope = "self"
        if access is not None and str(cadre_token or "").strip():
            _verify_cadre_stepup_session_conn(conn, int(user["id"] or 0), cadre_token)
            detail_scope = "cadre"
        query += " LIMIT 1"
        row = conn.execute(query, tuple(params)).fetchone()
        if row is None or (
            detail_scope != "cadre"
            and not _villager_request_visible_to_user(row, user)
        ):
            raise ValueError("request_not_found")
        if detail_scope == "cadre":
            _audit_log_conn(
                conn,
                int(user["id"] or 0),
                "villager_request_detail",
                "villager_request",
                str(request_id),
                {"scope": detail_scope},
            )
            conn.commit()
        return _villager_request_to_dict(row)
    finally:
        conn.close()


def update_villager_request(
    token: str,
    request_id: int,
    payload: dict[str, Any],
) -> dict[str, Any]:
    status = str(payload.get("status") or "").strip().lower()
    reply_text = str(payload.get("reply_text") or "").strip()
    reply_photo_urls = _normalize_photo_urls(
        payload.get("reply_photo_urls"),
        allowed_scenes=SCENES_REQUEST_REPLY,
        allow_external=False,
        strict_scene=True,
    )
    if status not in {"pending", "processing", "resolved", "closed"}:
        raise ValueError("invalid_request_status")
    if status in {"processing", "resolved", "closed"} and not reply_text and not reply_photo_urls:
        raise ValueError("request_handle_requires_feedback")

    conn = get_connection()
    try:
        user = _require_cadre_user_conn(conn, token)
        row = conn.execute(
            "SELECT id, handled_at FROM villager_requests WHERE id = ? LIMIT 1",
            (request_id,),
        ).fetchone()
        if row is None:
            raise ValueError("request_not_found")

        now = now_iso()
        handled_by = user["real_name"] or user["nickname"] or user["openid"]
        handled_at = (
            now
            if status in {"processing", "resolved", "closed"} or bool(reply_text) or bool(reply_photo_urls)
            else (row["handled_at"] or "")
        )
        conn.execute(
            """
            UPDATE villager_requests
            SET status = ?, reply_text = ?, reply_photo_urls = ?, handled_by = ?, updated_at = ?, handled_at = ?
            WHERE id = ?
            """,
            (status, reply_text, _json_text(reply_photo_urls), handled_by, now, handled_at, request_id),
        )
        updated = conn.execute(
            """
            SELECT
                id, user_id, request_type, title, content, contact_phone,
                reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
            FROM villager_requests
            WHERE id = ?
            LIMIT 1
            """,
            (request_id,),
        ).fetchone()
        _audit_log_conn(
            conn,
            user["id"],
            "villager_request_update",
            "villager_request",
            str(request_id),
            {
                "status": status,
                "reply_length": len(reply_text),
                "reply_photo_count": len(reply_photo_urls),
            },
        )
        conn.commit()
        return _villager_request_to_dict(updated)
    finally:
        conn.close()


def cancel_villager_request(token: str, request_id: int) -> dict[str, str]:
    conn = get_connection()
    try:
        user = _require_verified_villager_user_conn(conn, token)
        row = conn.execute(
            """
            SELECT
                id, user_id, request_type, title, content, contact_phone,
                reporter_name, reporter_real_name, reporter_id_card_hash, photo_urls,
                reply_photo_urls, status, reply_text, handled_by, created_at, updated_at, handled_at
            FROM villager_requests
            WHERE id = ?
            LIMIT 1
            """,
            (request_id,),
        ).fetchone()
        if row is None or not _villager_request_visible_to_user(row, user):
            raise ValueError("request_not_found")
        status = str(row["status"] or "").strip().lower()
        if status not in {"pending", "processing"}:
            raise ValueError("request_cancel_not_allowed")

        now = now_iso()
        result_text = "村民已主动取消本次诉求。"
        conn.execute(
            """
            UPDATE villager_requests
            SET
                status = 'closed',
                reply_text = ?,
                handled_by = ?,
                updated_at = ?,
                handled_at = ?
            WHERE id = ?
            """,
            (result_text, "村民本人", now, now, request_id),
        )
        _audit_log_conn(
            conn,
            user["id"],
            "villager_request_cancel",
            "villager_request",
            str(request_id),
            {
                "request_type": str(row["request_type"] or ""),
                "title": str(row["title"] or ""),
            },
        )
        conn.commit()
        return {
            "status": "ok",
            "message": "诉求已取消",
        }
    finally:
        conn.close()


def _notice_to_dict(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["photo_urls"] = _parse_photo_urls_field(
        item.get("photo_urls"),
        allowed_scenes=SCENES_NOTICE,
        allow_external=False,
    )
    item["is_home_top"] = bool(int(item.get("is_home_top") or 0))
    return item


def get_notice_subscription_status(token: str) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        config = _notice_subscription_config()
        row = conn.execute(
            """
            SELECT template_id, subscription_status, is_enabled, updated_at
            FROM notice_subscriptions
            WHERE user_id = ?
            LIMIT 1
            """,
            (int(user["id"] or 0),),
        ).fetchone()
        current_status = "unknown"
        is_enabled = False
        updated_at = ""
        if row is not None and str(row["template_id"] or "").strip() == str(config["template_id"] or "").strip():
            current_status = _normalize_notice_subscription_status(row["subscription_status"])
            is_enabled = bool(int(row["is_enabled"] or 0)) and _notice_subscription_is_enabled(current_status)
            updated_at = str(row["updated_at"] or "").strip()
        if not config["available"]:
            button_text = "通知模板配置中"
        elif is_enabled:
            button_text = "已开启通知提醒"
        elif current_status in {"reject", "delivery_rejected"}:
            button_text = "重新开启通知提醒"
        else:
            button_text = "订阅服务消息"
        return {
            "available": bool(config["available"]),
            "is_enabled": is_enabled,
            "template_id": str(config["template_id"] or ""),
            "template_title": str(config["template_title"] or ""),
            "current_status": current_status,
            "button_text": button_text,
            "prompt_title": str(config["prompt_title"] or ""),
            "prompt_content": str(config["prompt_content"] or ""),
            "updated_at": updated_at,
        }
    finally:
        conn.close()


def record_notice_subscription_result(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    config = _notice_subscription_config()
    if not config["available"]:
        raise ValueError("wechat_notice_template_unconfigured")
    conn = get_connection()
    try:
        user = _get_user_by_token_conn(conn, token)
        template_statuses = payload.get("template_statuses") if isinstance(payload.get("template_statuses"), dict) else {}
        scene = _clean_single_line_text(payload.get("scene") or "manual")[:40] or "manual"
        template_id = str(config["template_id"] or "").strip()
        status_text = _normalize_notice_subscription_status(template_statuses.get(template_id))
        if status_text == "unknown" and template_statuses:
            status_text = _normalize_notice_subscription_status(next(iter(template_statuses.values())))
        is_enabled = _notice_subscription_is_enabled(status_text)
        row = _upsert_notice_subscription_conn(
            conn,
            user_id=int(user["id"] or 0),
            openid=_clean_single_line_text(user["openid"]),
            template_id=template_id,
            subscription_status=status_text,
            is_enabled=is_enabled,
            scene=scene,
            result_json=template_statuses,
            last_error_code="",
            last_error_message=_clean_single_line_text(payload.get("err_msg"))[:200],
        )
        _audit_log_conn(
            conn,
            int(user["id"] or 0),
            "notice_subscription_update",
            "notice_subscription",
            template_id,
            {
                "scene": scene,
                "status": status_text,
                "is_enabled": is_enabled,
            },
        )
        conn.commit()
        return {
            "available": bool(config["available"]),
            "is_enabled": bool(int(row["is_enabled"] or 0)),
            "template_id": str(row["template_id"] or ""),
            "template_title": str(config["template_title"] or ""),
            "current_status": _normalize_notice_subscription_status(row["subscription_status"]),
            "button_text": "已开启通知提醒" if bool(int(row["is_enabled"] or 0)) else "订阅服务消息",
            "prompt_title": str(config["prompt_title"] or ""),
            "prompt_content": str(config["prompt_content"] or ""),
            "updated_at": str(row["updated_at"] or ""),
        }
    finally:
        conn.close()


def _store_notice_delivery_result_conn(
    conn: sqlite3.Connection,
    *,
    notice_id: int,
    user_id: int,
    openid: str,
    template_id: str,
    delivery_status: str,
    errcode: int,
    errmsg: str,
    response_json: dict[str, Any],
    sent_at: str = "",
) -> None:
    now = now_iso()
    existing = conn.execute(
        """
        SELECT id
        FROM notice_delivery_logs
        WHERE notice_id = ? AND user_id = ? AND template_id = ?
        LIMIT 1
        """,
        (int(notice_id or 0), int(user_id or 0), template_id),
    ).fetchone()
    if existing is None:
        conn.execute(
            """
            INSERT INTO notice_delivery_logs(
                notice_id, user_id, openid, template_id, delivery_status,
                errcode, errmsg, response_json, sent_at, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(notice_id or 0),
                int(user_id or 0),
                openid,
                template_id,
                delivery_status,
                int(errcode or 0),
                errmsg[:255],
                _json_text(response_json),
                sent_at,
                now,
                now,
            ),
        )
        return
    conn.execute(
        """
        UPDATE notice_delivery_logs
        SET
            openid = ?,
            delivery_status = ?,
            errcode = ?,
            errmsg = ?,
            response_json = ?,
            sent_at = ?,
            updated_at = ?
        WHERE id = ?
        """,
        (
            openid,
            delivery_status,
            int(errcode or 0),
            errmsg[:255],
            _json_text(response_json),
            sent_at,
            now,
            int(existing["id"] or 0),
        ),
    )


def dispatch_notice_subscription_messages(notice: dict[str, Any]) -> dict[str, Any]:
    config = _notice_subscription_config()
    template_id = str(config["template_id"] or "").strip()
    template_fields = config.get("template_fields") if isinstance(config.get("template_fields"), dict) else {}
    if not config["available"] or not template_id or not template_fields:
        return {
            "status": "skipped",
            "reason": "unconfigured",
            "attempted": 0,
            "sent": 0,
            "failed": 0,
            "disabled": 0,
        }
    notice_id = int(notice.get("id") or 0)
    if notice_id <= 0:
        return {
            "status": "skipped",
            "reason": "invalid_notice",
            "attempted": 0,
            "sent": 0,
            "failed": 0,
            "disabled": 0,
        }
    payload_data = _render_notice_template_data(notice, template_fields)
    if not payload_data:
        return {
            "status": "skipped",
            "reason": "invalid_template_fields",
            "attempted": 0,
            "sent": 0,
            "failed": 0,
            "disabled": 0,
        }

    conn = get_connection()
    try:
        recipients = conn.execute(
            """
            SELECT
                ns.user_id,
                ns.template_id,
                COALESCE(NULLIF(u.openid, ''), ns.openid) AS openid
            FROM notice_subscriptions ns
            JOIN users u ON u.id = ns.user_id
            WHERE ns.is_enabled = 1 AND ns.template_id = ?
            ORDER BY ns.updated_at DESC, ns.user_id DESC
            """,
            (template_id,),
        ).fetchall()
        attempted = 0
        sent = 0
        failed = 0
        disabled = 0
        for row in recipients:
            user_id = int(row["user_id"] or 0)
            openid = _clean_single_line_text(row["openid"])
            if user_id <= 0 or not openid:
                continue
            attempted += 1
            request_payload = {
                "touser": openid,
                "template_id": template_id,
                "page": _render_notice_subscription_page(notice_id),
                "miniprogram_state": str(config["miniprogram_state"] or "formal"),
                "lang": str(config["lang"] or "zh_CN"),
                "data": payload_data,
            }
            sent_at = ""
            try:
                response = _request_wechat_openapi_json(
                    url=WECHAT_SUBSCRIBE_MESSAGE_SEND_URL,
                    payload=request_payload,
                    timeout=8,
                    unconfigured_error="wechat_notice_template_unconfigured",
                    unavailable_error="wechat_notice_send_unavailable",
                )
            except ValueError as exc:
                failed += 1
                _store_notice_delivery_result_conn(
                    conn,
                    notice_id=notice_id,
                    user_id=user_id,
                    openid=openid,
                    template_id=template_id,
                    delivery_status="failed",
                    errcode=-1,
                    errmsg=str(exc),
                    response_json={"detail": str(exc)},
                    sent_at="",
                )
                continue
            errcode = int(response.get("errcode") or 0)
            errmsg = _clean_single_line_text(response.get("errmsg"))
            delivery_status = "sent" if errcode == 0 else "failed"
            if errcode == 0:
                sent += 1
                sent_at = now_iso()
            else:
                failed += 1
                if errcode in {40003, 43101}:
                    disabled += 1
                    _upsert_notice_subscription_conn(
                        conn,
                        user_id=user_id,
                        openid=openid,
                        template_id=template_id,
                        subscription_status="delivery_rejected",
                        is_enabled=False,
                        scene="delivery_feedback",
                        result_json=response,
                        last_error_code=str(errcode),
                        last_error_message=errmsg,
                    )
            _store_notice_delivery_result_conn(
                conn,
                notice_id=notice_id,
                user_id=user_id,
                openid=openid,
                template_id=template_id,
                delivery_status=delivery_status,
                errcode=errcode,
                errmsg=errmsg,
                response_json=response,
                sent_at=sent_at,
            )
        conn.commit()
        return {
            "status": "ok",
            "reason": "",
            "attempted": attempted,
            "sent": sent,
            "failed": failed,
            "disabled": disabled,
        }
    finally:
        conn.close()


def safe_dispatch_notice_subscription_messages(notice: dict[str, Any]) -> dict[str, Any]:
    try:
        return dispatch_notice_subscription_messages(notice)
    except ValueError as exc:
        logger.warning("notice_subscription_dispatch_skipped: detail=%s", exc)
        return {
            "status": "skipped",
            "reason": str(exc),
            "attempted": 0,
            "sent": 0,
            "failed": 0,
            "disabled": 0,
        }
    except Exception:
        logger.exception("notice_subscription_dispatch_failed")
        return {
            "status": "failed",
            "reason": "request_failed",
            "attempted": 0,
            "sent": 0,
            "failed": 0,
            "disabled": 0,
        }


def list_notices() -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT id, title, content, publisher, photo_urls, is_home_top, created_at "
            "FROM notices ORDER BY created_at DESC, id DESC"
        ).fetchall()
        return [_notice_to_dict(row) for row in rows]
    finally:
        conn.close()


def create_notice(data: dict[str, Any]) -> dict[str, Any]:
    created_notice: dict[str, Any] | None = None
    conn = get_connection()
    try:
        created_at = now_iso()
        is_home_top = 1 if bool(data.get("is_home_top")) else 0
        photo_urls = _normalize_photo_urls(
            data.get("photo_urls"),
            allowed_scenes=SCENES_NOTICE,
            allow_external=False,
            strict_scene=True,
        )
        cursor = conn.execute(
            "INSERT INTO notices(title, content, publisher, photo_urls, is_home_top, created_at) "
            "VALUES(?, ?, ?, ?, ?, ?)",
            (
                data["title"],
                data["content"],
                data["publisher"],
                _json_text(photo_urls),
                is_home_top,
                created_at,
            ),
        )
        conn.commit()
        new_id = cursor.lastrowid
        row = conn.execute(
            "SELECT id, title, content, publisher, photo_urls, is_home_top, created_at "
            "FROM notices WHERE id = ?",
            (new_id,),
        ).fetchone()
        created_notice = _notice_to_dict(row)
    finally:
        conn.close()
    if created_notice is None:
        raise ValueError("request_failed")
    safe_dispatch_notice_subscription_messages(created_notice)
    return created_notice


def update_notice(token: str, notice_id: int, data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        cadre_user = _require_cadre_user_conn(conn, token)
        row = conn.execute(
            "SELECT id FROM notices WHERE id = ? LIMIT 1",
            (notice_id,),
        ).fetchone()
        if row is None:
            raise ValueError("notice_not_found")
        photo_urls = _normalize_photo_urls(
            data.get("photo_urls"),
            allowed_scenes=SCENES_NOTICE,
            allow_external=False,
            strict_scene=True,
        )
        is_home_top = 1 if bool(data.get("is_home_top")) else 0
        conn.execute(
            """
            UPDATE notices
            SET title = ?, content = ?, publisher = ?, photo_urls = ?, is_home_top = ?
            WHERE id = ?
            """,
            (
                data["title"],
                data["content"],
                data["publisher"],
                _json_text(photo_urls),
                is_home_top,
                notice_id,
            ),
        )
        _audit_log_conn(
            conn,
            cadre_user["id"],
            "notice_update",
            "notice",
            str(notice_id),
            {
                "title": str(data["title"] or ""),
                "is_home_top": bool(is_home_top),
            },
        )
        conn.commit()
        updated_row = conn.execute(
            "SELECT id, title, content, publisher, photo_urls, is_home_top, created_at "
            "FROM notices WHERE id = ?",
            (notice_id,),
        ).fetchone()
        return _notice_to_dict(updated_row)
    finally:
        conn.close()


def delete_notice(token: str, notice_id: int) -> dict[str, Any]:
    conn = get_connection()
    try:
        cadre_user = _require_cadre_user_conn(conn, token)
        row = conn.execute(
            "SELECT id, title FROM notices WHERE id = ? LIMIT 1",
            (notice_id,),
        ).fetchone()
        if row is None:
            raise ValueError("notice_not_found")
        conn.execute("DELETE FROM notices WHERE id = ?", (notice_id,))
        _audit_log_conn(
            conn,
            cadre_user["id"],
            "notice_delete",
            "notice",
            str(notice_id),
            {"title": str(row["title"] or "")},
        )
        conn.commit()
        return {"status": "deleted", "message": "通知已删除"}
    finally:
        conn.close()


def get_notice_detail(notice_id: int) -> dict[str, Any]:
    conn = get_connection()
    try:
        notice_row = conn.execute(
            "SELECT id, title, content, publisher, photo_urls, is_home_top, created_at "
            "FROM notices WHERE id = ? LIMIT 1",
            (notice_id,),
        ).fetchone()
        if notice_row is None:
            raise ValueError("notice_not_found")
        return _notice_to_dict(notice_row)
    finally:
        conn.close()


def _parse_training_date(value: Any) -> date | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _today_china_date() -> date:
    return datetime.now(CHINA_TIMEZONE).date()


def _normalize_training_status(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"ended", "closed", "inactive"}:
        return "ended"
    return "open"


def _normalize_training_program_payload(data: dict[str, Any]) -> dict[str, str]:
    title = _clean_single_line_text(data.get("title"))[:100]
    if not title:
        raise ValueError("invalid_training_program_title")
    signup_deadline = _clean_single_line_text(data.get("signup_deadline"))[:20]
    end_date = _clean_single_line_text(data.get("end_date"))[:20]
    deadline_date = _parse_training_date(signup_deadline)
    finish_date = _parse_training_date(end_date)
    if signup_deadline and deadline_date is None:
        raise ValueError("invalid_training_date")
    if end_date and finish_date is None:
        raise ValueError("invalid_training_date")
    if deadline_date is not None and finish_date is not None and finish_date < deadline_date:
        raise ValueError("invalid_training_date_range")
    return {
        "title": title,
        "schedule_text": _clean_single_line_text(data.get("schedule_text"))[:120],
        "location": _clean_single_line_text(data.get("location"))[:120],
        "target": _clean_single_line_text(data.get("target"))[:200],
        "signup_deadline": signup_deadline[:10] if deadline_date is not None else "",
        "end_date": end_date[:10] if finish_date is not None else "",
        "contact_phone": _clean_single_line_text(data.get("contact_phone"))[:20],
        "status": _normalize_training_status(data.get("status")),
    }


def _training_program_status(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    today = _today_china_date()
    manual_status = _normalize_training_status((dict(row)).get("status"))
    signup_deadline = _parse_training_date((dict(row)).get("signup_deadline"))
    end_date = _parse_training_date((dict(row)).get("end_date"))
    if manual_status == "ended" or (end_date is not None and today > end_date):
        return {
            "status": "ended",
            "status_label": "已结束",
            "is_open": False,
            "is_ended": True,
            "days_remaining": None,
            "deadline_text": "已结束",
        }
    if signup_deadline is not None:
        days_remaining = (signup_deadline - today).days
        if days_remaining >= 0:
            return {
                "status": "open",
                "status_label": "报名中",
                "is_open": True,
                "is_ended": False,
                "days_remaining": days_remaining,
                "deadline_text": "今天截止" if days_remaining == 0 else f"剩余{days_remaining}天",
            }
        return {
            "status": "cutoff",
            "status_label": "报名截止",
            "is_open": False,
            "is_ended": True,
            "days_remaining": None,
            "deadline_text": "报名已截止",
        }
    return {
        "status": "open",
        "status_label": "报名中",
        "is_open": True,
        "is_ended": False,
        "days_remaining": None,
        "deadline_text": "报名中",
    }


def _training_program_to_dict(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["id"] = int(item.get("id") or 0)
    item["title"] = _clean_single_line_text(item.get("title"))[:100]
    item["schedule_text"] = _clean_single_line_text(item.get("schedule_text"))[:120]
    item["location"] = _clean_single_line_text(item.get("location"))[:120]
    item["target"] = _clean_single_line_text(item.get("target"))[:200]
    item["signup_deadline"] = _clean_single_line_text(item.get("signup_deadline"))[:20]
    item["end_date"] = _clean_single_line_text(item.get("end_date"))[:20]
    item["contact_phone"] = _clean_single_line_text(item.get("contact_phone"))[:20]
    item["created_at"] = str(item.get("created_at") or "")
    item["updated_at"] = str(item.get("updated_at") or "")
    item["signup_count"] = int(item.get("signup_count") or 0)
    item.update(_training_program_status(item))
    return item


def _fetch_training_program_with_count_conn(conn: sqlite3.Connection, program_id: int) -> sqlite3.Row | None:
    return conn.execute(
        """
        SELECT
            p.id,
            p.title,
            p.schedule_text,
            p.location,
            p.target,
            p.signup_deadline,
            p.end_date,
            p.contact_phone,
            p.status,
            p.created_at,
            p.updated_at,
            COUNT(s.id) AS signup_count
        FROM training_programs p
        LEFT JOIN training_signups s ON s.program_id = p.id
        WHERE p.id = ?
        GROUP BY p.id
        """,
        (int(program_id or 0),),
    ).fetchone()


def list_training_programs(*, keyword: str = "", openid: str = "") -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        rows = conn.execute(
            """
            SELECT
                p.id,
                p.title,
                p.schedule_text,
                p.location,
                p.target,
                p.signup_deadline,
                p.end_date,
                p.contact_phone,
                p.status,
                p.created_at,
                p.updated_at,
                COUNT(s.id) AS signup_count
            FROM training_programs p
            LEFT JOIN training_signups s ON s.program_id = p.id
            GROUP BY p.id
            ORDER BY p.id DESC
            """
        ).fetchall()
        items = [_training_program_to_dict(row) for row in rows]
        viewer_openid = str(openid or "").strip()
        signed_ids: set[int] = set()
        if viewer_openid:
            signed_rows = conn.execute(
                "SELECT DISTINCT program_id FROM training_signups WHERE openid = ?",
                (viewer_openid,),
            ).fetchall()
            signed_ids = {int(r["program_id"] or 0) for r in signed_rows}
        for item in items:
            item["signed_up"] = int(item.get("id") or 0) in signed_ids
        keyword_text = _clean_single_line_text(keyword).lower()
        if keyword_text:
            items = [
                item
                for item in items
                if keyword_text in item["title"].lower()
                or keyword_text in item["schedule_text"].lower()
                or keyword_text in item["location"].lower()
                or keyword_text in item["target"].lower()
            ]
        return sorted(items, key=lambda item: (0 if item["is_open"] else 1, -int(item["id"] or 0)))
    finally:
        conn.close()


def create_training_program(data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        payload = _normalize_training_program_payload(data)
        created_at = now_iso()
        cursor = conn.execute(
            """
            INSERT INTO training_programs(
                title, schedule_text, location, target, signup_deadline, end_date,
                contact_phone, status, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload["title"],
                payload["schedule_text"],
                payload["location"],
                payload["target"],
                payload["signup_deadline"],
                payload["end_date"],
                payload["contact_phone"],
                payload["status"],
                created_at,
                created_at,
            ),
        )
        conn.commit()
        row = _fetch_training_program_with_count_conn(conn, int(cursor.lastrowid or 0))
        if row is None:
            raise ValueError("training_program_not_found")
        return _training_program_to_dict(row)
    finally:
        conn.close()


def update_training_program(program_id: int, data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        existing = conn.execute(
            "SELECT id FROM training_programs WHERE id = ? LIMIT 1",
            (int(program_id or 0),),
        ).fetchone()
        if existing is None:
            raise ValueError("training_program_not_found")
        payload = _normalize_training_program_payload(data)
        updated_at = now_iso()
        conn.execute(
            """
            UPDATE training_programs
            SET title = ?,
                schedule_text = ?,
                location = ?,
                target = ?,
                signup_deadline = ?,
                end_date = ?,
                contact_phone = ?,
                status = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                payload["title"],
                payload["schedule_text"],
                payload["location"],
                payload["target"],
                payload["signup_deadline"],
                payload["end_date"],
                payload["contact_phone"],
                payload["status"],
                updated_at,
                int(program_id or 0),
            ),
        )
        conn.commit()
        row = _fetch_training_program_with_count_conn(conn, int(program_id or 0))
        if row is None:
            raise ValueError("training_program_not_found")
        return _training_program_to_dict(row)
    finally:
        conn.close()


def end_training_program(program_id: int) -> dict[str, Any]:
    conn = get_connection()
    try:
        existing = conn.execute(
            "SELECT id FROM training_programs WHERE id = ? LIMIT 1",
            (int(program_id or 0),),
        ).fetchone()
        if existing is None:
            raise ValueError("training_program_not_found")
        conn.execute(
            "UPDATE training_programs SET status = 'ended', updated_at = ? WHERE id = ?",
            (now_iso(), int(program_id or 0)),
        )
        conn.commit()
        row = _fetch_training_program_with_count_conn(conn, int(program_id or 0))
        if row is None:
            raise ValueError("training_program_not_found")
        return _training_program_to_dict(row)
    finally:
        conn.close()


def delete_training_program(program_id: int) -> dict[str, str]:
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id, title FROM training_programs WHERE id = ? LIMIT 1",
            (int(program_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("training_program_not_found")
        conn.execute("DELETE FROM training_signups WHERE program_id = ?", (int(program_id or 0),))
        conn.execute("DELETE FROM training_programs WHERE id = ?", (int(program_id or 0),))
        conn.commit()
        return {"status": "deleted", "message": "就业培训已删除"}
    finally:
        conn.close()


def signup_training_program(token: str, program_id: int, data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _require_villager_user_conn(conn, token)
        name = _clean_single_line_text(data.get("name"))[:50]
        phone = normalize_phone_number(str(data.get("phone") or ""))
        if not name:
            raise ValueError("invalid_training_signup_name")
        if not is_valid_phone(phone):
            raise ValueError("invalid_phone")
        program_row = _fetch_training_program_with_count_conn(conn, int(program_id or 0))
        if program_row is None:
            raise ValueError("training_program_not_found")
        program = _training_program_to_dict(program_row)
        if not bool(program.get("is_open")):
            raise ValueError("training_signup_closed")
        openid = str(user["openid"] or "").strip()
        if openid:
            existing = conn.execute(
                "SELECT id FROM training_signups WHERE openid = ? AND program_id = ? LIMIT 1",
                (openid, int(program_id or 0)),
            ).fetchone()
            if existing is not None:
                raise ValueError("signup_conflict")
        created_at = now_iso()
        protected_phone = _protected_business_phone_fields(
            phone,
            purpose="training_signup_phone",
        )
        try:
            cursor = conn.execute(
                """
                INSERT INTO training_signups(
                    program_id, name, phone, phone_ciphertext, phone_hash, phone_masked, openid, created_at
                )
                VALUES(?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    int(program_id or 0),
                    name,
                    protected_phone["legacy_token"],
                    protected_phone["ciphertext"],
                    protected_phone["lookup_hash"],
                    protected_phone["masked"],
                    openid,
                    created_at,
                ),
            )
        except DBIntegrityError as exc:
            raise ValueError("signup_conflict") from exc
        conn.commit()
        row = conn.execute(
            """
            SELECT id, program_id, name, phone_masked AS phone, created_at
            FROM training_signups
            WHERE id = ?
            LIMIT 1
            """,
            (int(cursor.lastrowid or 0),),
        ).fetchone()
        return dict(row)
    finally:
        conn.close()


def list_training_signups_by_program(program_id: int) -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        program_row = conn.execute(
            "SELECT id FROM training_programs WHERE id = ? LIMIT 1",
            (int(program_id or 0),),
        ).fetchone()
        if program_row is None:
            raise ValueError("training_program_not_found")
        rows = conn.execute(
            """
            SELECT id, program_id, name, phone_masked AS phone, created_at
            FROM training_signups
            WHERE program_id = ?
            ORDER BY created_at ASC, id ASC
            """,
            (int(program_id or 0),),
        ).fetchall()
        return [{**dict(row), "openid": ""} for row in rows]
    finally:
        conn.close()


def export_training_signups_csv(program_id: int) -> bytes:
    conn = get_connection()
    try:
        program_row = conn.execute(
            """
            SELECT id, title, schedule_text, location, signup_deadline, contact_phone
            FROM training_programs
            WHERE id = ?
            LIMIT 1
            """,
            (int(program_id or 0),),
        ).fetchone()
        if program_row is None:
            raise ValueError("training_program_not_found")
        rows = conn.execute(
            """
            SELECT name, phone_masked AS phone, created_at
            FROM training_signups
            WHERE program_id = ?
            ORDER BY created_at ASC, id ASC
            """,
            (int(program_id or 0),),
        ).fetchall()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["培训标题", "培训时间", "培训地点", "报名截止", "联系人电话", "报名姓名", "联系电话（脱敏）", "报名时间"])
        for row in rows:
            writer.writerow(
                [
                    str(program_row["title"] or ""),
                    str(program_row["schedule_text"] or ""),
                    str(program_row["location"] or ""),
                    str(program_row["signup_deadline"] or ""),
                    str(program_row["contact_phone"] or ""),
                    str(row["name"] or ""),
                    str(row["phone"] or ""),
                    str(row["created_at"] or ""),
                ]
            )
        return output.getvalue().encode("utf-8-sig")
    finally:
        conn.close()


# ==========================================================================
# 志愿活动（与就业培训同构的报名型；独立表 volunteer_activities / volunteer_activity_signups）
# ==========================================================================

def _normalize_volunteer_activity_payload(data: dict[str, Any]) -> dict[str, str]:
    title = _clean_single_line_text(data.get("title"))[:100]
    if not title:
        raise ValueError("invalid_volunteer_activity_title")
    signup_deadline = _clean_single_line_text(data.get("signup_deadline"))[:20]
    end_date = _clean_single_line_text(data.get("end_date"))[:20]
    deadline_date = _parse_training_date(signup_deadline)
    finish_date = _parse_training_date(end_date)
    if signup_deadline and deadline_date is None:
        raise ValueError("invalid_volunteer_activity_date")
    if end_date and finish_date is None:
        raise ValueError("invalid_volunteer_activity_date")
    if deadline_date is not None and finish_date is not None and finish_date < deadline_date:
        raise ValueError("invalid_volunteer_activity_date_range")
    return {
        "title": title,
        "schedule_text": _clean_single_line_text(data.get("schedule_text"))[:120],
        "location": _clean_single_line_text(data.get("location"))[:120],
        "target": _clean_single_line_text(data.get("target"))[:200],
        "detail": str(data.get("detail") or "").strip()[:2000],
        "signup_deadline": signup_deadline[:10] if deadline_date is not None else "",
        "end_date": end_date[:10] if finish_date is not None else "",
        "contact_phone": _clean_single_line_text(data.get("contact_phone"))[:20],
        "status": _normalize_training_status(data.get("status")),
    }


def _volunteer_activity_to_dict(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["id"] = int(item.get("id") or 0)
    item["title"] = _clean_single_line_text(item.get("title"))[:100]
    item["schedule_text"] = _clean_single_line_text(item.get("schedule_text"))[:120]
    item["location"] = _clean_single_line_text(item.get("location"))[:120]
    item["target"] = _clean_single_line_text(item.get("target"))[:200]
    item["detail"] = str(item.get("detail") or "").strip()[:2000]
    item["signup_deadline"] = _clean_single_line_text(item.get("signup_deadline"))[:20]
    item["end_date"] = _clean_single_line_text(item.get("end_date"))[:20]
    item["contact_phone"] = _clean_single_line_text(item.get("contact_phone"))[:20]
    item["created_at"] = str(item.get("created_at") or "")
    item["updated_at"] = str(item.get("updated_at") or "")
    item["signup_count"] = int(item.get("signup_count") or 0)
    item.update(_training_program_status(item))
    return item


def _fetch_volunteer_activity_with_count_conn(conn: sqlite3.Connection, activity_id: int) -> sqlite3.Row | None:
    return conn.execute(
        """
        SELECT
            a.id,
            a.title,
            a.schedule_text,
            a.location,
            a.target,
            a.detail,
            a.signup_deadline,
            a.end_date,
            a.contact_phone,
            a.status,
            a.created_at,
            a.updated_at,
            COUNT(s.id) AS signup_count
        FROM volunteer_activities a
        LEFT JOIN volunteer_activity_signups s ON s.activity_id = a.id
        WHERE a.id = ?
        GROUP BY a.id
        """,
        (int(activity_id or 0),),
    ).fetchone()


def list_volunteer_activities(*, keyword: str = "", openid: str = "") -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        rows = conn.execute(
            """
            SELECT
                a.id,
                a.title,
                a.schedule_text,
                a.location,
                a.target,
                a.detail,
                a.signup_deadline,
                a.end_date,
                a.contact_phone,
                a.status,
                a.created_at,
                a.updated_at,
                COUNT(s.id) AS signup_count
            FROM volunteer_activities a
            LEFT JOIN volunteer_activity_signups s ON s.activity_id = a.id
            GROUP BY a.id
            ORDER BY a.id DESC
            """
        ).fetchall()
        items = [_volunteer_activity_to_dict(row) for row in rows]
        viewer_openid = str(openid or "").strip()
        signed_ids: set[int] = set()
        if viewer_openid:
            signed_rows = conn.execute(
                "SELECT DISTINCT activity_id FROM volunteer_activity_signups WHERE openid = ?",
                (viewer_openid,),
            ).fetchall()
            signed_ids = {int(r["activity_id"] or 0) for r in signed_rows}
        for item in items:
            item["signed_up"] = int(item.get("id") or 0) in signed_ids
        keyword_text = _clean_single_line_text(keyword).lower()
        if keyword_text:
            items = [
                item
                for item in items
                if keyword_text in item["title"].lower()
                or keyword_text in item["schedule_text"].lower()
                or keyword_text in item["location"].lower()
                or keyword_text in item["target"].lower()
            ]
        return sorted(items, key=lambda item: (0 if item["is_open"] else 1, -int(item["id"] or 0)))
    finally:
        conn.close()


def create_volunteer_activity(data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        payload = _normalize_volunteer_activity_payload(data)
        created_at = now_iso()
        cursor = conn.execute(
            """
            INSERT INTO volunteer_activities(
                title, schedule_text, location, target, detail, signup_deadline, end_date,
                contact_phone, status, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload["title"],
                payload["schedule_text"],
                payload["location"],
                payload["target"],
                payload["detail"],
                payload["signup_deadline"],
                payload["end_date"],
                payload["contact_phone"],
                payload["status"],
                created_at,
                created_at,
            ),
        )
        conn.commit()
        row = _fetch_volunteer_activity_with_count_conn(conn, int(cursor.lastrowid or 0))
        if row is None:
            raise ValueError("volunteer_activity_not_found")
        return _volunteer_activity_to_dict(row)
    finally:
        conn.close()


def update_volunteer_activity(activity_id: int, data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        existing = conn.execute(
            "SELECT id FROM volunteer_activities WHERE id = ? LIMIT 1",
            (int(activity_id or 0),),
        ).fetchone()
        if existing is None:
            raise ValueError("volunteer_activity_not_found")
        payload = _normalize_volunteer_activity_payload(data)
        updated_at = now_iso()
        conn.execute(
            """
            UPDATE volunteer_activities
            SET title = ?,
                schedule_text = ?,
                location = ?,
                target = ?,
                detail = ?,
                signup_deadline = ?,
                end_date = ?,
                contact_phone = ?,
                status = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                payload["title"],
                payload["schedule_text"],
                payload["location"],
                payload["target"],
                payload["detail"],
                payload["signup_deadline"],
                payload["end_date"],
                payload["contact_phone"],
                payload["status"],
                updated_at,
                int(activity_id or 0),
            ),
        )
        conn.commit()
        row = _fetch_volunteer_activity_with_count_conn(conn, int(activity_id or 0))
        if row is None:
            raise ValueError("volunteer_activity_not_found")
        return _volunteer_activity_to_dict(row)
    finally:
        conn.close()


def end_volunteer_activity(activity_id: int) -> dict[str, Any]:
    conn = get_connection()
    try:
        existing = conn.execute(
            "SELECT id FROM volunteer_activities WHERE id = ? LIMIT 1",
            (int(activity_id or 0),),
        ).fetchone()
        if existing is None:
            raise ValueError("volunteer_activity_not_found")
        conn.execute(
            "UPDATE volunteer_activities SET status = 'ended', updated_at = ? WHERE id = ?",
            (now_iso(), int(activity_id or 0)),
        )
        conn.commit()
        row = _fetch_volunteer_activity_with_count_conn(conn, int(activity_id or 0))
        if row is None:
            raise ValueError("volunteer_activity_not_found")
        return _volunteer_activity_to_dict(row)
    finally:
        conn.close()


def delete_volunteer_activity(activity_id: int) -> dict[str, str]:
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id, title FROM volunteer_activities WHERE id = ? LIMIT 1",
            (int(activity_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("volunteer_activity_not_found")
        conn.execute("DELETE FROM volunteer_activity_signups WHERE activity_id = ?", (int(activity_id or 0),))
        conn.execute("DELETE FROM volunteer_activities WHERE id = ?", (int(activity_id or 0),))
        conn.commit()
        return {"status": "deleted", "message": "志愿活动已删除"}
    finally:
        conn.close()


def signup_volunteer_activity(token: str, activity_id: int, data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _require_villager_user_conn(conn, token)
        name = _clean_single_line_text(data.get("name"))[:50]
        phone = normalize_phone_number(str(data.get("phone") or ""))
        if not name:
            raise ValueError("invalid_volunteer_signup_name")
        if not is_valid_phone(phone):
            raise ValueError("invalid_phone")
        activity_row = _fetch_volunteer_activity_with_count_conn(conn, int(activity_id or 0))
        if activity_row is None:
            raise ValueError("volunteer_activity_not_found")
        activity = _volunteer_activity_to_dict(activity_row)
        if not bool(activity.get("is_open")):
            raise ValueError("volunteer_signup_closed")
        openid = str(user["openid"] or "").strip()
        if openid:
            existing = conn.execute(
                "SELECT id FROM volunteer_activity_signups WHERE openid = ? AND activity_id = ? LIMIT 1",
                (openid, int(activity_id or 0)),
            ).fetchone()
            if existing is not None:
                raise ValueError("signup_conflict")
        created_at = now_iso()
        protected_phone = _protected_business_phone_fields(
            phone,
            purpose="volunteer_signup_phone",
        )
        try:
            cursor = conn.execute(
                """
                INSERT INTO volunteer_activity_signups(
                    activity_id, name, phone, phone_ciphertext, phone_hash, phone_masked, openid, created_at
                )
                VALUES(?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    int(activity_id or 0),
                    name,
                    protected_phone["legacy_token"],
                    protected_phone["ciphertext"],
                    protected_phone["lookup_hash"],
                    protected_phone["masked"],
                    openid,
                    created_at,
                ),
            )
        except DBIntegrityError as exc:
            raise ValueError("signup_conflict") from exc
        conn.commit()
        row = conn.execute(
            """
            SELECT id, activity_id, name, phone_masked AS phone, created_at
            FROM volunteer_activity_signups
            WHERE id = ?
            LIMIT 1
            """,
            (int(cursor.lastrowid or 0),),
        ).fetchone()
        return dict(row)
    finally:
        conn.close()


def list_volunteer_activity_signups_by_activity(activity_id: int) -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        activity_row = conn.execute(
            "SELECT id FROM volunteer_activities WHERE id = ? LIMIT 1",
            (int(activity_id or 0),),
        ).fetchone()
        if activity_row is None:
            raise ValueError("volunteer_activity_not_found")
        rows = conn.execute(
            """
            SELECT id, activity_id, name, phone_masked AS phone, created_at
            FROM volunteer_activity_signups
            WHERE activity_id = ?
            ORDER BY created_at ASC, id ASC
            """,
            (int(activity_id or 0),),
        ).fetchall()
        return [{**dict(row), "openid": ""} for row in rows]
    finally:
        conn.close()


def export_volunteer_activity_signups_csv(activity_id: int) -> bytes:
    conn = get_connection()
    try:
        activity_row = conn.execute(
            """
            SELECT id, title, schedule_text, location, signup_deadline, contact_phone
            FROM volunteer_activities
            WHERE id = ?
            LIMIT 1
            """,
            (int(activity_id or 0),),
        ).fetchone()
        if activity_row is None:
            raise ValueError("volunteer_activity_not_found")
        rows = conn.execute(
            """
            SELECT name, phone_masked AS phone, created_at
            FROM volunteer_activity_signups
            WHERE activity_id = ?
            ORDER BY created_at ASC, id ASC
            """,
            (int(activity_id or 0),),
        ).fetchall()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["活动标题", "活动时间", "活动地点", "报名截止", "联系人电话", "报名姓名", "联系电话（脱敏）", "报名时间"])
        for row in rows:
            writer.writerow(
                [
                    str(activity_row["title"] or ""),
                    str(activity_row["schedule_text"] or ""),
                    str(activity_row["location"] or ""),
                    str(activity_row["signup_deadline"] or ""),
                    str(activity_row["contact_phone"] or ""),
                    str(row["name"] or ""),
                    str(row["phone"] or ""),
                    str(row["created_at"] or ""),
                ]
            )
        return output.getvalue().encode("utf-8-sig")
    finally:
        conn.close()


# ==========================================================================
# 村民建议（纯收集箱：实名村民可提交；前台不公开身份、后端记录提交人供溯源；非办件、无状态）
# ==========================================================================

def _proposal_submitted_today_conn(conn: sqlite3.Connection, openid: str) -> bool:
    """该 openid 今天(北京时间自然日)是否已提交过建议。"""
    viewer_openid = str(openid or "").strip()
    if not viewer_openid:
        return False
    row = conn.execute(
        "SELECT created_at FROM proposals WHERE user_openid = ? ORDER BY id DESC LIMIT 1",
        (viewer_openid,),
    ).fetchone()
    if row is None:
        return False
    try:
        last_dt = datetime.fromisoformat(str(row["created_at"] or ""))
    except ValueError:
        return False
    if last_dt.tzinfo is None:
        last_dt = last_dt.replace(tzinfo=timezone.utc)
    return last_dt.astimezone(CHINA_TIMEZONE).date() == _today_china_date()


def get_openid_by_token_optional(token: str) -> str:
    if not str(token or "").strip():
        return ""
    conn = get_connection()
    try:
        row = _get_user_by_token_conn_optional(conn, token)
        return str(row["openid"] or "").strip() if row is not None else ""
    finally:
        conn.close()


def proposal_status(token: str) -> dict[str, Any]:
    conn = get_connection()
    try:
        user = _require_verified_villager_user_conn(conn, token)
        openid = str(user["openid"] or "").strip()
        return {"submitted_today": _proposal_submitted_today_conn(conn, openid)}
    finally:
        conn.close()


def create_proposal(token: str, data: dict[str, Any]) -> dict[str, Any]:
    content = str(data.get("content") or "").strip()
    if not content:
        raise ValueError("invalid_proposal_content")
    if len(content) > 500:
        content = content[:500]
    conn = get_connection()
    try:
        user = _require_verified_villager_user_conn(conn, token)
        openid = str(user["openid"] or "").strip()
        if _proposal_submitted_today_conn(conn, openid):
            raise ValueError("proposal_daily_limit")
        villager_name = str(user["real_name"] or "").strip() or str(user["nickname"] or "").strip()
        created_at = now_iso()
        cursor = conn.execute(
            """
            INSERT INTO proposals(content, user_openid, villager_name, created_at)
            VALUES(?, ?, ?, ?)
            """,
            (content, openid, villager_name, created_at),
        )
        conn.commit()
        return {"id": int(cursor.lastrowid or 0), "created_at": created_at}
    finally:
        conn.close()


def list_proposals() -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        rows = conn.execute(
            """
            SELECT id, content, user_openid, villager_name, created_at
            FROM proposals
            ORDER BY id DESC
            """
        ).fetchall()
        return [
            {
                "id": int(row["id"] or 0),
                "content": str(row["content"] or ""),
                "villager_name": str(row["villager_name"] or ""),
                "user_openid": str(row["user_openid"] or ""),
                "created_at": str(row["created_at"] or ""),
            }
            for row in rows
        ]
    finally:
        conn.close()


# ==========================================================================
# 本地好物（local_product）：纯展示，无价格/库存/购买；干部后台维护，村民端只读
# ==========================================================================

def _normalize_local_product_payload(data: dict[str, Any]) -> dict[str, Any]:
    name = _clean_single_line_text(data.get("name"))[:120]
    if not name:
        raise ValueError("invalid_local_product_name")
    try:
        sort_order = int(data.get("sort_order") or 0)
    except (TypeError, ValueError):
        sort_order = 0
    sort_order = max(0, min(sort_order, 9999))
    return {
        "name": name,
        "subtitle": _clean_single_line_text(data.get("subtitle"))[:255],
        "cover_url": str(data.get("cover_url") or "").strip()[:2048],
        "detail_img_url": str(data.get("detail_img_url") or "").strip()[:2048],
        "history": str(data.get("history") or "").strip()[:5000],
        "craft": str(data.get("craft") or "").strip()[:5000],
        "flavor": str(data.get("flavor") or "").strip()[:5000],
        "sort_order": sort_order,
    }


def _local_product_to_dict(row: sqlite3.Row | dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    return {
        "id": int(item.get("id") or 0),
        "name": str(item.get("name") or ""),
        "subtitle": str(item.get("subtitle") or ""),
        "cover_url": build_upload_access_url(str(item.get("cover_url") or "")),
        "detail_img_url": build_upload_access_url(str(item.get("detail_img_url") or "")),
        "history": str(item.get("history") or ""),
        "craft": str(item.get("craft") or ""),
        "flavor": str(item.get("flavor") or ""),
        "sort_order": int(item.get("sort_order") or 0),
        "created_at": str(item.get("created_at") or ""),
        "updated_at": str(item.get("updated_at") or ""),
    }


_LOCAL_PRODUCT_COLUMNS = (
    "id, name, subtitle, cover_url, detail_img_url, history, craft, flavor, sort_order, created_at, updated_at"
)


def list_local_products() -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        rows = conn.execute(
            f"SELECT {_LOCAL_PRODUCT_COLUMNS} FROM local_product ORDER BY sort_order ASC, id ASC"
        ).fetchall()
        return [_local_product_to_dict(row) for row in rows]
    finally:
        conn.close()


def get_local_product(product_id: int) -> dict[str, Any]:
    conn = get_connection()
    try:
        row = conn.execute(
            f"SELECT {_LOCAL_PRODUCT_COLUMNS} FROM local_product WHERE id = ? LIMIT 1",
            (int(product_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("local_product_not_found")
        return _local_product_to_dict(row)
    finally:
        conn.close()


def create_local_product(data: dict[str, Any]) -> dict[str, Any]:
    payload = _normalize_local_product_payload(data)
    conn = get_connection()
    try:
        created_at = now_iso()
        cursor = conn.execute(
            """
            INSERT INTO local_product(
                name, subtitle, cover_url, detail_img_url, history, craft, flavor,
                sort_order, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload["name"],
                payload["subtitle"],
                payload["cover_url"],
                payload["detail_img_url"],
                payload["history"],
                payload["craft"],
                payload["flavor"],
                payload["sort_order"],
                created_at,
                created_at,
            ),
        )
        conn.commit()
        row = conn.execute(
            f"SELECT {_LOCAL_PRODUCT_COLUMNS} FROM local_product WHERE id = ? LIMIT 1",
            (int(cursor.lastrowid or 0),),
        ).fetchone()
        return _local_product_to_dict(row)
    finally:
        conn.close()


def update_local_product(product_id: int, data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        existing = conn.execute(
            "SELECT id FROM local_product WHERE id = ? LIMIT 1",
            (int(product_id or 0),),
        ).fetchone()
        if existing is None:
            raise ValueError("local_product_not_found")
        payload = _normalize_local_product_payload(data)
        conn.execute(
            """
            UPDATE local_product
            SET name = ?,
                subtitle = ?,
                cover_url = ?,
                detail_img_url = ?,
                history = ?,
                craft = ?,
                flavor = ?,
                sort_order = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                payload["name"],
                payload["subtitle"],
                payload["cover_url"],
                payload["detail_img_url"],
                payload["history"],
                payload["craft"],
                payload["flavor"],
                payload["sort_order"],
                now_iso(),
                int(product_id or 0),
            ),
        )
        conn.commit()
        row = conn.execute(
            f"SELECT {_LOCAL_PRODUCT_COLUMNS} FROM local_product WHERE id = ? LIMIT 1",
            (int(product_id or 0),),
        ).fetchone()
        return _local_product_to_dict(row)
    finally:
        conn.close()


def delete_local_product(product_id: int) -> dict[str, str]:
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id FROM local_product WHERE id = ? LIMIT 1",
            (int(product_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("local_product_not_found")
        conn.execute("DELETE FROM local_product WHERE id = ?", (int(product_id or 0),))
        conn.commit()
        return {"status": "deleted", "message": "好物已删除"}
    finally:
        conn.close()


def list_ledgers() -> list[dict[str, Any]]:
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT id, category, content, status, owner, photo_urls, created_at "
            "FROM ledgers ORDER BY id DESC"
        ).fetchall()
        items: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            item["photo_urls"] = _parse_photo_urls_field(
                item.get("photo_urls"),
                allowed_scenes=SCENES_LEDGER,
                allow_external=False,
            )
            items.append(item)
        return items
    finally:
        conn.close()


def create_ledger(data: dict[str, Any]) -> dict[str, Any]:
    conn = get_connection()
    try:
        created_at = now_iso()
        photo_urls = _normalize_photo_urls(
            data.get("photo_urls"),
            allowed_scenes=SCENES_LEDGER,
            allow_external=False,
            strict_scene=True,
        )
        cursor = conn.execute(
            "INSERT INTO ledgers(category, content, status, owner, photo_urls, created_at) "
            "VALUES(?, ?, ?, ?, ?, ?)",
            (
                data["category"],
                data["content"],
                data["status"],
                data["owner"],
                _json_text(photo_urls),
                created_at,
            ),
        )
        conn.commit()
        new_id = cursor.lastrowid
        row = conn.execute(
            "SELECT id, category, content, status, owner, photo_urls, created_at "
            "FROM ledgers WHERE id = ?",
            (new_id,),
        ).fetchone()
        item = dict(row)
        item["photo_urls"] = _parse_photo_urls_field(
            item.get("photo_urls"),
            allowed_scenes=SCENES_LEDGER,
            allow_external=False,
        )
        return item
    finally:
        conn.close()
