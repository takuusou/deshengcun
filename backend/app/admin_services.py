from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
import csv
import io
import json
import logging
import mimetypes
import re
from urllib.parse import unquote, urlsplit
from urllib.request import urlopen
import zipfile
from typing import Any

from .admin_auth import bootstrap_admin_user, get_current_admin, reset_admin_account_password
from .db import get_connection
from .services import (
    DEFAULT_CADRE_MANAGER_NAMES,
    HIDDEN_GOV_SERVICE_KEYS,
    SCENES_GOV_REPLY,
    SCENES_NOTICE,
    SCENES_OPEN,
    SCENES_REQUEST_REPLY,
    VALID_GOV_DETAIL_TYPES,
    _find_existing_identity_binding_conn,
    _find_roster_match_conn,
    _gov_service_to_dict,
    _require_active_roster_conn,
    _id_hash,
    _is_phone_verified_user,
    _normalize_id_card,
    _parse_csv_rows,
    _parse_xlsx_rows,
    _gov_application_to_dict,
    _map_roster_row,
    _normalize_photo_urls,
    _notice_to_dict,
    _parse_json_object,
    _resolve_gender,
    _sanitize_gov_detail_payload,
    _sanitize_metric_key,
    _seed_default_gov_services_conn,
    _seed_default_metrics_conn,
    _sync_villager_registry_conn,
    _villager_request_to_dict,
    build_upload_access_url,
    calc_auth_status,
    create_notice,
    create_training_program,
    delete_training_program,
    decrypt_id_card_number,
    decrypt_phone_number,
    encrypt_id_card_number,
    end_training_program,
    export_training_signups_csv,
    get_village_open_profile,
    get_notice_detail,
    get_stored_file_record,
    is_valid_id_card,
    list_notices,
    list_training_programs,
    list_training_signups_by_program,
    list_village_metrics,
    create_volunteer_activity,
    delete_volunteer_activity,
    end_volunteer_activity,
    export_volunteer_activity_signups_csv,
    list_volunteer_activities,
    list_volunteer_activity_signups_by_activity,
    update_volunteer_activity,
    list_proposals,
    list_local_products,
    create_local_product,
    update_local_product,
    delete_local_product,
    mask_id_card,
    mask_phone,
    normalize_phone_number,
    normalize_cadre_role_level,
    store_uploaded_file_record,
    update_training_program,
    update_village_open_profile,
)
from .runtime_config import get_backend_upload_dir
from .storage import build_download_url, resolve_local_file_path


security_logger = logging.getLogger("deshengcun.security")


REQUEST_STATUS_SET = {"pending", "processing", "resolved", "closed"}
REQUEST_TYPE_SET = {"complaint", "service", "dibao", "tekun"}
REQUEST_WRITE_ROLE_SET = {"super_admin", "cadre"}
GOV_APPLICATION_STATUS_SET = {"pending", "processing", "resolved", "rejected", "closed"}
REQUEST_STATUS_TRANSITIONS = {
    "pending": REQUEST_STATUS_SET,
    "processing": {"processing", "resolved", "closed"},
    "resolved": {"resolved"},
    "closed": {"closed"},
}
GOV_APPLICATION_STATUS_TRANSITIONS = {
    "pending": GOV_APPLICATION_STATUS_SET,
    "processing": {"processing", "resolved", "rejected", "closed"},
    "resolved": {"resolved"},
    "rejected": {"rejected"},
    "closed": {"closed"},
}
CADRE_AUTH_STATUS_SET = {"active", "disabled"}
CADRE_ROLE_LEVEL_SET = {"super_admin", "cadre"}
ADMIN_ROLE_KEY_SET = {"super_admin", "cadre", "readonly"}
def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _extract_birth_date_from_id_card(id_card_number: str) -> date | None:
    text = _normalize_id_card(id_card_number)
    if not text:
        return None
    birth_text = ""
    birth_format = ""
    if len(text) == 18:
        birth_text = text[6:14]
        birth_format = "%Y%m%d"
    elif len(text) == 15 and text.isdigit():
        birth_text = f"19{text[6:12]}"
        birth_format = "%Y%m%d"
    if not birth_text:
        return None
    try:
        return datetime.strptime(birth_text, birth_format).date()
    except ValueError:
        return None


def _calculate_age_from_id_card(id_card_number: str, *, today: date | None = None) -> int | None:
    birth_date = _extract_birth_date_from_id_card(id_card_number)
    if birth_date is None:
        return None
    current_date = today or datetime.now().date()
    age = current_date.year - birth_date.year
    if (current_date.month, current_date.day) < (birth_date.month, birth_date.day):
        age -= 1
    if age < 0 or age > 130:
        return None
    return age


def _safe_decrypt_phone(encrypted_payload: str) -> str:
    try:
        phone_number = decrypt_phone_number(encrypted_payload)
    except ValueError:
        return ""
    return phone_number if len(phone_number) == 11 else ""


def _safe_decrypt_id_card(encrypted_payload: str) -> str:
    try:
        return decrypt_id_card_number(encrypted_payload)
    except ValueError:
        return ""


def _serialize_admin_villager_identity(row: Any) -> dict[str, Any]:
    phone_number = _safe_decrypt_phone(str(row["phone_ciphertext"] or ""))
    id_card_number = _safe_decrypt_id_card(str(row["id_card_encrypted"] or ""))
    gender = _resolve_gender(str(row["gender"] or ""), id_card_number)
    birth_date = _extract_birth_date_from_id_card(id_card_number)
    current_date = datetime.now().date()
    age = _calculate_age_from_id_card(id_card_number, today=current_date)
    is_turning_60_this_year = bool(
        birth_date is not None
        and birth_date.year == current_date.year - 60
        and (current_date.month, current_date.day) < (birth_date.month, birth_date.day)
    )
    return {
        "user_id": int(row["id"] or 0),
        "real_name": str(row["real_name"] or ""),
        "gender": gender,
        "phone_masked": str(row["phone_masked"] or "") or mask_phone(phone_number),
        "id_card_masked": mask_id_card(id_card_number or str(row["id_card_masked"] or "")),
        "auth_status": str(row["auth_status"] or "unverified"),
        "realname_status": str(row["realname_status"] or "unsubmitted"),
        "roster_match_status": str(row["roster_match_status"] or "unmatched"),
        "village_group": str(row["village_group"] or ""),
        "age": age,
        "is_turning_60_this_year": is_turning_60_this_year,
        "is_age_60_plus": bool(age is not None and age >= 60),
        "is_age_80_plus": bool(age is not None and age >= 80),
        "updated_at": str(row["updated_at"] or ""),
        "realname_submitted_at": str(row["realname_submitted_at"] or "") or str(row["updated_at"] or ""),
    }


def _format_export_phone_value(full_phone: str, masked_phone: str) -> str:
    _ = full_phone
    masked_text = str(masked_phone or "").strip()
    if not masked_text:
        return ""
    if "****" in masked_text:
        return masked_text.replace("****", "-****-")
    if len(masked_text) == 11 and masked_text.isdigit():
        return mask_phone(masked_text).replace("****", "-****-")
    return masked_text


def _format_export_datetime(value: str) -> datetime | str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return text.replace("T", " ")
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone(timedelta(hours=8))).replace(tzinfo=None)
    return parsed


def _build_admin_villager_identity_query(keyword_text: str = "") -> tuple[str, list[Any]]:
    query = """
        SELECT
            id,
            real_name,
            gender,
            phone_masked,
            phone_ciphertext,
            id_card_masked,
            id_card_encrypted,
            auth_status,
            realname_status,
            roster_match_status,
            village_group,
            updated_at,
            (
                SELECT vr.realname_submitted_at
                FROM villager_registry vr
                WHERE vr.user_id = users.id
            ) AS realname_submitted_at
        FROM users
        WHERE role = 'villager'
          AND (real_name != '' OR id_card_hash != '' OR realname_status != 'unsubmitted')
    """
    params: list[Any] = []
    if keyword_text:
        like_text = f"%{keyword_text}%"
        query += " AND (real_name LIKE ? OR phone_masked LIKE ? OR id_card_masked LIKE ? OR village_group LIKE ?)"
        params.extend([like_text, like_text, like_text, like_text])
    return query, params


def _json_text(value: dict[str, Any]) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _audit_admin_action_conn(
    conn: Any,
    *,
    admin_user_id: int,
    action: str,
    target_type: str,
    target_id: str,
    detail: dict[str, Any],
) -> None:
    conn.execute(
        """
        INSERT INTO audit_logs(
            actor_user_id, actor_admin_user_id, actor_type,
            action, target_type, target_id, detail_json, ip, user_agent, created_at
        )
        VALUES(?, ?, 'admin', ?, ?, ?, ?, '', '', ?)
        """,
        (None, admin_user_id, action, target_type, target_id, _json_text(detail), _now_iso()),
    )


def _normalize_request_status(value: str) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    if text not in REQUEST_STATUS_SET:
        raise ValueError("invalid_request_status")
    return text


def _require_status_transition(
    current_status: str,
    next_status: str,
    transitions: dict[str, set[str]],
    error_code: str,
) -> None:
    current = str(current_status or "").strip().lower()
    if current not in transitions or next_status not in transitions[current]:
        raise ValueError(error_code)


def _normalize_request_type(value: str) -> str:
    text = str(value or "").strip().lower()
    if not text or text == "all":
        return ""
    if text not in REQUEST_TYPE_SET:
        raise ValueError("invalid_request_type")
    return text


def _require_admin_write_role(admin_user: dict[str, Any]) -> None:
    role_key = str(admin_user.get("role_key") or "").strip().lower()
    if role_key not in REQUEST_WRITE_ROLE_SET:
        raise ValueError("admin_write_forbidden")


def _require_admin_super_role(admin_user: dict[str, Any]) -> None:
    role_key = str(admin_user.get("role_key") or "").strip().lower()
    if role_key != "super_admin":
        raise ValueError("admin_super_required")


def _build_in_clause(values: list[int]) -> tuple[str, list[Any]]:
    ids = [int(value) for value in values if int(value or 0) > 0]
    if not ids:
        return "(NULL)", []
    return f"({', '.join('?' for _ in ids)})", ids


def _count_rows(conn: Any, table_name: str, where_sql: str = "", params: list[Any] | None = None) -> int:
    query = f"SELECT COUNT(1) AS c FROM {table_name}"
    if where_sql:
        query += f" WHERE {where_sql}"
    row = conn.execute(query, params or []).fetchone()
    return int((row["c"] if row else 0) or 0)


def _delete_rows(conn: Any, table_name: str, where_sql: str = "", params: list[Any] | None = None) -> int:
    deleted_count = _count_rows(conn, table_name, where_sql, params)
    if deleted_count <= 0:
        return 0
    query = f"DELETE FROM {table_name}"
    if where_sql:
        query += f" WHERE {where_sql}"
    conn.execute(query, params or [])
    return deleted_count


def _delete_rows_by_user_ids(conn: Any, table_name: str, user_column: str, user_ids: list[int]) -> int:
    in_clause, params = _build_in_clause(user_ids)
    if not params:
        return 0
    return _delete_rows(conn, table_name, f"{user_column} IN {in_clause}", params)


def purge_admin_test_data(token: str, confirm_text: str) -> dict[str, Any]:
    if str(confirm_text or "").strip().upper() != "DELETE_TEST_DATA":
        raise ValueError("invalid_cleanup_confirmation")

    conn = get_connection()
    try:
        admin_user = get_current_admin(token)
        _require_admin_super_role(admin_user)

        preserved_cadre_user_ids = sorted(
            {
                int(row["user_id"] or 0)
                for row in conn.execute(
                    "SELECT user_id FROM cadre_authorizations WHERE user_id IS NOT NULL AND user_id != 0"
                ).fetchall()
                if int(row["user_id"] or 0) > 0
            }
        )
        villager_user_ids = [
            int(row["id"] or 0)
            for row in conn.execute("SELECT id FROM users WHERE role = 'villager'").fetchall()
            if int(row["id"] or 0) > 0
        ]
        removable_user_ids = [user_id for user_id in villager_user_ids if user_id not in preserved_cadre_user_ids]

        deleted_counts = {
            "notice_delivery_logs_deleted": _delete_rows(conn, "notice_delivery_logs"),
            "notice_subscriptions_deleted": _delete_rows(conn, "notice_subscriptions"),
            "notices_deleted": _delete_rows(conn, "notices"),
            "villager_requests_deleted": _delete_rows(conn, "villager_requests"),
            "gov_service_applications_deleted": _delete_rows(conn, "gov_service_applications"),
            "auth_sessions_deleted": _delete_rows_by_user_ids(conn, "auth_sessions", "user_id", villager_user_ids),
            "privacy_consents_deleted": _delete_rows_by_user_ids(conn, "privacy_consents", "user_id", villager_user_ids),
            "villager_registry_deleted": _delete_rows_by_user_ids(conn, "villager_registry", "user_id", villager_user_ids),
            "stored_files_deleted": _delete_rows(
                conn,
                "stored_files",
                "scene IN (?, ?, ?, ?, ?)",
                ["request", "request_reply", "gov_apply", "gov_reply", "notice"],
            ),
        }

        villager_users_scrubbed = 0
        if villager_user_ids:
            in_clause, params = _build_in_clause(villager_user_ids)
            villager_users_scrubbed = _count_rows(conn, "users", f"id IN {in_clause} AND role = 'villager'", params)
            conn.execute(
                f"""
                UPDATE users
                SET
                    phone = '',
                    phone_ciphertext = NULL,
                    phone_hash = '',
                    phone_masked = '',
                    phone_last4 = '',
                    phone_bound = 0,
                    phone_verified_at = '',
                    real_name = '',
                    gender = '',
                    id_card_masked = '',
                    id_card_last4 = '',
                    id_card_hash = '',
                    id_card_encrypted = NULL,
                    face_photo_url = '',
                    face_verified_at = '',
                    realname_status = 'unsubmitted',
                    face_status = 'unsubmitted',
                    auth_status = 'unverified',
                    roster_match_status = 'unmatched',
                    village_group = '',
                    updated_at = ?
                WHERE id IN {in_clause} AND role = 'villager'
                """,
                [_now_iso(), *params],
            )

        villager_users_deleted = _delete_rows_by_user_ids(conn, "users", "id", removable_user_ids)
        audit_logs_deleted = _delete_rows(conn, "audit_logs")

        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_purge_test_data",
            target_type="maintenance",
            target_id="test_data",
            detail={
                **deleted_counts,
                "villager_users_scrubbed": villager_users_scrubbed,
                "villager_users_deleted": villager_users_deleted,
                "preserved_cadre_user_ids": preserved_cadre_user_ids,
            },
        )
        conn.commit()

        return {
            "status": "ok",
            "message": "业务数据与村民实名留痕已清理。",
            **deleted_counts,
            "villager_users_deleted": villager_users_deleted,
            "villager_users_scrubbed": villager_users_scrubbed,
            "audit_logs_deleted": audit_logs_deleted,
            "preserved_cadre_user_ids": preserved_cadre_user_ids,
        }
    finally:
        conn.close()


def import_admin_villager_roster(token: str, file_name: str, file_bytes: bytes) -> dict[str, Any]:
    ext = (str(file_name or "").rsplit(".", 1)[-1] if "." in str(file_name or "") else "").lower()
    if ext not in {"csv", "xlsx"}:
        raise ValueError("unsupported_file_type")

    if ext == "csv":
        raw_rows = _parse_csv_rows(file_bytes)
    else:
        raw_rows = _parse_xlsx_rows(file_bytes)

    conn = get_connection()
    try:
        admin_user = get_current_admin(token)
        _require_admin_super_role(admin_user)

        entries: list[dict[str, str]] = []
        skipped_count = 0
        dedup: set[tuple[str, str, str, str]] = set()
        for raw_row in raw_rows:
            row = _map_roster_row(raw_row)
            if row is None:
                skipped_count += 1
                continue
            key = (
                row["full_name"],
                row["id_last4"],
                row["id_hash"],
                row["group_name"],
            )
            if key in dedup:
                skipped_count += 1
                continue
            dedup.add(key)
            entries.append(row)

        if not entries:
            raise ValueError("empty_roster_file")

        now = _now_iso()
        imported_by = str(admin_user.get("real_name") or admin_user.get("display_name") or admin_user.get("login_name") or "admin").strip()
        cursor = conn.execute(
            """
            INSERT INTO roster_import_batches(filename, imported_by, imported_count, skipped_count, created_at)
            VALUES(?, ?, 0, 0, ?)
            """,
            (str(file_name or "roster.csv")[:255], imported_by, now),
        )
        batch_id = cursor.lastrowid

        conn.execute("UPDATE villager_roster SET status = 'inactive' WHERE status = 'active'")
        conn.executemany(
            """
            INSERT INTO villager_roster(
                batch_id, full_name, id_last4, id_hash, group_name, status, created_at
            )
            VALUES(?, ?, ?, ?, ?, 'active', ?)
            """,
            [
                (
                    batch_id,
                    row["full_name"],
                    row["id_last4"],
                    row["id_hash"],
                    row["group_name"],
                    now,
                )
                for row in entries
            ],
        )
        conn.execute(
            """
            UPDATE roster_import_batches
            SET imported_count = ?, skipped_count = ?
            WHERE id = ?
            """,
            (len(entries), skipped_count, batch_id),
        )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user.get("id") or 0),
            action="admin_roster_import",
            target_type="villager_roster",
            target_id=str(batch_id),
            detail={
                "filename": str(file_name or "roster.csv")[:255],
                "imported_count": len(entries),
                "skipped_count": skipped_count,
            },
        )
        conn.commit()
        return {
            "batch_id": int(batch_id or 0),
            "filename": str(file_name or "roster.csv")[:255],
            "imported_count": len(entries),
            "skipped_count": skipped_count,
            "created_at": now,
        }
    finally:
        conn.close()


def _normalize_cadre_auth_status(value: str) -> str:
    text = str(value or "").strip().lower()
    if not text or text == "all":
        return ""
    if text not in CADRE_AUTH_STATUS_SET:
        raise ValueError("invalid_cadre_authorization_status")
    return text


def _normalize_cadre_role_level(value: str) -> str:
    text = normalize_cadre_role_level(value)
    if text not in CADRE_ROLE_LEVEL_SET:
        raise ValueError("invalid_cadre_role_level")
    return text


def _normalize_admin_role_key(value: str) -> str:
    text = str(value or "").strip().lower()
    if text not in ADMIN_ROLE_KEY_SET:
        raise ValueError("invalid_admin_role_key")
    return text


def _parse_detail_json(raw: Any) -> dict[str, Any]:
    payload = _parse_json_object(raw)
    return payload if isinstance(payload, dict) else {}


def _detail_preview(detail_json: Any) -> str:
    payload = _parse_detail_json(detail_json)
    if not payload:
        return "-"
    text = json.dumps(payload, ensure_ascii=False, separators=(", ", ": "))
    return text[:120] + ("..." if len(text) > 120 else "")


def _admin_actor_name(row: Any) -> str:
    item = dict(row) if not isinstance(row, dict) else row
    name = str(item.get("actor_display_name") or "").strip()
    if name:
        return name
    return "系统记录"


def _normalize_gov_application_status(value: str) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    if text not in GOV_APPLICATION_STATUS_SET:
        raise ValueError("invalid_gov_application_status")
    return text


def get_admin_nav_badges(token: str) -> dict[str, int]:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        request_queue_count = int(
            conn.execute(
                """
                SELECT COUNT(*) AS count
                FROM villager_requests
                WHERE status IN ('pending', 'processing')
                """
            ).fetchone()["count"]
            or 0
        )
        gov_queue_count = int(
            conn.execute(
                """
                SELECT COUNT(*) AS count
                FROM gov_service_applications
                WHERE status IN ('pending', 'processing')
                """
            ).fetchone()["count"]
            or 0
        )
        notice_count = int(
            conn.execute(
                """
                SELECT COUNT(*) AS count
                FROM notices
                """
            ).fetchone()["count"]
            or 0
        )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_nav_badges",
            target_type="dashboard",
            target_id="sidebar",
            detail={
                "requests": request_queue_count,
                "gov_applications": gov_queue_count,
                "notices": notice_count,
            },
        )
        conn.commit()
        return {
            "requests": request_queue_count,
            "gov_applications": gov_queue_count,
            "notices": notice_count,
        }
    finally:
        conn.close()


def list_admin_gov_services(token: str) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        _seed_default_gov_services_conn(conn)
        rows = conn.execute(
            """
            SELECT
                id, service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, detail_type, detail_payload,
                action_label, is_enabled, sort_order, updated_at, created_at
            FROM gov_services
            ORDER BY sort_order ASC, id ASC
            """
        ).fetchall()
        items = [
            _gov_service_to_dict(row, include_payload=False)
            for row in rows
            if str(row["service_key"] or "") not in HIDDEN_GOV_SERVICE_KEYS
        ]
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_gov_service_list",
            target_type="gov_service",
            target_id="all",
            detail={"count": len(items)},
        )
        conn.commit()
        return items
    finally:
        conn.close()


def get_admin_gov_service_detail(token: str, service_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        _seed_default_gov_services_conn(conn)
        row = conn.execute(
            """
            SELECT
                id, service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, detail_type, detail_payload,
                action_label, is_enabled, sort_order, updated_at, created_at
            FROM gov_services
            WHERE id = ?
            LIMIT 1
            """,
            (service_id,),
        ).fetchone()
        if row is None:
            raise ValueError("gov_service_not_found")
        item = _gov_service_to_dict(row, include_payload=True)
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_gov_service_detail",
            target_type="gov_service",
            target_id=str(service_id),
            detail={"service_key": str(item.get("service_key") or "")},
        )
        conn.commit()
        return item
    finally:
        conn.close()


def update_admin_gov_service(token: str, service_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    title = str(payload.get("title") or "").strip()
    subtitle = str(payload.get("subtitle") or "").strip()
    category = str(payload.get("category") or "government").strip() or "government"
    icon_emoji = str(payload.get("icon_emoji") or "").strip()
    theme_color = str(payload.get("theme_color") or "").strip()
    contact_name = str(payload.get("contact_name") or "").strip()
    contact_phone = re.sub(r"\D", "", str(payload.get("contact_phone") or "")).strip()[:20]
    intro = str(payload.get("intro") or "").strip()
    detail_type = str(payload.get("detail_type") or "info_apply").strip()
    action_label = str(payload.get("action_label") or "").strip() or "我要申请"
    sort_order = int(payload.get("sort_order") or 0)
    is_enabled = bool(payload.get("is_enabled", True))
    detail_payload = _sanitize_gov_detail_payload(payload.get("detail_payload"), strict_scene=True)

    if not title:
        raise ValueError("invalid_gov_service_payload")
    if detail_type not in VALID_GOV_DETAIL_TYPES:
        raise ValueError("invalid_gov_service_payload")
    if contact_phone and not re.fullmatch(r"\d{6,20}", contact_phone):
        raise ValueError("invalid_phone")

    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        _seed_default_gov_services_conn(conn)
        row = conn.execute(
            """
            SELECT id, service_key
            FROM gov_services
            WHERE id = ?
            LIMIT 1
            """,
            (service_id,),
        ).fetchone()
        if row is None:
            raise ValueError("gov_service_not_found")

        now = _now_iso()
        conn.execute(
            """
            UPDATE gov_services
            SET
                title = ?,
                subtitle = ?,
                category = ?,
                icon_emoji = ?,
                theme_color = ?,
                contact_name = ?,
                contact_phone = ?,
                intro = ?,
                detail_type = ?,
                detail_payload = ?,
                action_label = ?,
                is_enabled = ?,
                sort_order = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                title[:60],
                subtitle[:120],
                category[:40],
                icon_emoji[:20],
                theme_color[:20],
                contact_name[:50],
                contact_phone[:20],
                intro[:2000],
                detail_type[:40],
                _json_text(detail_payload),
                action_label[:30],
                1 if is_enabled else 0,
                max(0, min(sort_order, 1000)),
                now,
                service_id,
            ),
        )
        updated_row = conn.execute(
            """
            SELECT
                id, service_key, title, subtitle, category, icon_emoji, theme_color,
                contact_name, contact_phone, intro, detail_type, detail_payload,
                action_label, is_enabled, sort_order, updated_at, created_at
            FROM gov_services
            WHERE id = ?
            LIMIT 1
            """,
            (service_id,),
        ).fetchone()
        updated = _gov_service_to_dict(updated_row, include_payload=True)
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_gov_service_update",
            target_type="gov_service",
            target_id=str(service_id),
            detail={
                "service_key": str(updated.get("service_key") or str(row["service_key"] or "")),
                "title": str(updated.get("title") or ""),
            },
        )
        conn.commit()
        return updated
    finally:
        conn.close()


def list_admin_requests(
    token: str,
    *,
    status: str = "",
    request_type: str = "",
    keyword: str = "",
) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    status_value = _normalize_request_status(status)
    type_value = _normalize_request_type(request_type)
    keyword_text = str(keyword or "").strip()

    conn = get_connection()
    try:
        query = """
            SELECT
                r.id,
                r.user_id,
                r.request_type,
                r.title,
                r.content,
                r.contact_phone,
                r.reporter_name,
                r.photo_urls,
                r.reply_photo_urls,
                r.status,
                r.reply_text,
                r.handled_by,
                r.created_at,
                r.updated_at,
                r.handled_at,
                COALESCE(u.village_group, '') AS village_group
            FROM villager_requests r
            LEFT JOIN users u ON u.id = r.user_id
            WHERE 1 = 1
        """
        params: list[Any] = []
        if status_value:
            query += " AND r.status = ?"
            params.append(status_value)
        if type_value:
            query += " AND r.request_type = ?"
            params.append(type_value)
        if keyword_text:
            like_text = f"%{keyword_text}%"
            query += " AND (r.title LIKE ? OR r.content LIKE ? OR r.reporter_name LIKE ? OR COALESCE(u.village_group, '') LIKE ?)"
            params.extend([like_text, like_text, like_text, like_text])
        query += " ORDER BY r.id DESC"

        rows = conn.execute(query, tuple(params)).fetchall()
        if len(rows) >= 50:
            security_logger.warning(
                "security_alert event=admin_bulk_request_view admin_user_id=%s count=%s",
                int(admin_user["id"] or 0),
                len(rows),
            )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_request_list",
            target_type="villager_request",
            target_id="all",
            detail={
                "status": status_value or "all",
                "request_type": type_value or "all",
                "keyword": keyword_text,
                "count": len(rows),
            },
        )
        conn.commit()
        return [_villager_request_to_dict(row) for row in rows]
    finally:
        conn.close()


def get_admin_request_detail(token: str, request_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT
                r.id,
                r.user_id,
                r.request_type,
                r.title,
                r.content,
                r.contact_phone,
                r.reporter_name,
                r.photo_urls,
                r.reply_photo_urls,
                r.status,
                r.reply_text,
                r.handled_by,
                r.created_at,
                r.updated_at,
                r.handled_at,
                COALESCE(u.village_group, '') AS village_group
            FROM villager_requests r
            LEFT JOIN users u ON u.id = r.user_id
            WHERE r.id = ?
            LIMIT 1
            """,
            (int(request_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("request_not_found")
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_request_detail",
            target_type="villager_request",
            target_id=str(request_id),
            detail={},
        )
        conn.commit()
        return _villager_request_to_dict(row)
    finally:
        conn.close()


def _extract_request_material_filename(url: str) -> str:
    path = unquote(urlsplit(str(url or "").strip()).path or "")
    parts = [part for part in path.split("/") if part]
    for index, part in enumerate(parts[:-1]):
        if part == "request":
            filename = str(parts[index + 1] or "").strip()
            if re.fullmatch(r"[A-Za-z0-9_.-]{1,160}", filename):
                return filename
    return ""


def _read_request_material_bytes(url: str) -> tuple[bytes, str]:
    filename = _extract_request_material_filename(url)
    if not filename:
        raise FileNotFoundError("request_material_file_not_found")
    record = get_stored_file_record("request", filename)
    if record is not None:
        if str(record.get("storage_provider") or "local").strip().lower() == "s3":
            signed_url = build_download_url(record=record, expires_in_seconds=300)
            with urlopen(signed_url, timeout=30) as response:
                return response.read(), filename
        file_path = resolve_local_file_path(record).resolve()
    else:
        file_path = (get_backend_upload_dir() / "request" / filename).resolve()
        scene_root = (get_backend_upload_dir() / "request").resolve()
        if not str(file_path).startswith(str(scene_root)):
            raise FileNotFoundError("request_material_file_not_found")
    if not file_path.is_file():
        raise FileNotFoundError("request_material_file_not_found")
    return file_path.read_bytes(), filename


def _safe_archive_name(value: str, fallback: str) -> str:
    text = re.sub(r"[\\/:*?\"<>|\r\n]+", "_", str(value or "").strip()).strip(" ._")
    return text[:60] or fallback


def build_admin_request_materials_zip(token: str, request_id: int) -> tuple[bytes, str]:
    detail = get_admin_request_detail(token, request_id)
    request_type = str(detail.get("request_type") or "").strip().lower()
    if request_type not in {"dibao", "tekun"}:
        raise ValueError("request_materials_not_supported")

    photo_urls = [str(item or "").strip() for item in detail.get("photo_urls") or [] if str(item or "").strip()]
    if not photo_urls:
        raise ValueError("request_materials_empty")

    try:
        content_payload = json.loads(str(detail.get("content") or "{}"))
    except (TypeError, ValueError, json.JSONDecodeError):
        content_payload = {}
    manifest = content_payload.get("materials") if isinstance(content_payload, dict) else []
    manifest = manifest if isinstance(manifest, list) else []

    grouped_urls: list[tuple[str, list[str]]] = []
    cursor = 0
    for material_index, material in enumerate(manifest, start=1):
        if not isinstance(material, dict):
            continue
        try:
            count = max(0, int(material.get("count") or 0))
        except (TypeError, ValueError):
            count = 0
        name = _safe_archive_name(material.get("name"), f"材料{material_index}")
        urls = photo_urls[cursor : cursor + count]
        cursor += count
        if urls:
            grouped_urls.append((name, urls))
    if cursor < len(photo_urls):
        grouped_urls.append(("其他材料", photo_urls[cursor:]))
    if not grouped_urls:
        grouped_urls.append(("申请材料", photo_urls))

    archive_buffer = io.BytesIO()
    written_count = 0
    missing_files: list[str] = []
    with zipfile.ZipFile(archive_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for group_index, (material_name, urls) in enumerate(grouped_urls, start=1):
            for file_index, url in enumerate(urls, start=1):
                try:
                    file_bytes, source_filename = _read_request_material_bytes(url)
                except (FileNotFoundError, OSError, RuntimeError):
                    missing_files.append(f"{material_name} 第{file_index}张")
                    continue
                suffix = str(source_filename.rsplit(".", 1)[-1] if "." in source_filename else "").lower()
                if not re.fullmatch(r"[a-z0-9]{1,8}", suffix):
                    content_type = mimetypes.guess_type(source_filename)[0] or "image/jpeg"
                    suffix = mimetypes.guess_extension(content_type) or ".jpg"
                    suffix = suffix.lstrip(".")
                archive_name = f"{group_index:02d}_{material_name}_{file_index:02d}.{suffix}"
                archive.writestr(archive_name, file_bytes)
                written_count += 1
        if missing_files:
            archive.writestr("下载说明.txt", "以下材料文件未能找到：\n" + "\n".join(missing_files))

    if written_count <= 0:
        raise ValueError("request_materials_empty")
    type_label = "低保申请" if request_type == "dibao" else "特困认证"
    applicant = _safe_archive_name(detail.get("reporter_name"), "村民")
    filename = f"{type_label}_{applicant}_{int(request_id)}_申请材料.zip"
    return archive_buffer.getvalue(), filename


def update_admin_request(
    token: str,
    request_id: int,
    *,
    status: str,
    reply_text: str = "",
    reply_photo_urls: list[str] | None = None,
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)

    next_status = _normalize_request_status(status)
    reply_text_value = str(reply_text or "").strip()
    reply_photo_urls_value = _normalize_photo_urls(
        reply_photo_urls,
        allowed_scenes=SCENES_REQUEST_REPLY,
        allow_external=False,
        strict_scene=True,
    )

    if next_status in {"processing", "resolved", "closed"} and not reply_text_value and not reply_photo_urls_value:
        raise ValueError("request_handle_requires_feedback")

    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id, status, handled_at FROM villager_requests WHERE id = ? LIMIT 1",
            (int(request_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("request_not_found")
        _require_status_transition(
            str(row["status"] or ""),
            next_status,
            REQUEST_STATUS_TRANSITIONS,
            "invalid_request_status_transition",
        )

        now = _now_iso()
        handled_by = (
            str(admin_user.get("display_name") or "").strip()
            or str(admin_user.get("real_name") or "").strip()
            or str(admin_user.get("login_name") or "").strip()
        )
        handled_at = (
            now
            if next_status in {"processing", "resolved", "closed"} or reply_text_value or reply_photo_urls_value
            else str(row["handled_at"] or "")
        )
        conn.execute(
            """
            UPDATE villager_requests
            SET
                status = ?,
                reply_text = ?,
                reply_photo_urls = ?,
                handled_by = ?,
                updated_at = ?,
                handled_at = ?
            WHERE id = ?
            """,
            (
                next_status,
                reply_text_value,
                json.dumps(reply_photo_urls_value, ensure_ascii=False),
                handled_by,
                now,
                handled_at,
                int(request_id or 0),
            ),
        )
        updated = conn.execute(
            """
            SELECT
                r.id,
                r.user_id,
                r.request_type,
                r.title,
                r.content,
                r.contact_phone,
                r.reporter_name,
                r.photo_urls,
                r.reply_photo_urls,
                r.status,
                r.reply_text,
                r.handled_by,
                r.created_at,
                r.updated_at,
                r.handled_at,
                COALESCE(u.village_group, '') AS village_group
            FROM villager_requests r
            LEFT JOIN users u ON u.id = r.user_id
            WHERE r.id = ?
            LIMIT 1
            """,
            (int(request_id or 0),),
        ).fetchone()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_request_update",
            target_type="villager_request",
            target_id=str(request_id),
            detail={
                "status": next_status,
                "reply_length": len(reply_text_value),
                "reply_photo_count": len(reply_photo_urls_value),
            },
        )
        conn.commit()
        return _villager_request_to_dict(updated)
    finally:
        conn.close()


def list_admin_gov_applications(
    token: str,
    *,
    status: str = "",
    service_key: str = "",
    keyword: str = "",
    village_group: str = "",
) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    status_value = _normalize_gov_application_status(status)
    service_key_value = str(service_key or "").strip().lower()
    keyword_text = str(keyword or "").strip()
    village_group_value = str(village_group or "").strip()

    conn = get_connection()
    try:
        query = """
            SELECT
                g.id, g.service_id, g.service_key, g.service_title, g.user_id,
                g.applicant_name, g.applicant_phone, g.selected_option, g.note,
                g.extra_fields, g.reply_photo_urls, g.status, g.cadre_reply, g.handled_by,
                g.created_at, g.updated_at, g.handled_at,
                COALESCE(u.village_group, '') AS village_group
            FROM gov_service_applications g
            LEFT JOIN users u ON u.id = g.user_id
            WHERE 1 = 1
        """
        params: list[Any] = []
        if status_value:
            query += " AND g.status = ?"
            params.append(status_value)
        if service_key_value and service_key_value != "all":
            query += " AND g.service_key = ?"
            params.append(service_key_value)
        if village_group_value and village_group_value != "all":
            if village_group_value == "__ungrouped__":
                query += " AND COALESCE(u.village_group, '') = ''"
            else:
                query += " AND COALESCE(u.village_group, '') = ?"
                params.append(village_group_value)
        if keyword_text:
            like_text = f"%{keyword_text}%"
            query += " AND (g.service_title LIKE ? OR g.applicant_name LIKE ? OR g.selected_option LIKE ? OR g.note LIKE ? OR COALESCE(u.village_group, '') LIKE ?)"
            params.extend([like_text, like_text, like_text, like_text, like_text])
        query += " ORDER BY g.id DESC"

        rows = conn.execute(query, tuple(params)).fetchall()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_gov_application_list",
            target_type="gov_application",
            target_id="all",
            detail={
                "status": status_value or "all",
                "service_key": service_key_value or "all",
                "village_group": village_group_value or "all",
                "keyword": keyword_text,
                "count": len(rows),
            },
        )
        conn.commit()
        return [_gov_application_to_dict(row) for row in rows]
    finally:
        conn.close()


def get_admin_gov_application_detail(token: str, application_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT
                g.id, g.service_id, g.service_key, g.service_title, g.user_id,
                g.applicant_name, g.applicant_phone, g.selected_option, g.note,
                g.extra_fields, g.reply_photo_urls, g.status, g.cadre_reply, g.handled_by,
                g.created_at, g.updated_at, g.handled_at,
                COALESCE(u.village_group, '') AS village_group
            FROM gov_service_applications g
            LEFT JOIN users u ON u.id = g.user_id
            WHERE g.id = ?
            LIMIT 1
            """,
            (int(application_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("gov_application_not_found")
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_gov_application_detail",
            target_type="gov_application",
            target_id=str(application_id),
            detail={},
        )
        conn.commit()
        return _gov_application_to_dict(row)
    finally:
        conn.close()


def update_admin_gov_application(
    token: str,
    application_id: int,
    *,
    status: str,
    cadre_reply: str = "",
    reply_photo_urls: list[str] | None = None,
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)

    next_status = _normalize_gov_application_status(status)
    cadre_reply_value = str(cadre_reply or "").strip()
    reply_photo_urls_value = _normalize_photo_urls(
        reply_photo_urls,
        allowed_scenes=SCENES_GOV_REPLY,
        allow_external=False,
        strict_scene=True,
    )

    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT id, status, handled_at, service_key
            FROM gov_service_applications
            WHERE id = ?
            LIMIT 1
            """,
            (int(application_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("gov_application_not_found")
        _require_status_transition(
            str(row["status"] or ""),
            next_status,
            GOV_APPLICATION_STATUS_TRANSITIONS,
            "invalid_gov_status_transition",
        )
        if next_status in {"processing", "resolved", "rejected", "closed"} and not cadre_reply_value and not reply_photo_urls_value:
            raise ValueError("gov_handle_requires_feedback")
        if str(row["service_key"] or "") == "certificate" and next_status in {"resolved", "closed"} and not reply_photo_urls_value:
            raise ValueError("certificate_reply_photo_required")

        now = _now_iso()
        handled_by = (
            str(admin_user.get("display_name") or "").strip()
            or str(admin_user.get("real_name") or "").strip()
            or str(admin_user.get("login_name") or "").strip()
        )
        handled_at = (
            now
            if next_status in {"processing", "resolved", "rejected", "closed"} or cadre_reply_value or reply_photo_urls_value
            else str(row["handled_at"] or "")
        )
        conn.execute(
            """
            UPDATE gov_service_applications
            SET
                status = ?,
                cadre_reply = ?,
                reply_photo_urls = ?,
                handled_by = ?,
                updated_at = ?,
                handled_at = ?
            WHERE id = ?
            """,
            (
                next_status,
                cadre_reply_value[:2000],
                json.dumps(reply_photo_urls_value, ensure_ascii=False),
                handled_by[:50],
                now,
                handled_at,
                int(application_id or 0),
            ),
        )
        updated = conn.execute(
            """
            SELECT
                g.id, g.service_id, g.service_key, g.service_title, g.user_id,
                g.applicant_name, g.applicant_phone, g.selected_option, g.note,
                g.extra_fields, g.reply_photo_urls, g.status, g.cadre_reply, g.handled_by,
                g.created_at, g.updated_at, g.handled_at,
                COALESCE(u.village_group, '') AS village_group
            FROM gov_service_applications g
            LEFT JOIN users u ON u.id = g.user_id
            WHERE g.id = ?
            LIMIT 1
            """,
            (int(application_id or 0),),
        ).fetchone()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_gov_application_update",
            target_type="gov_application",
            target_id=str(application_id),
            detail={
                "status": next_status,
                "reply_length": len(cadre_reply_value),
                "reply_photo_count": len(reply_photo_urls_value),
            },
        )
        conn.commit()
        return _gov_application_to_dict(updated)
    finally:
        conn.close()


def list_admin_notices(token: str, *, keyword: str = "", home_top_only: bool = False) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_notices()
    keyword_text = str(keyword or "").strip().lower()
    items = rows
    if home_top_only:
        items = [item for item in items if bool(item.get("is_home_top"))]
    if keyword_text:
        items = [
            item
            for item in items
            if keyword_text in str(item.get("title") or "").lower()
            or keyword_text in str(item.get("content") or "").lower()
            or keyword_text in str(item.get("publisher") or "").lower()
        ]
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_notice_list",
            target_type="notice",
            target_id="all",
            detail={
                "keyword": keyword_text,
                "home_top_only": bool(home_top_only),
                "count": len(items),
            },
        )
        conn.commit()
    finally:
        conn.close()
    return items


def get_admin_notice_detail(token: str, notice_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    detail = get_notice_detail(notice_id)
    payload = {
        "id": int(detail["id"] or 0),
        "title": str(detail.get("title") or ""),
        "content": str(detail.get("content") or ""),
        "publisher": str(detail.get("publisher") or ""),
        "photo_urls": list(detail.get("photo_urls") or []),
        "is_home_top": bool(detail.get("is_home_top")),
        "created_at": str(detail.get("created_at") or ""),
    }
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_notice_detail",
            target_type="notice",
            target_id=str(notice_id),
            detail={},
        )
        conn.commit()
    finally:
        conn.close()
    return payload


def create_admin_notice(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    created = create_notice(payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_notice_create",
            target_type="notice",
            target_id=str(created.get("id") or ""),
            detail={
                "title": str(created.get("title") or ""),
                "is_home_top": bool(created.get("is_home_top")),
            },
        )
        conn.commit()
    finally:
        conn.close()
    return created


def update_admin_notice(token: str, notice_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id FROM notices WHERE id = ? LIMIT 1",
            (int(notice_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("notice_not_found")
        photo_urls = _normalize_photo_urls(
            payload.get("photo_urls"),
            allowed_scenes=SCENES_NOTICE,
            allow_external=False,
            strict_scene=True,
        )
        is_home_top = 1 if bool(payload.get("is_home_top")) else 0
        conn.execute(
            """
            UPDATE notices
            SET title = ?, content = ?, publisher = ?, photo_urls = ?, is_home_top = ?
            WHERE id = ?
            """,
            (
                payload["title"],
                payload["content"],
                payload["publisher"],
                _json_text(photo_urls),
                is_home_top,
                int(notice_id or 0),
            ),
        )
        updated = conn.execute(
            "SELECT id, title, content, publisher, photo_urls, is_home_top, created_at FROM notices WHERE id = ? LIMIT 1",
            (int(notice_id or 0),),
        ).fetchone()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_notice_update",
            target_type="notice",
            target_id=str(notice_id),
            detail={
                "title": str(payload.get("title") or ""),
                "is_home_top": bool(is_home_top),
            },
        )
        conn.commit()
        return _notice_to_dict(updated)
    finally:
        conn.close()


def delete_admin_notice(token: str, notice_id: int) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id, title FROM notices WHERE id = ? LIMIT 1",
            (int(notice_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("notice_not_found")
        conn.execute("DELETE FROM notices WHERE id = ?", (int(notice_id or 0),))
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_notice_delete",
            target_type="notice",
            target_id=str(notice_id),
            detail={"title": str(row["title"] or "")},
        )
        conn.commit()
        return {"status": "deleted", "message": "通知已删除"}
    finally:
        conn.close()


def upload_admin_image(
    token: str,
    *,
    scene: str,
    filename: str,
    content_type: str,
    file_bytes: bytes,
    public_origin: str = "",
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    record = store_uploaded_file_record(
        scene=scene,
        filename=filename,
        content_type=content_type,
        file_bytes=file_bytes,
        uploaded_by_token=None,
    )
    canonical_url = f"{str(public_origin or '').rstrip('/')}/api/v1/uploads/access/{scene}/{filename}"
    signed_url = build_upload_access_url(canonical_url)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_upload_image",
            target_type="stored_file",
            target_id=str(record.get("id") or f"{scene}/{filename}"),
            detail={"scene": scene, "filename": filename},
        )
        conn.commit()
    finally:
        conn.close()
    return {
        "url": signed_url,
        "filename": filename,
        "size": len(file_bytes),
        "content_type": content_type,
    }


def list_admin_metrics(token: str) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_village_metrics()
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_metric_list",
            target_type="village_metric",
            target_id="all",
            detail={"count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def update_admin_metrics(token: str, payload: dict[str, Any]) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    items = payload.get("items") or []
    if not isinstance(items, list) or not items:
        raise ValueError("invalid_metric_payload")

    conn = get_connection()
    try:
        _seed_default_metrics_conn(conn)
        now = _now_iso()
        updated_by = (
            str(admin_user.get("display_name") or "").strip()
            or str(admin_user.get("real_name") or "").strip()
            or str(admin_user.get("login_name") or "").strip()
        )
        normalized: list[dict[str, Any]] = []
        used_keys: set[str] = set()
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                raise ValueError("invalid_metric_payload")
            metric_label = str(item.get("metric_label") or "").strip()
            metric_value = str(item.get("metric_value") or "").strip()
            metric_key = _sanitize_metric_key(str(item.get("metric_key") or ""))
            if not metric_key:
                metric_key = f"metric_{index}"
            while metric_key in used_keys:
                metric_key = f"{metric_key}_{index}"
            if not metric_label or not metric_value:
                raise ValueError("invalid_metric_payload")
            used_keys.add(metric_key)
            normalized.append(
                {
                    "metric_key": metric_key,
                    "metric_label": metric_label[:40],
                    "metric_value": metric_value[:40],
                    "sort_order": index,
                }
            )

        conn.execute("DELETE FROM village_metrics")
        conn.executemany(
            """
            INSERT INTO village_metrics(
                metric_key, metric_label, metric_value, sort_order, updated_by, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    item["metric_key"],
                    item["metric_label"],
                    item["metric_value"],
                    item["sort_order"],
                    updated_by,
                    now,
                    now,
                )
                for item in normalized
            ],
        )
        rows = conn.execute(
            """
            SELECT metric_key, metric_label, metric_value, sort_order, updated_at
            FROM village_metrics
            ORDER BY sort_order ASC, id ASC
            """
        ).fetchall()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_metric_update",
            target_type="village_metric",
            target_id="all",
            detail={"count": len(normalized)},
        )
        conn.commit()
        output: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            if item.get("metric_key") == "groups":
                item["metric_label"] = "村组"
            output.append(item)
        return output
    finally:
        conn.close()


def get_admin_open_profile(token: str) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    profile = get_village_open_profile()
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_open_profile_view",
            target_type="open_profile",
            target_id="1",
            detail={"has_photo": bool(profile.get("village_photo_url"))},
        )
        conn.commit()
    finally:
        conn.close()
    return profile


def update_admin_open_profile(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    updated_by = (
        str(admin_user.get("display_name") or "").strip()
        or str(admin_user.get("real_name") or "").strip()
        or str(admin_user.get("login_name") or "").strip()
    )
    profile = update_village_open_profile(updated_by, payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_open_profile_update",
            target_type="open_profile",
            target_id="1",
            detail={
                "has_photo": bool(profile.get("village_photo_url")),
                "intro_length": len(str(profile.get("village_intro") or "")),
            },
        )
        conn.commit()
    finally:
        conn.close()
    return profile


def _admin_cadre_authorization_to_dict(row: Any) -> dict[str, Any]:
    return {
        "id": int(row["id"] or 0),
        "real_name": str(row["real_name"] or ""),
        "office_title": str(row["office_title"] or ""),
        "user_id": int(row["user_id"] or 0),
        "openid_bound": bool(str(row["openid"] or "").strip()),
        "unionid_bound": bool(str(row["unionid"] or "").strip()),
        "role_level": normalize_cadre_role_level(str(row["role_level"] or "")),
        "can_export_full_id": False,
        "is_default": str(row["real_name"] or "").strip() in DEFAULT_CADRE_MANAGER_NAMES,
        "status": str(row["status"] or "active"),
        "last_bound_at": str(row["last_bound_at"] or ""),
        "admin_account_exists": bool(row["admin_user_id"] or 0),
        "admin_login_name": str(row["admin_login_name"] or ""),
        "admin_display_name": str(row["admin_display_name"] or ""),
        "admin_role_key": str(row["admin_role_key"] or ""),
        "admin_status": str(row["admin_status"] or ""),
        "admin_locked_until": str(row["admin_locked_until"] or ""),
        "admin_last_login_at": str(row["admin_last_login_at"] or ""),
    }


def _fetch_admin_cadre_authorization_conn(conn: Any, authorization_id: int) -> dict[str, Any]:
    row = conn.execute(
        """
        SELECT
            c.id, c.real_name, c.office_title, c.user_id, c.openid, c.unionid, c.role_level, c.can_export_full_id,
            c.status, c.last_bound_at,
            a.id AS admin_user_id,
            a.login_name AS admin_login_name,
            a.display_name AS admin_display_name,
            a.role_key AS admin_role_key,
            a.status AS admin_status,
            a.locked_until AS admin_locked_until,
            a.last_login_at AS admin_last_login_at
        FROM cadre_authorizations c
        LEFT JOIN admin_users a ON a.cadre_authorization_id = c.id
        WHERE c.id = ?
        LIMIT 1
        """,
        (int(authorization_id or 0),),
    ).fetchone()
    if row is None:
        raise ValueError("cadre_authorization_not_found")
    return _admin_cadre_authorization_to_dict(row)


def list_admin_cadre_authorizations(
    token: str,
    *,
    status: str = "",
    keyword: str = "",
) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    status_value = _normalize_cadre_auth_status(status)
    keyword_text = str(keyword or "").strip()
    conn = get_connection()
    try:
        query = """
            SELECT
                c.id, c.real_name, c.office_title, c.user_id, c.openid, c.unionid, c.role_level, c.can_export_full_id,
                c.status, c.last_bound_at,
                a.id AS admin_user_id,
                a.login_name AS admin_login_name,
                a.display_name AS admin_display_name,
                a.role_key AS admin_role_key,
                a.status AS admin_status,
                a.locked_until AS admin_locked_until,
                a.last_login_at AS admin_last_login_at
            FROM cadre_authorizations c
            LEFT JOIN admin_users a ON a.cadre_authorization_id = c.id
            WHERE 1 = 1
        """
        params: list[Any] = []
        if status_value:
            query += " AND c.status = ?"
            params.append(status_value)
        if keyword_text:
            like_text = f"%{keyword_text}%"
            query += " AND (c.real_name LIKE ? OR a.login_name LIKE ? OR a.display_name LIKE ?)"
            params.extend([like_text, like_text, like_text])
        query += """
            ORDER BY
                CASE WHEN c.role_level = 'super_admin' THEN 0 ELSE 1 END ASC,
                c.id ASC
        """
        rows = conn.execute(query, tuple(params)).fetchall()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_cadre_authorization_list",
            target_type="cadre_authorization",
            target_id="all",
            detail={"status": status_value or "all", "keyword": keyword_text, "count": len(rows)},
        )
        conn.commit()
        return [_admin_cadre_authorization_to_dict(row) for row in rows]
    finally:
        conn.close()


def update_admin_cadre_authorization(
    token: str,
    authorization_id: int,
    *,
    office_title: str = "",
    role_level: str,
    can_export_full_id: bool,
    status: str,
    admin_role_key: str,
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    role_level_value = _normalize_cadre_role_level(role_level)
    office_title_value = str(office_title or "").strip()[:60]
    status_value = _normalize_cadre_auth_status(status)
    admin_role_value = _normalize_admin_role_key(admin_role_key)
    if role_level_value != "super_admin" and admin_role_value == "super_admin":
        raise ValueError("invalid_admin_role_key")
    export_full_id = 0

    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id FROM cadre_authorizations WHERE id = ? LIMIT 1",
            (int(authorization_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("cadre_authorization_not_found")
        now = _now_iso()
        conn.execute(
            """
            UPDATE cadre_authorizations
            SET office_title = ?, role_level = ?, can_export_full_id = ?, status = ?, updated_at = ?
            WHERE id = ?
            """,
            (office_title_value, role_level_value, export_full_id, status_value, now, int(authorization_id or 0)),
        )
        admin_row = conn.execute(
            "SELECT id FROM admin_users WHERE cadre_authorization_id = ? LIMIT 1",
            (int(authorization_id or 0),),
        ).fetchone()
        if admin_row is not None:
            next_admin_status = "active" if status_value == "active" else "disabled"
            conn.execute(
                """
                UPDATE admin_users
                SET role_key = ?, status = ?, updated_at = ?
                WHERE id = ?
                """,
                (admin_role_value, next_admin_status, now, int(admin_row["id"] or 0)),
            )
            if next_admin_status != "active":
                conn.execute(
                    """
                    UPDATE admin_sessions
                    SET revoked_at = ?
                    WHERE admin_user_id = ? AND revoked_at = ''
                    """,
                    (now, int(admin_row["id"] or 0)),
                )
        result = _fetch_admin_cadre_authorization_conn(conn, int(authorization_id or 0))
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_cadre_authorization_update",
            target_type="cadre_authorization",
            target_id=str(authorization_id),
            detail={
                "office_title": office_title_value,
                "role_level": role_level_value,
                "can_export_full_id": bool(export_full_id),
                "status": status_value,
                "admin_role_key": admin_role_value,
            },
        )
        conn.commit()
        return result
    finally:
        conn.close()


def create_admin_cadre_authorization(
    token: str,
    *,
    real_name: str,
    office_title: str = "",
    role_level: str = "cadre",
    can_export_full_id: bool = False,
    status: str = "active",
    create_account: bool = False,
    login_name: str = "",
    password: str = "",
    admin_role_key: str = "cadre",
    force_reset: bool = True,
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    real_name_value = str(real_name or "").strip()[:50]
    if not real_name_value:
        raise ValueError("invalid_cadre_real_name")
    office_title_value = str(office_title or "").strip()[:60]
    role_level_value = _normalize_cadre_role_level(role_level)
    status_value = _normalize_cadre_auth_status(status) or "active"
    export_full_id = 0
    if create_account and status_value != "active":
        raise ValueError("admin_authorization_disabled")

    conn = get_connection()
    try:
        exists = conn.execute(
            "SELECT id FROM cadre_authorizations WHERE real_name = ? LIMIT 1",
            (real_name_value,),
        ).fetchone()
        if exists is not None:
            raise ValueError("cadre_authorization_exists")
        now = _now_iso()
        cursor = conn.execute(
            """
            INSERT INTO cadre_authorizations(
                real_name, office_title, role_level, can_export_full_id, status, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?)
            """,
            (real_name_value, office_title_value, role_level_value, export_full_id, status_value, now, now),
        )
        new_id = int(cursor.lastrowid or 0)
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_cadre_authorization_create",
            target_type="cadre_authorization",
            target_id=str(new_id),
            detail={
                "real_name": real_name_value,
                "office_title": office_title_value,
                "role_level": role_level_value,
                "can_export_full_id": bool(export_full_id),
                "status": status_value,
                "create_account": bool(create_account),
            },
        )
        conn.commit()
    finally:
        conn.close()

    if create_account:
        # 复用既有账号创建流程（含密码校验、登录名查重、角色校验与审计）
        return create_admin_account_for_authorization(
            token,
            new_id,
            login_name=login_name,
            password=password,
            display_name=real_name_value,
            role_key=admin_role_key,
            force_reset=force_reset,
        )

    conn = get_connection()
    try:
        return _fetch_admin_cadre_authorization_conn(conn, new_id)
    finally:
        conn.close()


def delete_admin_cadre_authorization(token: str, authorization_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id, real_name FROM cadre_authorizations WHERE id = ? LIMIT 1",
            (int(authorization_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("cadre_authorization_not_found")
        real_name_value = str(row["real_name"] or "").strip()
        # 6 个默认干部是种子名单，启动时会自动补回，删除无意义——直接拦截
        if real_name_value in DEFAULT_CADRE_MANAGER_NAMES:
            raise ValueError("cadre_authorization_default_protected")
        admin_row = conn.execute(
            "SELECT id FROM admin_users WHERE cadre_authorization_id = ? LIMIT 1",
            (int(authorization_id or 0),),
        ).fetchone()
        if admin_row is not None:
            raise ValueError("cadre_authorization_has_account")
        conn.execute(
            "DELETE FROM cadre_authorizations WHERE id = ?",
            (int(authorization_id or 0),),
        )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_cadre_authorization_delete",
            target_type="cadre_authorization",
            target_id=str(authorization_id),
            detail={"real_name": real_name_value},
        )
        conn.commit()
        return {"status": "deleted", "message": "干部已删除"}
    finally:
        conn.close()


def create_admin_account_for_authorization(
    token: str,
    authorization_id: int,
    *,
    login_name: str,
    password: str,
    display_name: str = "",
    role_key: str = "cadre",
    force_reset: bool = False,
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    conn = get_connection()
    try:
        auth_row = conn.execute(
            "SELECT id, role_level, status FROM cadre_authorizations WHERE id = ? LIMIT 1",
            (int(authorization_id or 0),),
        ).fetchone()
        if auth_row is None:
            raise ValueError("cadre_authorization_not_found")
        if str(auth_row["status"] or "").strip() != "active":
            raise ValueError("admin_authorization_disabled")
        role_key_value = _normalize_admin_role_key(role_key)
        if normalize_cadre_role_level(str(auth_row["role_level"] or "")) != "super_admin" and role_key_value == "super_admin":
            raise ValueError("invalid_admin_role_key")
        created = bootstrap_admin_user(
            login_name=login_name,
            password=password,
            authorization_id=int(authorization_id or 0),
            display_name=display_name,
            role_key=role_key_value,
            force_reset=force_reset,
        )
        fresh = _fetch_admin_cadre_authorization_conn(conn, int(authorization_id or 0))
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_cadre_account_create",
            target_type="admin_user",
            target_id=str(created.get("id") or ""),
            detail={"authorization_id": authorization_id, "login_name": created.get("login_name", "")},
        )
        conn.commit()
        return fresh
    finally:
        conn.close()


def reset_admin_account_for_authorization(
    token: str,
    authorization_id: int,
    *,
    new_password: str,
    force_reset: bool = True,
) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    updated = reset_admin_account_password(
        authorization_id=int(authorization_id or 0),
        new_password=new_password,
        force_reset=force_reset,
    )
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_cadre_account_reset_password",
            target_type="admin_user",
            target_id=str(updated.get("id") or ""),
            detail={"authorization_id": authorization_id, "login_name": updated.get("login_name", "")},
        )
        conn.commit()
    finally:
        conn.close()
    return {"status": "ok", "message": "后台账号密码已重置"}


def list_admin_audit_logs(
    token: str,
    *,
    actor_type: str = "",
    action: str = "",
    target_type: str = "",
    keyword: str = "",
    limit: int = 200,
) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    actor_type_value = str(actor_type or "").strip().lower()
    if actor_type_value not in {"", "villager", "admin"}:
        raise ValueError("invalid_audit_actor_type")
    action_value = str(action or "").strip().lower()
    target_type_value = str(target_type or "").strip().lower()
    keyword_text = str(keyword or "").strip()
    limit_value = max(1, min(int(limit or 200), 500))

    conn = get_connection()
    try:
        query = """
            SELECT
                l.id, l.actor_type, l.actor_user_id, l.actor_admin_user_id, l.action, l.target_type, l.target_id,
                l.detail_json, l.ip, l.user_agent, l.created_at,
                COALESCE(au.display_name, au.real_name, u.real_name, u.nickname, '') AS actor_display_name
            FROM audit_logs l
            LEFT JOIN users u ON u.id = l.actor_user_id
            LEFT JOIN admin_users au ON au.id = l.actor_admin_user_id
            WHERE 1 = 1
        """
        params: list[Any] = []
        if actor_type_value:
            query += " AND l.actor_type = ?"
            params.append(actor_type_value)
        if action_value:
            query += " AND l.action = ?"
            params.append(action_value)
        if target_type_value:
            query += " AND l.target_type = ?"
            params.append(target_type_value)
        if keyword_text:
            like_text = f"%{keyword_text}%"
            query += " AND (l.action LIKE ? OR l.target_type LIKE ? OR l.target_id LIKE ? OR l.detail_json LIKE ? OR COALESCE(au.display_name, au.real_name, u.real_name, u.nickname, '') LIKE ?)"
            params.extend([like_text, like_text, like_text, like_text, like_text])
        query += " ORDER BY l.id DESC LIMIT ?"
        params.append(limit_value)
        rows = conn.execute(query, tuple(params)).fetchall()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_audit_log_list",
            target_type="audit_log",
            target_id="all",
            detail={
                "actor_type": actor_type_value or "all",
                "action": action_value or "all",
                "target_type": target_type_value or "all",
                "keyword": keyword_text,
                "limit": limit_value,
                "count": len(rows),
            },
        )
        conn.commit()
        return [
            {
                "id": int(row["id"] or 0),
                "actor_type": str(row["actor_type"] or ""),
                "actor_user_id": int(row["actor_user_id"] or 0) or None,
                "actor_admin_user_id": int(row["actor_admin_user_id"] or 0) or None,
                "actor_display_name": _admin_actor_name(row),
                "action": str(row["action"] or ""),
                "target_type": str(row["target_type"] or ""),
                "target_id": str(row["target_id"] or ""),
                "detail_preview": _detail_preview(row["detail_json"]),
                "ip": str(row["ip"] or ""),
                "user_agent": str(row["user_agent"] or ""),
                "created_at": str(row["created_at"] or ""),
            }
            for row in rows
        ]
    finally:
        conn.close()


def get_admin_audit_log_detail(token: str, log_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT
                l.id, l.actor_type, l.actor_user_id, l.actor_admin_user_id, l.action, l.target_type, l.target_id,
                l.detail_json, l.ip, l.user_agent, l.created_at,
                COALESCE(au.display_name, au.real_name, u.real_name, u.nickname, '') AS actor_display_name
            FROM audit_logs l
            LEFT JOIN users u ON u.id = l.actor_user_id
            LEFT JOIN admin_users au ON au.id = l.actor_admin_user_id
            WHERE l.id = ?
            LIMIT 1
            """,
            (int(log_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("audit_log_not_found")
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_audit_log_detail",
            target_type="audit_log",
            target_id=str(log_id),
            detail={},
        )
        conn.commit()
        return {
            "id": int(row["id"] or 0),
            "actor_type": str(row["actor_type"] or ""),
            "actor_user_id": int(row["actor_user_id"] or 0) or None,
            "actor_admin_user_id": int(row["actor_admin_user_id"] or 0) or None,
            "actor_display_name": _admin_actor_name(row),
            "action": str(row["action"] or ""),
            "target_type": str(row["target_type"] or ""),
            "target_id": str(row["target_id"] or ""),
            "detail_json": _parse_detail_json(row["detail_json"]),
            "ip": str(row["ip"] or ""),
            "user_agent": str(row["user_agent"] or ""),
            "created_at": str(row["created_at"] or ""),
        }
    finally:
        conn.close()


def list_admin_meetings(token: str, *, keyword: str = "") -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    keyword_text = str(keyword or "").strip().lower()
    rows = list_meetings()
    if keyword_text:
        rows = [
            item
            for item in rows
            if keyword_text in str(item.get("topic") or "").lower()
            or keyword_text in str(item.get("summary") or "").lower()
            or keyword_text in str(item.get("owner") or "").lower()
        ]
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_meeting_list",
            target_type="meeting",
            target_id="all",
            detail={"keyword_used": bool(keyword_text), "count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def create_admin_meeting(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _ensure_admin_meeting_mutation_enabled()
    _require_admin_write_role(admin_user)
    created = create_meeting(payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_meeting_create",
            target_type="meeting",
            target_id=str(created.get("id") or ""),
            detail={"topic": str(created.get("topic") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return created


def update_admin_meeting(token: str, meeting_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _ensure_admin_meeting_mutation_enabled()
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute("SELECT id FROM meetings WHERE id = ? LIMIT 1", (int(meeting_id or 0),)).fetchone()
        if row is None:
            raise ValueError("meeting_not_found")
        photo_urls = _normalize_photo_urls(
            payload.get("photo_urls"),
            allowed_scenes={"meeting"},
            allow_external=False,
            strict_scene=True,
        )
        conn.execute(
            """
            UPDATE meetings
            SET topic = ?, summary = ?, owner = ?, meeting_time = ?, photo_urls = ?
            WHERE id = ?
            """,
            (
                payload["topic"],
                payload["summary"],
                payload["owner"],
                payload["meeting_time"],
                _json_text(photo_urls),
                int(meeting_id or 0),
            ),
        )
        updated = conn.execute(
            "SELECT id, topic, summary, owner, meeting_time, photo_urls, created_at FROM meetings WHERE id = ? LIMIT 1",
            (int(meeting_id or 0),),
        ).fetchone()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_meeting_update",
            target_type="meeting",
            target_id=str(meeting_id),
            detail={"topic": str(payload.get("topic") or "")},
        )
        conn.commit()
        return _meeting_to_dict(updated)
    finally:
        conn.close()


def delete_admin_meeting(token: str, meeting_id: int) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _ensure_admin_meeting_mutation_enabled()
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute("SELECT id, topic FROM meetings WHERE id = ? LIMIT 1", (int(meeting_id or 0),)).fetchone()
        if row is None:
            raise ValueError("meeting_not_found")
        conn.execute("DELETE FROM meetings WHERE id = ?", (int(meeting_id or 0),))
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_meeting_delete",
            target_type="meeting",
            target_id=str(meeting_id),
            detail={"topic": str(row["topic"] or "")},
        )
        conn.commit()
        return {"status": "deleted", "message": "会议记录已删除"}
    finally:
        conn.close()


def list_admin_volunteer_events(token: str, *, keyword: str = "") -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    keyword_text = str(keyword or "").strip().lower()
    rows = list_volunteer_events()
    if keyword_text:
        rows = [
            item
            for item in rows
            if keyword_text in str(item.get("title") or "").lower()
            or keyword_text in str(item.get("detail") or "").lower()
            or keyword_text in str(item.get("created_by") or "").lower()
        ]
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_event_list",
            target_type="volunteer_event",
            target_id="all",
            detail={"keyword": keyword_text, "count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def create_admin_volunteer_event(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    created = create_volunteer_event(payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_event_create",
            target_type="volunteer_event",
            target_id=str(created.get("id") or ""),
            detail={"title": str(created.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return created


def update_admin_volunteer_event(token: str, event_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute("SELECT id FROM volunteer_events WHERE id = ? LIMIT 1", (int(event_id or 0),)).fetchone()
        if row is None:
            raise ValueError("event_not_found")
        photo_urls = _normalize_photo_urls(
            payload.get("photo_urls"),
            allowed_scenes={"volunteer"},
            allow_external=False,
            strict_scene=True,
        )
        conn.execute(
            """
            UPDATE volunteer_events
            SET title = ?, detail = ?, event_time = ?, quota = ?, created_by = ?, photo_urls = ?
            WHERE id = ?
            """,
            (
                payload["title"],
                payload["detail"],
                payload["event_time"],
                payload["quota"],
                payload["created_by"],
                _json_text(photo_urls),
                int(event_id or 0),
            ),
        )
        updated = conn.execute(
            """
            SELECT
                e.id, e.title, e.detail, e.event_time, e.quota, e.created_by, e.photo_urls, e.created_at,
                COUNT(s.id) AS signup_count
            FROM volunteer_events e
            LEFT JOIN volunteer_signups s ON s.event_id = e.id
            WHERE e.id = ?
            GROUP BY e.id
            """,
            (int(event_id or 0),),
        ).fetchone()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_event_update",
            target_type="volunteer_event",
            target_id=str(event_id),
            detail={"title": str(payload.get("title") or "")},
        )
        conn.commit()
        return _volunteer_event_to_dict(updated)
    finally:
        conn.close()


def delete_admin_volunteer_event(token: str, event_id: int) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT id, title FROM volunteer_events WHERE id = ? LIMIT 1",
            (int(event_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("event_not_found")
        conn.execute("DELETE FROM volunteer_signups WHERE event_id = ?", (int(event_id or 0),))
        conn.execute("DELETE FROM volunteer_events WHERE id = ?", (int(event_id or 0),))
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_event_delete",
            target_type="volunteer_event",
            target_id=str(event_id),
            detail={"title": str(row["title"] or "")},
        )
        conn.commit()
        return {"status": "deleted", "message": "志愿活动已删除"}
    finally:
        conn.close()


def list_admin_volunteer_signups(token: str, event_id: int) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        event_row = conn.execute(
            "SELECT id FROM volunteer_events WHERE id = ? LIMIT 1",
            (int(event_id or 0),),
        ).fetchone()
        if event_row is None:
            raise ValueError("event_not_found")
        rows = conn.execute(
            """
            SELECT
                s.id,
                s.event_id,
                s.participant_name,
                COALESCE(u.phone, '') AS participant_phone,
                COALESCE(u.village_group, '') AS village_group,
                s.created_at
            FROM volunteer_signups s
            LEFT JOIN users u ON u.openid = s.participant_openid
            WHERE s.event_id = ?
            ORDER BY s.created_at ASC, s.id ASC
            """,
            (int(event_id or 0),),
        ).fetchall()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_signups_view",
            target_type="volunteer_event",
            target_id=str(event_id),
            detail={"count": len(rows)},
        )
        conn.commit()
        return [
            {
                **dict(row),
                "participant_phone": mask_phone(
                    normalize_phone_number(str(row["participant_phone"] or "")) or str(row["participant_phone"] or "")
                ),
            }
            for row in rows
        ]
    finally:
        conn.close()


def export_admin_volunteer_signups_csv(token: str, event_id: int) -> bytes:
    admin_user = get_current_admin(token)
    conn = get_connection()
    try:
        event_row = conn.execute(
            "SELECT id, title, event_time, created_by FROM volunteer_events WHERE id = ? LIMIT 1",
            (int(event_id or 0),),
        ).fetchone()
        if event_row is None:
            raise ValueError("event_not_found")
        rows = conn.execute(
            """
            SELECT
                s.participant_name,
                COALESCE(u.phone, '') AS participant_phone,
                COALESCE(u.village_group, '') AS village_group,
                s.created_at
            FROM volunteer_signups s
            LEFT JOIN users u ON u.openid = s.participant_openid
            WHERE s.event_id = ?
            ORDER BY s.created_at ASC, s.id ASC
            """,
            (int(event_id or 0),),
        ).fetchall()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["活动标题", "活动时间", "发布人", "报名姓名", "联系电话", "所属村组", "报名时间"])
        for row in rows:
            writer.writerow(
                [
                    str(event_row["title"] or ""),
                    str(event_row["event_time"] or ""),
                    str(event_row["created_by"] or ""),
                    str(row["participant_name"] or ""),
                    mask_phone(
                        normalize_phone_number(str(row["participant_phone"] or "")) or str(row["participant_phone"] or "")
                    ),
                    str(row["village_group"] or ""),
                    str(row["created_at"] or ""),
                ]
            )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_signups_export",
            target_type="volunteer_event",
            target_id=str(event_id),
            detail={"count": len(rows)},
        )
        conn.commit()
        return output.getvalue().encode("utf-8-sig")
    finally:
        conn.close()


def list_admin_training_programs(token: str, *, keyword: str = "") -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_training_programs(keyword=keyword)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_program_list",
            target_type="training_program",
            target_id="all",
            detail={"keyword": str(keyword or "").strip(), "count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def create_admin_training_program(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    created = create_training_program(payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_program_create",
            target_type="training_program",
            target_id=str(created.get("id") or ""),
            detail={"title": str(created.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return created


def update_admin_training_program(token: str, program_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    updated = update_training_program(program_id, payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_program_update",
            target_type="training_program",
            target_id=str(program_id),
            detail={"title": str(updated.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return updated


def end_admin_training_program(token: str, program_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    updated = end_training_program(program_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_program_end",
            target_type="training_program",
            target_id=str(program_id),
            detail={"title": str(updated.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return updated


def delete_admin_training_program(token: str, program_id: int) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    row = delete_training_program(program_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_program_delete",
            target_type="training_program",
            target_id=str(program_id),
            detail={},
        )
        conn.commit()
    finally:
        conn.close()
    return row


def list_admin_training_signups(token: str, program_id: int) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_training_signups_by_program(program_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_signups_view",
            target_type="training_program",
            target_id=str(program_id),
            detail={"count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def export_admin_training_signups_csv(token: str, program_id: int) -> bytes:
    admin_user = get_current_admin(token)
    content = export_training_signups_csv(program_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_training_signups_export",
            target_type="training_program",
            target_id=str(program_id),
            detail={},
        )
        conn.commit()
    finally:
        conn.close()
    return content


def list_admin_volunteer_activities(token: str, *, keyword: str = "") -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_volunteer_activities(keyword=keyword)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_list",
            target_type="volunteer_activity",
            target_id="all",
            detail={"keyword": str(keyword or "").strip(), "count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def create_admin_volunteer_activity(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    created = create_volunteer_activity(payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_create",
            target_type="volunteer_activity",
            target_id=str(created.get("id") or ""),
            detail={"title": str(created.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return created


def update_admin_volunteer_activity(token: str, activity_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    updated = update_volunteer_activity(activity_id, payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_update",
            target_type="volunteer_activity",
            target_id=str(activity_id),
            detail={"title": str(updated.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return updated


def end_admin_volunteer_activity(token: str, activity_id: int) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    updated = end_volunteer_activity(activity_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_end",
            target_type="volunteer_activity",
            target_id=str(activity_id),
            detail={"title": str(updated.get("title") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return updated


def delete_admin_volunteer_activity(token: str, activity_id: int) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    row = delete_volunteer_activity(activity_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_delete",
            target_type="volunteer_activity",
            target_id=str(activity_id),
            detail={},
        )
        conn.commit()
    finally:
        conn.close()
    return row


def list_admin_volunteer_activity_signups(token: str, activity_id: int) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_volunteer_activity_signups_by_activity(activity_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_signups_view",
            target_type="volunteer_activity",
            target_id=str(activity_id),
            detail={"count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def export_admin_volunteer_activity_signups_csv(token: str, activity_id: int) -> bytes:
    admin_user = get_current_admin(token)
    content = export_volunteer_activity_signups_csv(activity_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_volunteer_activity_signups_export",
            target_type="volunteer_activity",
            target_id=str(activity_id),
            detail={},
        )
        conn.commit()
    finally:
        conn.close()
    return content


def list_admin_proposals(token: str) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    rows = list_proposals()
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_proposal_list",
            target_type="proposal",
            target_id="all",
            detail={"count": len(rows)},
        )
        conn.commit()
    finally:
        conn.close()
    return rows


def list_admin_local_products(token: str) -> list[dict[str, Any]]:
    get_current_admin(token)
    return list_local_products()


def create_admin_local_product(token: str, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    created = create_local_product(payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_local_product_create",
            target_type="local_product",
            target_id=str(created.get("id") or ""),
            detail={"name": str(created.get("name") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return created


def update_admin_local_product(token: str, product_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    updated = update_local_product(product_id, payload)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_local_product_update",
            target_type="local_product",
            target_id=str(product_id),
            detail={"name": str(updated.get("name") or "")},
        )
        conn.commit()
    finally:
        conn.close()
    return updated


def delete_admin_local_product(token: str, product_id: int) -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_write_role(admin_user)
    row = delete_local_product(product_id)
    conn = get_connection()
    try:
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_local_product_delete",
            target_type="local_product",
            target_id=str(product_id),
            detail={},
        )
        conn.commit()
    finally:
        conn.close()
    return row


def list_admin_villager_identities(
    token: str,
    *,
    keyword: str = "",
    limit: int = 200,
) -> list[dict[str, Any]]:
    admin_user = get_current_admin(token)
    keyword_text = str(keyword or "").strip()

    conn = get_connection()
    try:
        query, params = _build_admin_villager_identity_query(keyword_text)
        query += " ORDER BY updated_at DESC, id DESC LIMIT ?"
        params.append(max(1, min(int(limit or 200), 500)))
        rows = conn.execute(query, tuple(params)).fetchall()
        if len(rows) >= 50:
            security_logger.warning(
                "security_alert event=admin_bulk_identity_view admin_user_id=%s count=%s",
                int(admin_user["id"] or 0),
                len(rows),
            )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_villager_identity_list",
            target_type="user",
            target_id="all",
            detail={"keyword": keyword_text, "count": len(rows)},
        )
        conn.commit()
        return [_serialize_admin_villager_identity(row) for row in rows]
    finally:
        conn.close()


def get_admin_villager_identity_summary(token: str, *, keyword: str = "") -> dict[str, int]:
    admin_user = get_current_admin(token)
    keyword_text = str(keyword or "").strip()

    conn = get_connection()
    try:
        query, params = _build_admin_villager_identity_query(keyword_text)
        query += " ORDER BY updated_at DESC, id DESC"
        rows = conn.execute(query, tuple(params)).fetchall()
        records = [_serialize_admin_villager_identity(row) for row in rows]
        summary = {
            "total_count": len(records),
            "matched_count": sum(1 for row in records if row["roster_match_status"] == "matched"),
            "grouped_count": sum(1 for row in records if str(row["village_group"] or "").strip()),
            "turning_60_this_year_count": sum(1 for row in records if row["is_turning_60_this_year"]),
            "age_60_plus_count": sum(1 for row in records if row["is_age_60_plus"]),
            "age_80_plus_count": sum(1 for row in records if row["is_age_80_plus"]),
        }
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_villager_identity_summary",
            target_type="user",
            target_id="all",
            detail={"keyword_used": bool(keyword_text), **summary},
        )
        conn.commit()
        return summary
    finally:
        conn.close()


def export_admin_villager_identities_xlsx(token: str, *, keyword: str = "") -> bytes:
    admin_user = get_current_admin(token)
    keyword_text = str(keyword or "").strip()

    conn = get_connection()
    try:
        query, params = _build_admin_villager_identity_query(keyword_text)
        query += " ORDER BY updated_at DESC, id DESC"
        rows = conn.execute(query, tuple(params)).fetchall()
        records = [_serialize_admin_villager_identity(row) for row in rows]

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "村民脱敏名册"
        headers = [
            "姓名",
            "性别",
            "身份证号（脱敏）",
            "手机号（脱敏）",
            "村组",
            "年龄",
            "实名状态",
            "名单状态",
            "最后更新时间",
        ]
        sheet.append(headers)
        for row_index, row in enumerate(records, start=2):
            sheet.append(
                [
                    row["real_name"],
                    row["gender"],
                    row["id_card_masked"],
                    _format_export_phone_value("", row["phone_masked"]),
                    row["village_group"] or "",
                    row["age"] if row["age"] is not None else "",
                    {
                        "unverified": "未完成",
                        "partial": "进行中",
                        "verified": "已完成",
                    }.get(row["auth_status"], row["auth_status"] or "-"),
                    {
                        "unmatched": "未匹配",
                        "matched": "已匹配",
                        "not_configured": "未配置名单",
                    }.get(row["roster_match_status"], row["roster_match_status"] or "-"),
                    _format_export_datetime(row["updated_at"]),
                ]
            )
            sheet.cell(row=row_index, column=3).number_format = "@"
            sheet.cell(row=row_index, column=4).number_format = "@"
            sheet.cell(row=row_index, column=9).number_format = "yyyy-mm-dd hh:mm"

        header_fill = PatternFill("solid", fgColor="B93A2E")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(bottom=Side(style="thin", color="E6DDD4"))
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30

        for row in sheet.iter_rows(min_row=2, max_col=len(headers)):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            row[0].alignment = Alignment(horizontal="left", vertical="center")
            row[4].alignment = Alignment(horizontal="left", vertical="center")
            sheet.row_dimensions[row[0].row].height = 24

        column_widths = {
            "A": 12,
            "B": 8,
            "C": 24,
            "D": 18,
            "E": 20,
            "F": 10,
            "G": 12,
            "H": 14,
            "I": 20,
        }
        for column, width in column_widths.items():
            sheet.column_dimensions[column].width = width

        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.print_title_rows = "1:1"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        output = io.BytesIO()
        workbook.save(output)
        security_logger.warning(
            "security_alert event=admin_masked_identity_export admin_user_id=%s count=%s",
            int(admin_user["id"] or 0),
            len(records),
        )
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_villager_identity_export",
            target_type="user",
            target_id="all",
            detail={"keyword_used": bool(keyword_text), "count": len(records)},
        )
        conn.commit()
        return output.getvalue()
    finally:
        conn.close()


def unbind_admin_villager_identity(token: str, user_id: int, *, reason: str = "") -> dict[str, str]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)

    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT id, real_name, phone_masked, phone_verified_at, face_status
            FROM users
            WHERE id = ? AND role = 'villager'
            LIMIT 1
            """,
            (int(user_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("villager_identity_not_found")
        now = _now_iso()
        auth_status = calc_auth_status(
            bool(str(row["phone_verified_at"] or "").strip()),
            "unsubmitted",
            str(row["face_status"] or "unsubmitted"),
        )
        conn.execute(
            """
            UPDATE users
            SET
                real_name = '',
                gender = '',
                id_card_masked = '',
                id_card_last4 = '',
                id_card_hash = '',
                id_card_encrypted = '',
                realname_status = 'unsubmitted',
                auth_status = ?,
                roster_match_status = 'unmatched',
                village_group = '',
                updated_at = ?
            WHERE id = ?
            """,
            (auth_status, now, int(user_id or 0)),
        )
        _sync_villager_registry_conn(conn, int(user_id or 0))
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_villager_identity_unbind",
            target_type="user",
            target_id=str(user_id),
            detail={
                "real_name": str(row["real_name"] or ""),
                "phone_masked": str(row["phone_masked"] or ""),
                "reason": str(reason or "").strip(),
            },
        )
        conn.commit()
        return {"status": "ok", "message": "已解绑该村民的基础实名登记"}
    finally:
        conn.close()


def rebind_admin_villager_identity(
    token: str,
    user_id: int,
    *,
    real_name: str,
    id_card_number: str,
) -> dict[str, Any]:
    admin_user = get_current_admin(token)
    _require_admin_super_role(admin_user)

    name_text = str(real_name or "").strip()
    id_card_text = _normalize_id_card(id_card_number)
    if not name_text:
        raise ValueError("invalid_real_name")
    if not is_valid_id_card(id_card_text):
        raise ValueError("invalid_id_card")

    conn = get_connection()
    try:
        row = conn.execute(
            """
            SELECT id, gender, phone_verified_at, face_status
            FROM users
            WHERE id = ? AND role = 'villager'
            LIMIT 1
            """,
            (int(user_id or 0),),
        ).fetchone()
        if row is None:
            raise ValueError("villager_identity_not_found")

        id_card_hash = _id_hash(id_card_text)
        _require_active_roster_conn(conn)
        match_row = _find_roster_match_conn(conn, name_text, id_card_hash)
        if match_row is None:
            raise ValueError("identity_not_in_roster")
        conflict_row = _find_existing_identity_binding_conn(
            conn,
            id_card_hash,
            exclude_user_id=int(user_id or 0),
        )
        if conflict_row is not None:
            raise ValueError("id_card_already_verified")

        gender = _resolve_gender(str(row["gender"] or ""), id_card_text)
        roster_match_status = "matched"
        village_group = str(match_row["group_name"] or "")
        roster_match_id = str(match_row["id"] or "")

        now = _now_iso()
        auth_status = calc_auth_status(
            bool(str(row["phone_verified_at"] or "").strip()),
            "verified",
            str(row["face_status"] or "unsubmitted"),
        )
        conn.execute(
            """
            UPDATE users
            SET
                real_name = ?,
                gender = ?,
                id_card_masked = ?,
                id_card_last4 = ?,
                id_card_hash = ?,
                id_card_encrypted = ?,
                realname_status = 'verified',
                auth_status = ?,
                roster_match_status = ?,
                village_group = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                name_text,
                gender,
                mask_id_card(id_card_text),
                "",
                id_card_hash,
                encrypt_id_card_number(id_card_text),
                auth_status,
                roster_match_status,
                village_group,
                now,
                int(user_id or 0),
            ),
        )
        _sync_villager_registry_conn(conn, int(user_id or 0))
        updated_row = conn.execute(
            """
            SELECT
                id,
                real_name,
                phone_masked,
                id_card_masked,
                auth_status,
                realname_status,
                roster_match_status,
                village_group,
                updated_at
            FROM users
            WHERE id = ?
            LIMIT 1
            """,
            (int(user_id or 0),),
        ).fetchone()
        _audit_admin_action_conn(
            conn,
            admin_user_id=int(admin_user["id"] or 0),
            action="admin_villager_identity_rebind",
            target_type="user",
            target_id=str(user_id),
            detail={
                "real_name": name_text,
                "roster_match_status": roster_match_status,
                "roster_match_id": roster_match_id,
                "village_group": village_group,
            },
        )
        conn.commit()
        return {
            "user_id": int(updated_row["id"] or 0),
            "real_name": str(updated_row["real_name"] or ""),
            "phone_masked": str(updated_row["phone_masked"] or ""),
            "id_card_masked": str(updated_row["id_card_masked"] or ""),
            "auth_status": str(updated_row["auth_status"] or "unverified"),
            "realname_status": str(updated_row["realname_status"] or "unsubmitted"),
            "roster_match_status": str(updated_row["roster_match_status"] or "unmatched"),
            "village_group": str(updated_row["village_group"] or ""),
            "updated_at": str(updated_row["updated_at"] or ""),
        }
    finally:
        conn.close()
